"""WH (D57, D60, D61) and WH Serials (D62)."""
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from common import (AL_C, AL_L, AL_LW, C_CALC, C_CALC_HDR, C_HDR, C_INPUT, C_RED_BG, C_RED_FG, Col, DATE_FMT,
                    PCT_FMT, T, fill, font, register, style, title, tref, valid)

# item definitions (structure of the old WH: 21 TE items + 28 CW items)
_ITEMS = [
    # group, name, code, uom, tracking, progress code, as-built item, need unit
    ("WL", "cnWave 28 GHz BTS", "C280500A101A", "pcs", "Serial", "BTS", "Sector Qty", "Cambium Sector"),
    ("WL", "DC LPU Kit", "C000000L114A", "pcs", "Qty", None, None, "LPU Kit"),
    ("WL", "WiBAS G5 dual-BS", "WG5-DUAL-BS", "pcs", "Serial", "BTS", "Sector Qty", "Intracom Sector"),
    ("WL", "WiBAS Antenna", "HANT262890", "pcs", "Qty", None, None, "Intracom Sector"),
    ("CPE", "cnWave C100", "C280500C001A", "pcs", "Serial", "CPE", None, None),
    ("CPE", "WiBAS G5 Connect+", "WG5-CONN-PLUS", "pcs", "Serial", "CPE", None, None),
    ("MW", "UltraLink-GX80 Radio", "KZ00-65F", "pcs", "Serial", "PTP", None, "MW Half-link"),
    ("MW", "MW Antenna ANT-7080-60P", "KH14-056", "pcs", "Qty", None, "MW Antenna Qty", "MW Half-link"),
    ("IP", "Router S9500-30XS", "S9500", "pcs", "Serial", "ROUTER", "Router Qty", "Router"),
    ("IP", "Router S9501-18SMT", "S9501", "pcs", "Serial", "ROUTER", "Router Qty", "Router"),
    ("IP", "Router S9700-53DX", "S9700", "pcs", "Serial", "ROUTER", "Router Qty", "Router"),
    ("IP", "SFP-SM-10G-10KM", None, "pcs", "Qty", None, None, None),
    ("PWR", "Cabinet MTS9514A", "MTS9514A", "pcs", "Serial", "CABINET", None, "Cabinet"),
    ("PWR", "Li Battery SLB48-300-147", "800-00211", "pcs", "Serial", "BATTERY", "Battery Qty", "Battery"),
    ("CORE", "Server Dell R450", "R450", "pcs", "Serial", "CORE-SERVER", None, None),
    ("CORE", "Server Dell R650", "R650", "pcs", "Serial", "CORE-SERVER", None, None),
    ("Accessories", "ERKA Shielded Outdoor 2x2.5 mm²", None, "m", "Qty", None, None, None),
    ("Accessories", "ERKA Shielded Outdoor 4x0.75 mm²", None, "m", "Qty", None, None, None),
    ("Accessories", "Fiber Optic Cable Outdoor", None, "m", "Qty", None, None, None),
    ("Accessories", "Telecom Bracket", None, "pcs", "Qty", None, None, None),
    ("Accessories", "Fiber Cable Clamp 4 Layers 8 Runs", None, "pcs", "Qty", None, None, None),
    ("CW", "Tower", None, "m", "Qty", None, "Tower (m)", None),
    ("CW", "Mounting Pole", None, "m", "Qty", None, "Mounting Pole (m)", None),
    ("CW", "AC Cable 4x16 mm²", None, "m", "Qty", None, "AC Cable 4x16 (m)", None),
    ("CW", "AC Cable 4x10 mm²", None, "m", "Qty", None, "AC Cable 4x10 (m)", None),
    ("CW", "DC Cable 25 mm²", None, "m", "Qty", None, "DC Cable 25 (m)", None),
    ("CW", "DC Cable 35 mm²", None, "m", "Qty", None, "DC Cable 35 (m)", None),
    ("CW", "Fiber Cable – Patch Cord", None, "m", "Qty", None, "Fiber Patch Cord", None),
    ("CW", "Fiber Cable – Outdoor Shielded", None, "m", "Qty", None, "Outdoor Shield Fiber (m)", None),
    ("CW", "Indoor Cable Tray", None, "m", "Qty", None, "Indoor Cable Tray (m)", None),
    ("CW", "Outdoor Cable Tray", None, "m", "Qty", None, "Outdoor Cable Tray (m)", None),
    ("CW", 'PVC Pipe 1"', None, "m", "Qty", None, 'PVC 1" (m)', None),
    ("CW", 'PVC Pipe 1.5"', None, "m", "Qty", None, 'PVC 1.5" (m)', None),
    ("CW", 'PVC Pipe 2"', None, "m", "Qty", None, 'PVC 2" (m)', None),
    ("CW", "GND Cable 16 mm²", None, "m", "Qty", None, "GND 16 (m)", None),
    ("CW", "GND Cable 35 mm²", None, "m", "Qty", None, "GND 35 (m)", None),
    ("CW", "GND Cable 50 mm²", None, "m", "Qty", None, "GND 50 (m)", None),
    ("CW", "Grounding Busbar", None, "pcs", "Qty", None, "Busbar (pcs)", None),
    ("CW", "Grounding Copper Rod", None, "pcs", "Qty", None, "Copper Rod (pcs)", None),
    ("CW", "Grounding Copper-plated Rod", None, "pcs", "Qty", None, "Copper-plated Rod (pcs)", None),
    ("CW", "Electrical Switch 63A 3P AC", None, "pcs", "Qty", None, "63A 3P AC (pcs)", None),
    ("CW", "Electrical Switch 63A 1P DC", None, "pcs", "Qty", None, "63A 1P DC (pcs)", None),
    ("CW", "Electrical Switch 32A 1P DC", None, "pcs", "Qty", None, "32A 1P DC (pcs)", None),
    ("CW", "Electrical Switch 16A 1P DC", None, "pcs", "Qty", None, "16A 1P DC (pcs)", None),
    ("CW", "Electrical Switch 10A 1P DC", None, "pcs", "Qty", None, "10A 1P DC (pcs)", None),
    ("CW", "AC Box", None, "pcs", "Qty", None, "AC Box (pcs)", None),
    ("CW", "DC Box", None, "pcs", "Qty", None, "DC Box (pcs)", None),
    ("CW", "AC Box Stand", None, "pcs", "Qty", None, "AC Box Stand (pcs)", None),
    ("CW", "Indoor Rack", None, "pcs", "Qty", None, "Indoor Rack (pcs)", None),
]

ITEMS = [dict(group=g, name=n, code=c, uom=u, track=t, prog=p, ab=a, unit=nu, nm="cMov_I%02d" % (i + 1))
         for i, (g, n, c, u, t, p, a, nu) in enumerate(_ITEMS)]

HDR_ROWS = [("Group", "ItemCat", "WHGroup"), ("Item Name", None, "WHItemNames"),
            ("Manufacturer", None, "WHManufacturer"), ("Item Code", None, "WHItemCodes"),
            ("UoM", "UoM", "WHUoM"), ("Tracking", "TrackType", "WHTracking"),
            ("Contract Qty", ("decimal",), "WHContractQty"), ("Planned Delivery Date", ("date",), "WHPlannedDelivery"),
            ("Progress Item Code", None, "WHProgCode"), ("As-Built Item", "AsBuiltItems", "WHAsBuilt"),
            ("Need Unit", "NeedUnit", "WHNeedUnit")]
SUM_ROWS = ["Received", "Delivered %", "Issued", "Returned Good", "Faulty / Damaged", "On Hand", "Reserved",
            "Available", "Need", "Shortfall"]
SUM_NAMES = ["WHReceived", "WHDeliveredPct", "WHIssued", "WHReturnedGood", "WHFaulty", "WHOnHand", "WHReserved",
             "WHAvailable", "WHNeed", "WHShortfall"]
SUM_NOTE = ["PL + OPEN + ADJ rows (not Cancelled)", "Received ÷ Contract Qty", "DN rows, Status Dispatched",
            "RN rows, Received, Good", "RN rows, Received, Faulty / Damaged (not in On Hand)",
            "Received − Issued + Returned Good", "DN rows, Status Approved", "On Hand − Reserved",
            "TE need of the design (not cancelled sites) not yet dispatched", "Need − Available"]

I0 = 12                      # first item column (L)
R_HDR0 = 2                   # first header row
R_SUM0 = 14                  # first summary row
R_MOV_SUB, R_MOV_HDR = 26, 27
MOV_ROWS = 20

UNIT_NEED = ('CHOOSE(MATCH({u},NeedUnit,0),SUMIFS(SC_Cam,Master_FZ,"<>Cancelled"),SUMIFS(SC_Int,Master_FZ,'
             '"<>Cancelled"),SUMIFS(SC_MW,Master_FZ,"<>Cancelled"),COUNTIFS(SC_TE,"Yes",Master_FZ,"<>Cancelled"),'
             'COUNTIFS(SC_TE,"Yes",Master_FZ,"<>Cancelled"),2*COUNTIFS(SC_TE,"Yes",Master_FZ,"<>Cancelled"),'
             '2*COUNTIFS(SC_TE,"Yes",Master_FZ,"<>Cancelled"))')


def build_wh(wb, names):
    ws = wb.create_sheet("WH")
    title(ws, "A1", "WAREHOUSE · single main store · one column per item · stock summary and one movement table "
                    "(PL / OPEN / ADJ / DN / RN) · serial numbers in WH Serials")
    nI = len(ITEMS)
    lastI = I0 + nI - 1
    LI, LL = L(I0), L(lastI)
    # header rows
    for k, (lab, dv, nm) in enumerate(HDR_ROWS):
        r = R_HDR0 + k
        style(ws.cell(r, 1, k + 1), font(9, True), fill(C_HDR), AL_C)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=I0 - 1)
        style(ws.cell(r, 2, lab), font(9, True), fill(C_HDR), AL_L)
        for c in range(3, I0):
            style(ws.cell(r, c), None, fill(C_HDR), AL_L)
        names.add(nm, f"WH!${LI}${r}:${LL}${r}")
        if dv:
            if isinstance(dv, tuple):
                v = (DataValidation(type="date", operator="between", formula1="DateMin", formula2="DateMax",
                                    allow_blank=True) if dv[0] == "date" else
                     DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True))
            else:
                v = DataValidation(type="list", formula1="=" + dv, allow_blank=True)
            ws.add_data_validation(v)
            v.add(f"{LI}{r}:{LL}{r}")
    vals = {0: "group", 1: "name", 3: "code", 4: "uom", 5: "track", 8: "prog", 9: "ab", 10: "unit"}
    for j, it in enumerate(ITEMS):
        c = I0 + j
        for k in range(len(HDR_ROWS)):
            r = R_HDR0 + k
            v = it[vals[k]] if k in vals else None
            cell = ws.cell(r, c, v)
            yellow = k in (6, 7)
            style(cell, font(9, k == 1), fill(C_INPUT) if yellow else fill(C_HDR), AL_C,
                  DATE_FMT if k == 7 else "General")
        ws.column_dimensions[L(c)].width = 11
    ws.row_dimensions[3].height = 42
    # summary
    style(ws.cell(R_SUM0 - 1, 1, "STOCK SUMMARY · per item"), font(10, True), None, AL_L, border=False)
    for k, lab in enumerate(SUM_ROWS):
        r = R_SUM0 + k
        style(ws.cell(r, 1, k + 1), font(9, True), fill(C_CALC_HDR), AL_C)
        style(ws.cell(r, 2, lab), font(9, True), fill(C_CALC_HDR), AL_L)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=I0 - 1)
        style(ws.cell(r, 3, SUM_NOTE[k]), font(8, italic=True), fill(C_CALC_HDR), AL_L)
        for c in range(4, I0):
            style(ws.cell(r, c), None, fill(C_CALC_HDR))
        names.add(SUM_NAMES[k], f"WH!${LI}${r}:${LL}${r}")
    rC = R_HDR0 + 6
    rU = R_HDR0 + 10
    for j, it in enumerate(ITEMS):
        c = I0 + j
        X = L(c)
        nm = it["nm"]

        def S(*crit):
            return "SUMIFS(" + nm + "," + ",".join(crit) + ")"
        f = {
            0: "=" + "+".join(S('cMov_Type,"' + t + '"', 'cMov_Status,"<>Cancelled"') for t in ("PL", "OPEN", "ADJ")),
            1: f'=IF(N({X}{rC})=0,"n/a",{X}{R_SUM0}/{X}{rC})',
            2: "=" + "+".join(S('cMov_Type,"' + t + '"', 'cMov_Status,"Dispatched"') for t in ("TE DN", "CW DN")),
            3: "=" + "+".join(S('cMov_Type,"' + t + '"', 'cMov_Status,"Received"', 'cMov_Condition,"Good"')
                              for t in ("TE RN", "CW RN")),
            4: "=" + "+".join(S('cMov_Type,"' + t + '"', 'cMov_Status,"Received"', 'cMov_Condition,"Faulty / Damaged"')
                              for t in ("TE RN", "CW RN")),
            5: f"={X}{R_SUM0}-{X}{R_SUM0+2}+{X}{R_SUM0+3}",
            6: "=" + "+".join(S('cMov_Type,"' + t + '"', 'cMov_Status,"Approved"') for t in ("TE DN", "CW DN")),
            7: f"={X}{R_SUM0+5}-{X}{R_SUM0+6}",
        }
        netd = (f'IF({X}${rU}="Router",SUMIF(${LI}${rU}:${LL}${rU},"Router",${LI}${R_SUM0+2}:${LL}${R_SUM0+2})'
                f'-SUMIF(${LI}${rU}:${LL}${rU},"Router",${LI}${R_SUM0+3}:${LL}${R_SUM0+3})-SUMIF(${LI}${rU}:${LL}${rU},'
                f'"Router",${LI}${R_SUM0+4}:${LL}${R_SUM0+4}),{X}{R_SUM0+2}-{X}{R_SUM0+3}-{X}{R_SUM0+4})')
        f[8] = (f'=IF({X}${rU}="","n/a",IFERROR(MAX(0,' + UNIT_NEED.format(u=f"{X}${rU}") + '-' + netd + '),"n/a"))')
        f[9] = (f'=IF(ISNUMBER({X}{R_SUM0+8}),MAX(0,{X}{R_SUM0+8}-IF({X}${rU}="Router",SUMIF(${LI}${rU}:${LL}${rU},'
                f'"Router",${LI}${R_SUM0+7}:${LL}${R_SUM0+7}),{X}{R_SUM0+7})),"n/a")')
        for k in range(10):
            style(ws.cell(R_SUM0 + k, c, f[k]), font(9), fill(C_CALC), AL_C, PCT_FMT if k == 1 else "General")
    ws.conditional_formatting.add(f"{LI}{R_SUM0+7}:{LL}{R_SUM0+7}",
                                  CellIsRule(operator="lessThan", formula=["0"],
                                             fill=PatternFill("solid", fgColor="FF" + C_RED_BG),
                                             font=Font(color="FF" + C_RED_FG)))
    ws.conditional_formatting.add(f"{LI}{R_SUM0+9}:{LL}{R_SUM0+9}",
                                  FormulaRule(formula=[f"AND(ISNUMBER({LI}{R_SUM0+9}),{LI}{R_SUM0+9}>0)"],
                                              fill=PatternFill("solid", fgColor="FF" + C_RED_BG),
                                              font=Font(color="FF" + C_RED_FG)))
    # movement table
    style(ws.cell(R_MOV_SUB - 1, 1, "MOVEMENTS · one row per document · one column per item"), font(10, True),
          None, AL_L, border=False)
    r0 = R_MOV_HDR + 1
    items_rng = "$" + LI + "{r}:$" + LL + "{r}"
    whole = f"(${LI}${R_HDR0+4}:${LL}${R_HDR0+4}<>\"m\")"
    ser = (f"SUMPRODUCT((${LI}${R_HDR0+5}:${LL}${R_HDR0+5}=\"Serial\")*({items_rng}<>(([[Type]]=\"PL\")*"
           f"COUNTIFS(cSer_ItemName,WHItemNames,cSer_PLInvoiceNo,[[Doc No.]])+ISNUMBER(SEARCH(\"DN\",[[Type]]))*"
           f"COUNTIFS(cSer_ItemName,WHItemNames,cSer_DNNo,[[Doc No.]])+ISNUMBER(SEARCH(\"RN\",[[Type]]))*"
           f"COUNTIFS(cSer_ItemName,WHItemNames,cSer_RNNo,[[Doc No.]]))))")
    check = ('=IF(COUNTIF([[Type]]:[[Doc No.]],"?*")+COUNT([[Date]])+COUNTIF([[Site ID]],"?*")+COUNTIF('
             '[[Status]]:[[Remark]],"?*")+COUNT(' + items_rng + ')=0,"",IF([[Type]]="","Type missing",'
             'IF([[Doc No.]]="","Doc No. missing",IF(COUNTIF(cMov_DocNo,[[Doc No.]])>1,"Duplicate Doc No.",'
             'IF(NOT(' + valid("[[Date]]") + '),"Date missing",IF([[Status]]="","Status missing",IF(AND(OR('
             'ISNUMBER(SEARCH("DN",[[Type]])),ISNUMBER(SEARCH("RN",[[Type]]))),[[Site ID]]=""),"Site missing",'
             'IF(AND([[Site ID]]<>"",COUNTIF(Master_B,[[Site ID]])=0),"Site not in Master",IF(COUNT(' + items_rng
             + ')=0,"No quantity",IF(AND([[Type]]<>"ADJ",COUNTIF(' + items_rng + ',"<0")>0),"Negative qty",'
             'IF(IFERROR(SUMPRODUCT((' + items_rng + '-INT(' + items_rng + ')<>0)*' + whole + ')>0,TRUE),'
             '"Fraction in whole unit",IF(OR([[Type]]="OPEN",[[Type]]="ADJ",[[Status]]="Approved",[[Status]]='
             '"Cancelled"),"OK",IF(IFERROR(' + ser + '>0,TRUE),"Serial count ≠ qty","OK")))))))))))))')
    cols = [
        Col("#", "num", T('=IF([[Type]]="","",ROW()-' + str(R_MOV_HDR) + ')'), w=5),
        Col("Type", dv="RefWHType", w=8),
        Col("Doc No.", w=12),
        Col("Date", dv=("date",), fmt="date"),
        Col("Site ID", dv="SiteIDs", w=11),
        Col("Site Name", "calc", T('=IF([[Site ID]]="","",IFERROR(INDEX(cRF_SiteName,MATCH([[Site ID]],cRF_SiteID,0))'
                                   '&"",""))'), w=16),
        Col("Status", dv="TxnStatus", w=10),
        Col("Condition", dv="Condition", w=10),
        Col("Reason", dv="ReturnReason", w=11),
        Col("Remark", w=16),
        Col("Check", "calc", T(check), w=14),
    ]
    for it in ITEMS:
        cols.append(Col(it["name"], dv=("decimal", -100000, 100000), w=11, nm=it["nm"]))
    register(ws, names, "tblWHMovements", "cMov", cols, first_col=1, sub_row=R_MOV_SUB, hdr_row=R_MOV_HDR,
             nrows=MOV_ROWS, header_h=42)
    # item subtotals: sums of visible rows
    last = R_MOV_HDR + MOV_ROWS
    for j in range(len(ITEMS)):
        X = L(I0 + j)
        ws.cell(R_MOV_SUB, I0 + j).value = f"=SUBTOTAL(109,{X}{r0}:{X}{last})"
    ws.freeze_panes = "L4"
    ws.sheet_view.zoomScale = 90
    return ws


def build_serials(wb, names):
    ws = wb.create_sheet("WH Serials")
    title(ws, "A1", "WH SERIALS · one row per serial number · site and dates read from the WH movement documents")
    style(ws.cell(2, 2, "Find Serial No."), font(9, True), fill(C_HDR), AL_C)
    style(ws.cell(2, 3), font(9), fill(C_INPUT), AL_C)
    lab = [("Item", 'IFERROR(INDEX(cSer_ItemName,MATCH($C$2,cSer_SerialNo,0))&"","")'),
           ("Site ID", 'IFERROR(INDEX(cSer_SiteID,MATCH($C$2,cSer_SerialNo,0))&"","")'),
           ("Status", 'IFERROR(INDEX(cSer_Status,MATCH($C$2,cSer_SerialNo,0))&"","")'),
           ("Installed", 'IFERROR(INDEX(cSer_Installed,MATCH($C$2,cSer_SerialNo,0)),"")')]
    for k, (t, f) in enumerate(lab):
        c = 5 + 2 * k
        style(ws.cell(2, c, t), font(9, True), fill(C_CALC_HDR), AL_C)
        style(ws.cell(2, c + 1, '=IF($C$2="","",' + f + ')'), font(9), fill(C_CALC), AL_C,
              DATE_FMT if t == "Installed" else "General")
    lk = 'IFERROR(MATCH([[{d}]],cMov_DocNo,0),0)'
    cols = [
        Col("#", "num", T('=IF([[Serial No.]]="","",ROW()-6)'), w=5),
        Col("Serial No.", w=16),
        Col("Item Name", dv="WHItemNames", w=18),
        Col("Item Code", "calc", T('=IF([[Item Name]]="","",IFERROR(INDEX(WHItemCodes,MATCH([[Item Name]],WHItemNames,'
                                   '0))&"",""))'), w=12),
        Col("PL / Invoice No.", w=12),
        Col("Pack / Box No.", w=10),
        Col("DN No.", w=11),
        Col("RN No.", w=11),
        Col("Site ID", "calc", T('=IF([[DN No.]]="","",IFERROR(INDEX(cMov_SiteID,MATCH([[DN No.]],cMov_DocNo,0))&"",""))'),
            w=11),
        Col("Site Name", "calc", T('=IF([[Site ID]]="","",IFERROR(INDEX(cRF_SiteName,MATCH([[Site ID]],cRF_SiteID,0))'
                                   '&"",""))'), w=16),
        Col("Dispatch Date", "calc",
            T('=IF([[DN No.]]="","",IFERROR(IF(INDEX(cMov_Status,MATCH([[DN No.]],cMov_DocNo,0))="Dispatched",'
              'INDEX(cMov_Date,MATCH([[DN No.]],cMov_DocNo,0)),""),""))'), fmt="date"),
        Col("Return Condition", "calc",
            T('=IF([[RN No.]]="","",IFERROR(INDEX(cMov_Condition,MATCH([[RN No.]],cMov_DocNo,0))&"",""))'), w=11),
        Col("Status", "calc",
            T('=IF([[Serial No.]]="","",IF(AND([[RN No.]]<>"",IFERROR(INDEX(cMov_Status,MATCH([[RN No.]],cMov_DocNo,'
              '0))&"","")="Received"),"Returned",IF([[Dispatch Date]]<>"","Dispatched",IF([[DN No.]]<>"","DN pending",'
              'IF([[PL / Invoice No.]]<>"","In Stock","Unknown")))))'), w=10),
        Col("Installed", "calc",
            T('=IF([[Status]]<>"Dispatched","",IFERROR(IF(INDEX(Master_CU,MATCH([[Site ID]],Master_B,0))="","",'
              'INDEX(Master_CU,MATCH([[Site ID]],Master_B,0))),""))'), fmt="date"),
        Col("Check", "calc",
            T('=IF([[Serial No.]]="","",IF(COUNTIF(cSer_SerialNo,[[Serial No.]])>1,"Duplicate serial",IF([[Item Code]]'
              '="","Unknown item",IF(IFERROR(INDEX(WHTracking,MATCH([[Item Name]],WHItemNames,0))&"","")<>"Serial",'
              '"Item not serial-tracked",IF(AND([[DN No.]]<>"",' + lk.format(d="DN No.") + '=0),"DN not found",'
              'IF(AND([[RN No.]]<>"",[[DN No.]]=""),"RN without DN",IF(AND([[RN No.]]<>"",' + lk.format(d="RN No.")
              + '=0),"RN not found","OK")))))))'), w=14),
        Col("Installed Key (helper)", "calc",
            T('=IF(AND([[Site ID]]<>"",ISNUMBER([[Installed]])),[[Site ID]]&"|"&[[Item Code]],"")'), w=10),
    ]
    register(ws, names, "tblWHSerials", "cSer", cols, sub_row=5, hdr_row=6, nrows=20)
    ws.freeze_panes = "C7"
    return ws
