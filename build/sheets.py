"""Guide, Dashboard, Project Status, Progress & Plan, Data Quality (D65-D71)."""
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from common import (AL_C, AL_L, AL_LW, BORDER, C_CALC, C_CALC_HDR, C_DQ_HDR, C_GROUP, C_HDR, C_INPUT,
                    C_RED_BG, C_RED_FG, C_SLA_BLOCK, C_TITLE, C_YEL_BG, C_YEL_FG, Col, DATE_FMT, PCT_FMT, T,
                    fill, font, register, style, title, valid, M1, NSITES, SLA_R1)


def red():
    return dict(fill=PatternFill("solid", fgColor="FF" + C_RED_BG), font=Font(color="FF" + C_RED_FG))


def v(c):
    """Validity array of a Master column."""
    return "(Master_" + c + ">=DateMin)*(Master_" + c + "<=DateMax)"


def vr(name):
    return "(" + name + ">=DateMin)*(" + name + "<=DateMax)"


def band(ws, r, c1, c2, text):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    style(ws.cell(r, c1, text), font(10, True, "FFFFFF"), fill(C_SLA_BLOCK), AL_L)
    for c in range(c1 + 1, c2 + 1):
        style(ws.cell(r, c), None, fill(C_SLA_BLOCK))


def hdr(ws, r, c, text, calc=True):
    style(ws.cell(r, c, text), font(9, True), fill(C_CALC_HDR if calc else C_HDR), AL_C)


def cell(ws, r, c, f, fmt="General", bold=False, f_fill=C_CALC):
    style(ws.cell(r, c, f), font(9, bold), fill(f_fill) if f_fill else None, AL_C, fmt)


CW = '(SC_CW="Yes")'
TE = '(SC_TE="Yes")'
RFI = '(SC_RFI="Yes")'
MW = "(SC_MW>0)"
ANALYSIS = '((Master_AP="STA")+(Master_AP="ST")>0)'
PRED38 = '(((Master_AI="Done")+(SC_CW="N/A")*' + v("V") + ')>0)'


# ================================================================== Guide (D71)
SHEETS = [
    ("Control", "Master", "One row per site: identification, scope, all workflow dates, gates and milestones.",
     "Site ID (from RF & Design), scope Yes / N/A, phase, batch, scenario, dates and statuses with white headers. "
     "Grey headers are calculated."),
    ("Control", "Dashboard", "Six tables: milestones vs baseline, progress by region, checkpoint, exceptions, "
     "warehouse complete units, swap & relocation.", "Selectors in row 3 (blank = All)."),
    ("Control", "Project Status", "Weighted physical progress (actual, planned, SPI) by activity and item.",
     "Phase / Batch selectors, weights (yellow) and the Manual Deliverables Log."),
    ("Control", "Progress & Plan", "Daily page: today's progress and tomorrow's plan (10 rows each, no history).",
     "Site ID, SOW, planned / actual work, reason."),
    ("Control", "Data Quality", "27 checks that protect the calculations; Rows > 0 needs action.", "Nothing."),
    ("Control", "SLA", "Working-day duration of every workflow step per site, completed and open.",
     "Min / Max in rows 6-7 (the only source of durations)."),
    ("Control", "Baseline", "Proposed plan dates (calculated) and approved versions (pasted values).",
     "Paste approved rows into the Approved table; planning links Min / Max."),
    ("Registers", "RF & Design", "Site register: identification, RF design, TSSR, LOS, LLD.",
     "One row per site. MW Far-end Site (N/A for Core / MFN / FN)."),
    ("Registers", "Acquisition", "Candidates, selection, contract and lease.", "One row per candidate."),
    ("Registers", "PR-PO", "Orders / assignments, contractors and teams.", "One row per order; contractors; teams."),
    ("Registers", "GECOL", "GECOL power steps per site.", "One row per site with GECOL = Yes."),
    ("Registers", "HLC", "HLC fiber request and installation.", "One row per site with HLC = Yes."),
    ("Registers", "As-Built", "Installed quantities and configuration per site.", "One row per site."),
    ("Registers", "WH", "Item definitions, stock summary and all movements.", "Item header rows; one row per document."),
    ("Registers", "WH Serials", "Serial numbers and their documents.", "One row per serial."),
    ("Registers", "Risk Register", "Project-level risks.", "One row per risk."),
    ("Registers", "Permits & HSE", "Permits, access and PTW (HSE checks on PTW rows).", "One row per permit."),
    ("Registers", "Quality Inspections", "QC, Snag, ITP and NCR records.", "One row per record."),
    ("Registers", "Change Control", "Suspension, resumption, cancellation, cutover window / rollback, scope changes.",
     "One row per change."),
    ("Registers", "Readiness & Handover", "Power safety, uplink, O&M acceptance, HOTO package, closure.",
     "One row per site."),
    ("Reference", "Setting", "Parameters, holidays, staff and every drop-down list.", "Report Date and list values."),
    ("Reference", "Site Calc", "Shared per-site results (protected).", "Nothing."),
]
MILESTONES = [
    ("Site Release", "RF & Acquisition", "Master V", "Unique Site ID, one selected candidate with signed contract, "
     "TSSR approved, LLD approved, civil aviation approved when requested; HUB / Terminal: LOS Pass and link budget.",
     "Contract, TSSR, LLD", "RF / Acquisition"),
    ("PRE-CW Completion", "PRE-CW", "Master AW", "Pre-CW Readiness, scope resolved, PR / PO / WO, CW access, survey, "
     "analysis, layout and CW DN approved (CW sites).", "PO / WO, layout approval", "CW"),
    ("CW Completion", "CW", "Master BP", "CW Readiness (access + PTW, no hold / stop), required CW items Done / N/A, "
     "CW actual end.", "CW completion record", "CW"),
    ("Site RFI", "CW", "Master BZ", "RFI Readiness, CW PAC (snags cleared), GECOL Power-On and HLC TX link tested "
     "where required.", "RFI declaration", "CW · GECOL · HLC"),
    ("TE Completion", "TE", "Master CU", "TE Readiness (access + PTW), TE items Done / N/A, TE actual end.",
     "Installation record", "TE"),
    ("On-Air", "TE", "Master DP", "On-Air Readiness (power safe), commissioning, integration, MW alignment and "
     "frequency licence with MW, uplink ready, swap MOP and window approved.", "Integration log", "TE"),
    ("TE PAC", "TE", "Master EL", "TE PAC Readiness, SSV Pass, HWID, EIR, PAT, NOC, As-Built, snags cleared, TE PAC "
     "actual.", "TE PAC certificate", "TE"),
    ("KPI Acceptance", "TE", "Master ES", "One valid KPI source date for the cluster, after the site's readiness.",
     "KPI acceptance", "Optimization"),
    ("HOTO", "TE", "Master EY", "HOTO Readiness, launch, material reconciliation, HOTO package and O&M acceptance.",
     "Signed O&M acceptance", "TE · O&M"),
    ("Final Closure", "Post-HOTO", "Master FM", "FAC (CW / TE), obligations closed, no blocking NCR, closure "
     "approved.", "Closure approval", "Project"),
]
OM_ITEMS = ["OSS / NMS inventory", "Configuration baseline backup", "Software / firmware baseline",
            "License inventory", "Asset serial reconciliation", "Alarm baseline", "Spares status",
            "Warranty support", "O&M manuals", "Training transfer", "Access accounts handover",
            "Security hardening", "Escalation matrix", "Performance baseline"]
PROCEDURES = [
    ("Report Date", "Type the reporting cut-off date in Setting (Report Date). Overdue, open durations and "
                    "plan-to-date use it."),
    ("Approve a baseline", "1) Check Baseline › Proposal (Baseline Status = Proposed). 2) Copy the Site ID and "
                           "the nine proposed dates. 3) Paste Values into Baseline › Approved. 4) Type Version and "
                           "Approval Date on the pasted rows. 5) The newest Approval Date per site becomes Current and "
                           "Master reads it. Record the re-plan reason in Change Control (Baseline Re-plan)."),
    ("Add list values", "Type the new value in the next empty row inside its table in Setting; the drop-downs grow "
                        "automatically. Do not delete values that are already used."),
    ("More than " + str(NSITES) + " sites", "Insert copied rows inside Master (above row " + str(M1) + "), then extend "
                            "Site Calc, SLA and Baseline Proposal rows by the same number and the Master_* names."),
]


def build_guide(wb, names):
    ws = wb.create_sheet("Guide")
    ws.sheet_view.showGridLines = False
    title(ws, "A1", "GSDN SITE DELIVERY MASTER · GUIDE", 13)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 60
    for c in "EFG":
        ws.column_dimensions[c].width = 22
    r = 3
    band(ws, r, 1, 4, "1 · SHEET INDEX")
    r += 1
    for i, h in enumerate(["Group", "Sheet", "Purpose", "What you enter"]):
        hdr(ws, r, 1 + i, h, False)
    for g, s, p, w in SHEETS:
        r += 1
        for i, x in enumerate([g, s, p, w]):
            style(ws.cell(r, 1 + i, x), font(9, i == 1), None, AL_LW if i >= 2 else AL_C)
    r += 2
    band(ws, r, 1, 6, "2 · MILESTONE DEFINITIONS")
    r += 1
    for i, h in enumerate(["Milestone", "Stage", "Defining Column", "Achievement Criterion", "Evidence", "Owner"]):
        hdr(ws, r, 1 + i, h, False)
    for row in MILESTONES:
        r += 1
        for i, x in enumerate(row):
            style(ws.cell(r, 1 + i, x), font(9, i == 0), None, AL_LW if i == 3 else AL_C)
    r += 2
    band(ws, r, 1, 4, "3 · PROCEDURES")
    for p, t in PROCEDURES:
        r += 1
        style(ws.cell(r, 1, p), font(9, True), None, AL_C)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        style(ws.cell(r, 2, t), font(9), None, AL_LW)
        ws.row_dimensions[r].height = 42
    r += 2
    band(ws, r, 1, 4, "4 · O&M HANDOVER CHECKLIST (Readiness & Handover › O&M Checklist = Done when all apply)")
    for i, t in enumerate(OM_ITEMS):
        r += 1
        style(ws.cell(r, 1, i + 1), font(9), None, AL_C)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        style(ws.cell(r, 2, t), font(9), None, AL_L)
        style(ws.cell(r, 4, "☐ Done / N/A"), font(9), None, AL_C)
    return ws


# ================================================================== Dashboard (D66, D67, D72)
MS = [  # milestone, stage, actual, baseline, predecessor mask, applicability mask
    ("Site Release", "RF & Acquisition", "V", "U", "1", "1"),
    ("PRE-CW Completion", "PRE-CW", "AW", "AV", v("W"), "1"),
    ("CW Completion", "CW", "BP", "BO", v("AX"), CW),
    ("Site RFI", "CW", "BZ", "BY", v("BQ"), RFI),
    ("TE Completion", "TE", "CU", "CT", v("CA"), TE),
    ("On-Air", "TE", "DP", "DO", v("CV"), TE),
    ("TE PAC", "TE", "EL", "EK", v("DQ"), TE),
    ("KPI Acceptance", "TE", "ES", "ER", v("EM"), TE),
    ("HOTO", "TE", "EY", "EX", v("ET"), TE),
    ("Final Closure", "Post-HOTO", "FM", "FL", v("EZ"), "1"),
]


def step(name, done, pred, appl):
    return (name, done, pred, appl)


CW_STEPS = [
    step("RF Survey", v("S"), "1", "1"),
    step("Acquisition", v("T"), v("S"), "1"),
    step("Site Release", v("V"), v("T"), "1"),
    step("PR", '(Master_AH="Done")', v("V"), CW + '*(Master_AG="Outsourced")'),
    step("PO / WO", '(Master_AI="Done")', v("V") + '*(((Master_AH="Done")+(Master_AH="N/A"))>0)', CW),
    step("CW Access", v("AL"), '(Master_AI="Done")', CW),
    step("CW Survey", v("AO"), v("AL"), CW),
    step("STA / ST", v("AR"), v("AO"), CW + "*" + ANALYSIS),
    step("Report Submission", v("AR"), v("AO"), CW),
    step("Report Approval", v("AS"), v("AR"), CW),
    step("CW MOS", v("AZ"), v("AX"), CW),
    step("CW Start", v("BL"), v("AZ"), CW),
    step("CW Completion", v("BP"), v("BL"), CW),
    step("CW PAC", v("BV"), v("BP"), CW),
    step("GECOL Power", '(Master_BW="Done")', PRED38, '(Master_AC="Yes")'),
    step("HLC Fiber", '(Master_BX="Done")', PRED38, '(Master_AD="Yes")'),
    step("Site RFI", v("BZ"), v("BQ"), RFI),
]
TE_STEPS = [
    step("TE Access", v("CB"), v("V"), TE),
    step("TE Survey", v("CC"), v("CB"), TE),
    step("DN Submission", v("CG"), v("CC"), TE),
    step("DN Approval", v("CH"), v("CG"), TE),
    step("TE Dispatch", vr("SC_Dispatch"), v("CH"), TE),
    step("TE MOS", v("CJ"), vr("SC_Dispatch"), TE),
    step("TE Start", v("CQ"), v("CJ"), TE),
    step("TE Completion", v("CU"), v("CQ"), TE),
    step("Commissioning", v("CW"), v("CU"), TE),
    step("MW Alignment", v("CX"), v("CW"), TE + "*" + MW),
    step("Integration", v("DA"), "(" + MW + "*" + v("CX") + "+(1-" + MW + ")*" + v("CW") + ")", TE),
    step("On-Air", v("DP"), v("DA"), TE),
    step("SSV", v("DR"), v("DP"), TE),
    step("HWID", v("DT"), v("DP"), TE),
    step("EIR", v("DY"), v("DT"), TE),
    step("PAT", v("EE"), v("DY"), TE),
    step("As-Built", v("EF"), v("EE"), TE),
    step("TE PAC", v("EL"), v("EF"), TE),
    step("Optimization", v("EO"), v("EL"), TE),
    step("KPI Acceptance", v("ES"), v("EM"), TE),
    step("Material Reconciliation", v("EV"), v("ES"), TE),
    step("HOTO Package", v("EW"), v("EV"), TE),
    step("HOTO", v("EY"), v("EW"), TE),
]
POST_STEPS = [
    step("CW FAC", v("FE"), v("FB") + "*(Master_FB<=RptDate)", CW),
    step("TE FAC", v("FJ"), v("FG") + "*(Master_FG<=RptDate)", TE),
    step("Final Closure", v("FM"), v("EZ"), "1"),
]
PROG_CW = [("Site Release", v("V"), "1", "1"), ("PRE-CW Completion", v("AW"), v("W"), "1"),
           ("CW Completion", v("BP"), v("AX"), CW), ("CW PAC", v("BV"), v("BP"), CW),
           ("GECOL Power", '(Master_BW="Done")', PRED38, '(Master_AC="Yes")'),
           ("HLC Fiber", '(Master_BX="Done")', PRED38, '(Master_AD="Yes")'), ("Site RFI", v("BZ"), v("BQ"), RFI)]
PROG_TE = [("TE Completion", v("CU"), v("CA"), TE), ("On-Air", v("DP"), v("CV"), TE),
           ("TE PAC", v("EL"), v("DQ"), TE), ("KPI Acceptance", v("ES"), v("EM"), TE),
           ("HOTO", v("EY"), v("ET"), TE), ("Final Closure", v("FM"), v("EZ"), "1")]


def sp(*parts):
    return "SUMPRODUCT(" + "*".join(p for p in parts if p != "1") + ")" if any(p != "1" for p in parts) else "0"


def build_dashboard(wb, names):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 30
    for c in range(3, 80):
        ws.column_dimensions[L(c)].width = 10.5
    title(ws, "B1", "SITE DELIVERY DASHBOARD · CW / TE / HOTO", 13)
    # selectors (D66)
    sel = [("Project Phase", "C", "ProjectPhases", "DashPhase"), ("Delivery Batch", "E", "DeliveryBatches", "DashBatch"),
           ("Delivery Scenario", "G", "DeliveryScenario", "DashScenario")]
    style(ws["B3"], font(9, True), fill(C_HDR), AL_C)
    ws["B3"] = "Project Phase"
    for lab, col, lst, nm in sel:
        c0 = ws[col + "3"]
        lc = ws.cell(3, c0.column - 1)
        lc.value = lab
        style(lc, font(9, True), fill(C_HDR), AL_C)
        style(c0, font(9), fill(C_INPUT), AL_C)
        dv = DataValidation(type="list", formula1="=" + lst, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(col + "3")
        names.add(nm, f"Dashboard!${col}$3")
    style(ws["H3"], font(9, True), fill(C_HDR), AL_C)
    ws["H3"] = "Report Date"
    style(ws["I3"], font(9), fill(C_CALC), AL_C, DATE_FMT)
    ws["I3"] = '=IF(N(RptDate)=0,"Set in Setting",RptDate)'
    ws["B4"] = ("Blank selectors = All. Scope = sites the item applies to (not cancelled). Pending = its turn has "
                "arrived (previous step done). Not Due = previous step not done. Exceptions and WH are project-wide.")
    ws["B4"].font = font(8, italic=True)
    names.add("DashSel", f"('Site Calc'!$C$5:$C${M1}=1)*(Master!$FZ$5:$FZ${M1}<>\"Cancelled\")")
    # -------------------------------------------------- 1 Milestone Delivery
    r = 6
    band(ws, r, 2, 13, "1 · MILESTONE DELIVERY · approved baseline vs achievement")
    r += 1
    heads = ["Milestone", "Stage", "Scope", "Done", "Pending", "Not Due", "% Done", "Overdue Open", "Achieved Late",
             "Avg Delay (days)", "Baseline Missing", "Check"]
    for i, h in enumerate(heads):
        hdr(ws, r, 2 + i, h)
    ws.row_dimensions[r].height = 30
    for name, stage, a, b, pred, appl in MS:
        r += 1
        style(ws.cell(r, 2, name), font(9, True), None, AL_L)
        style(ws.cell(r, 3, stage), font(9), None, AL_C)
        late = sp("DashSel", appl, v(a), v(b), "(Master_" + a + ">Master_" + b + ")")
        lm = "DashSel" + ("" if appl == "1" else "*" + appl) + "*" + v(a) + "*" + v(b) + "*(Master_" + a + ">Master_" + b + ")"
        f = {
            4: "=" + sp("DashSel", appl),
            5: "=" + sp("DashSel", appl, v(a)),
            6: "=" + sp("DashSel", appl, "(1-" + v(a) + ")", pred),
            7: f"=D{r}-E{r}-F{r}" .replace("D", "D").replace("E", "E"),
            8: f'=IF(D{r}=0,"",E{r}/D{r})',
            9: '=IF(N(RptDate)=0,"",' + sp("DashSel", appl, "(1-" + v(a) + ")", v(b), "(Master_" + b + "<RptDate)") + ')',
            10: "=" + late,
            11: f'=IF(J{r}=0,"",(SUMPRODUCT({lm},Master_{a})-SUMPRODUCT({lm},Master_{b}))/J{r})',
            12: "=" + sp("DashSel", appl, "(1-" + v(b) + ")"),
            13: f'=IF(G{r}<0,"Check","OK")',
        }
        # column letters: Scope=D(4), Done=E(5), Pending=F(6), NotDue=G(7), %=H(8)...
        f[7] = f"=D{r}-E{r}-F{r}"
        f[8] = f'=IF(D{r}=0,"",E{r}/D{r})'
        f[11] = f'=IF(J{r}=0,"",(SUMPRODUCT({lm},Master_{a})-SUMPRODUCT({lm},Master_{b}))/J{r})'
        f[13] = f'=IF(G{r}<0,"Check","OK")'
        for c, x in f.items():
            cell(ws, r, c, x, PCT_FMT if c == 8 else ("0.0" if c == 11 else "General"))
    m_last = r
    # -------------------------------------------------- 4 Exceptions & Alerts (left) and 5 WH (right)
    r = m_last + 3
    band(ws, r, 2, 6, "4 · EXCEPTIONS & ALERTS · project-wide")
    band(ws, r, 9, 14, "5 · WH · COMPLETE UNITS · project-wide")
    ex_top = r + 1
    for i, h in enumerate(["Area · Indicator", "Base", "Count", "% of Base", ""]):
        if h:
            hdr(ws, ex_top, 2 + i, h)
    exc = exceptions()
    rr = ex_top
    area = None
    for ar, ind, base, cnt in exc:
        if ar != area:
            rr += 1
            style(ws.cell(rr, 2, ar), font(9, True, "FFFFFF"), fill(C_TITLE), AL_L)
            for c in (3, 4, 5):
                style(ws.cell(rr, c), None, fill(C_TITLE))
            area = ar
        rr += 1
        style(ws.cell(rr, 2, "   " + ind), font(9), None, AL_L)
        cell(ws, rr, 3, "=" + base)
        cell(ws, rr, 4, "=" + cnt, bold=True)
        cell(ws, rr, 5, f'=IF(N(C{rr})=0,"",D{rr}/C{rr})', PCT_FMT)
        ws.conditional_formatting.add(f"D{rr}", CellIsRule(operator="greaterThan", formula=["0"], **red()))
    ex_last = rr
    # WH complete units
    wr = ex_top
    for i, h in enumerate(["Unit", "Need (Scope)", "Dispatched", "Remaining Need", "Available", "Shortfall"]):
        hdr(ws, wr, 9 + i, h)
    from wh import UNIT_NEED
    units = ["Cambium Sector", "Intracom Sector", "MW Half-link", "Router", "Cabinet", "Battery", "LPU Kit"]
    for u in units:
        wr += 1
        style(ws.cell(wr, 9, u), font(9, True), None, AL_L)
        q = '"' + u + '"'
        ncomp = "MAX(1,COUNTIF(WHNeedUnit," + q + "))" if u != "Router" else "1"
        cell(ws, wr, 10, "=IFERROR(" + UNIT_NEED.format(u=q) + ",0)")
        cell(ws, wr, 11, "=(SUMIF(WHNeedUnit," + q + ",WHIssued)-SUMIF(WHNeedUnit," + q + ",WHReturnedGood)-SUMIF("
                         "WHNeedUnit," + q + ",WHFaulty))/" + ncomp)
        cell(ws, wr, 12, f"=MAX(0,J{wr}-K{wr})")
        avail = ("SUMIF(WHNeedUnit," + q + ",WHAvailable)" if u == "Router" else
                 "IF(COUNTIF(WHNeedUnit," + q + ")=0,0,_xlfn.MINIFS(WHAvailable,WHNeedUnit," + q + "))")
        cell(ws, wr, 13, "=" + avail)
        cell(ws, wr, 14, f"=MAX(0,L{wr}-M{wr})", bold=True)
        ws.conditional_formatting.add(f"N{wr}", CellIsRule(operator="greaterThan", formula=["0"], **red()))
    ws.cell(wr + 1, 9, "Router = sum of router models; other units = the scarcest component.").font = font(8, italic=True)
    # -------------------------------------------------- 2 Overall Progress by Region
    r = max(ex_last, wr) + 3
    band(ws, r, 2, 2 + 1 + 4 * len(PROG_CW), "2 · OVERALL PROGRESS BY REGION · Scope / Done / Pending / Not Due")
    for blk, items in (("CW", PROG_CW), ("TE", PROG_TE)):
        r += 1
        r = progress_block(ws, r, blk, items)
        r += 1
    # -------------------------------------------------- 3 Project Checkpoint
    r += 1
    band(ws, r, 2, 2 + len(TE_STEPS), "3 · PROJECT CHECKPOINT · detailed activities (Done / Pending)")
    for blk, items in (("CW", CW_STEPS), ("TE", TE_STEPS), ("POST-HOTO", POST_STEPS)):
        r += 1
        r = checkpoint_block(ws, r, blk, items)
        r += 1
    # -------------------------------------------------- 6 Swap & Relocation
    r += 1
    band(ws, r, 2, 13, "6 · SWAP & RELOCATION")
    r += 1
    heads = ["Scenario", "Scope", "On-Air Done", "Pending", "Not Due", "% Done", "First-time Success", "Rollbacks",
             "Window over Max", "CW Dismantle overdue", "TE Dismantle overdue", "Post-Swap KPI degraded"]
    for i, h in enumerate(heads):
        hdr(ws, r, 2 + i, h)
    ws.row_dimensions[r].height = 30
    top = r + 1
    for sc in ["New Build", "Expansion", "Hot Swap", "Cold Swap", "Relocation", "(not set)"]:
        r += 1
        style(ws.cell(r, 2, sc), font(9, True), None, AL_L)
        zm = '(Master_Z="")' if sc == "(not set)" else '(Master_Z="' + sc + '")'
        base = ["DashSel", zm, TE]
        cell(ws, r, 3, "=" + sp(*base))
        cell(ws, r, 4, "=" + sp(*base, v("DP")))
        cell(ws, r, 5, "=" + sp(*base, "(1-" + v("DP") + ")", v("CV")))
        cell(ws, r, 6, f"=C{r}-D{r}-E{r}")
        cell(ws, r, 7, f'=IF(C{r}=0,"",D{r}/C{r})', PCT_FMT)
        cell(ws, r, 8, "=" + sp(*base, v("DP"), '(Master_DF="Success")', "(N(+Master_DI)=0)")
             .replace("(N(+Master_DI)=0)", "(Master_DI=0)"))
        cell(ws, r, 9, "=SUMPRODUCT(DashSel*" + zm + ",Master_DI)")
        win = ("SUMPRODUCT(DashSel*" + zm + "*(ISNUMBER(SLA!{c}13:{c}" + str(SLA_R1) + "))*(SLA!{c}13:{c}" + str(SLA_R1) + ">N(SLA!{c}7)))")
        cell(ws, r, 10, "=" + win.format(c=sla_col("S-3")) + "+" + win.format(c=sla_col("S-4")))
        cell(ws, r, 11, '=IF(N(RptDate)=0,"",' + sp("DashSel", zm, '(Master_FN="Yes")', "(1-" + v("FQ") + ")",
                                                     v("FP"), "(Master_FP<RptDate)") + ')')
        cell(ws, r, 12, '=IF(N(RptDate)=0,"",' + sp("DashSel", zm, "(1-" + v("FV") + ")",
                                                     "(" + v("DH") + "*(Master_DH<RptDate)+(Master_FS=\"Yes\")*"
                                                     + v("FU") + "*(Master_FU<RptDate)>0)") + ')')
        cell(ws, r, 13, "=" + sp("DashSel", zm, '(Master_DL="Degraded")'))
    r += 1
    style(ws.cell(r, 2, "Total"), font(9, True), fill(C_CALC_HDR), AL_L)
    for c in range(3, 14):
        col = L(c)
        f = f"=SUM({col}{top}:{col}{r-1})" if c != 7 else f'=IF(C{r}=0,"",D{r}/C{r})'
        cell(ws, r, c, f, PCT_FMT if c == 7 else "General", bold=True, f_fill=C_CALC_HDR)
    ws.freeze_panes = "C5"
    return ws


def sla_col(code, block="C"):
    import calc_sheets as cs
    k = [st[0] for st in cs.STEPS].index(code)
    return L(4 + k + (len(cs.STEPS) if block == "O" else 0))


def progress_block(ws, r, blk, items):
    style(ws.cell(r, 2, blk + " · Region"), font(9, True), fill(C_CALC_HDR), AL_C)
    for k, (nm, *_rest) in enumerate(items):
        c = 3 + 4 * k
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 3)
        style(ws.cell(r, c, nm), font(9, True), fill(C_CALC_HDR), AL_C)
        for j in range(1, 4):
            style(ws.cell(r, c + j), None, fill(C_CALC_HDR))
    r += 1
    style(ws.cell(r, 2, "Region"), font(9, True), fill(C_CALC_HDR), AL_C)
    for k in range(len(items)):
        for j, h in enumerate(["Scope", "Done", "Pending", "Not Due"]):
            hdr(ws, r, 3 + 4 * k + j, h)
    top = r + 1
    rows = [("INDEX(Region,%d)" % i, '(Master_E=$B{r})') for i in range(1, 8)]
    rows.append(('"Other / missing region"', "(COUNTIF(Region,Master_E)=0)"))
    for lab, mask in rows:
        r += 1
        style(ws.cell(r, 2, '=IFERROR(' + lab + '&"","")' if lab.startswith("INDEX") else lab.strip('"')),
              font(9, True), None, AL_L)
        m = mask.format(r=r)
        for k, (nm, done, pred, appl) in enumerate(items):
            c = 3 + 4 * k
            cell(ws, r, c, "=" + sp("DashSel", m, appl))
            cell(ws, r, c + 1, "=" + sp("DashSel", m, appl, done))
            cell(ws, r, c + 2, "=" + sp("DashSel", m, appl, "(1-" + done + ")", pred))
            cell(ws, r, c + 3, f"={L(c)}{r}-{L(c+1)}{r}-{L(c+2)}{r}")
    r += 1
    style(ws.cell(r, 2, "Total"), font(9, True), fill(C_CALC_HDR), AL_L)
    for k in range(len(items)):
        for j in range(4):
            col = L(3 + 4 * k + j)
            cell(ws, r, 3 + 4 * k + j, f"=SUM({col}{top}:{col}{r-1})", bold=True, f_fill=C_CALC_HDR)
    return r


def checkpoint_block(ws, r, blk, items):
    style(ws.cell(r, 2, blk), font(9, True), fill(C_CALC_HDR), AL_C)
    for k, it in enumerate(items):
        hdr(ws, r, 3 + k, it[0])
    ws.row_dimensions[r].height = 30
    labs = ["Scope", "Done", "Pending (Due)", "Not Due", "% Done"]
    for j, lab in enumerate(labs):
        rr = r + 1 + j
        style(ws.cell(rr, 2, lab), font(9, True), fill(C_CALC_HDR if j in (1, 2) else None) if j in (1, 2) else None,
              AL_L)
        for k, (nm, done, pred, appl) in enumerate(items):
            c = 3 + k
            col = L(c)
            if j == 0:
                f = "=" + sp("DashSel", appl)
            elif j == 1:
                f = "=" + sp("DashSel", appl, done)
            elif j == 2:
                f = "=" + sp("DashSel", appl, "(1-" + done + ")", pred)
            elif j == 3:
                f = f"={col}{r+1}-{col}{r+2}-{col}{r+3}"
            else:
                f = f'=IF({col}{r+1}=0,"",{col}{r+2}/{col}{r+1})'
            cell(ws, rr, c, f, PCT_FMT if j == 4 else "General", bold=j in (1, 2))
    return r + len(labs)


def exceptions():
    """(area, indicator, base formula, count formula)."""
    rd = "N(RptDate)>0"
    valid_permit = '(cPrm_Status="Approved")*(cPrm_ValidTo>0)'
    E = []
    a = "Permits & Contracts"
    E += [(a, "Permits / access expiring", "SUMPRODUCT(" + valid_permit + ")", 'COUNTIF(cPrm_Validity,"Expiring")'),
          (a, "Permits / access expired", "SUMPRODUCT(" + valid_permit + ")", 'COUNTIF(cPrm_Validity,"Expired")'),
          (a, "Site leases expiring", 'COUNTIFS(cAcq_FeedsMaster,"Yes",cAcq_LeaseExpiryDate,">0")',
           'IF(' + rd + ',COUNTIFS(cAcq_FeedsMaster,"Yes",cAcq_LeaseExpiryDate,">="&RptDate,cAcq_LeaseExpiryDate,'
           '"<="&RptDate+ExpiringDays),0)'),
          (a, "Site leases expired", 'COUNTIFS(cAcq_FeedsMaster,"Yes",cAcq_LeaseExpiryDate,">0")',
           'IF(' + rd + ',COUNTIFS(cAcq_FeedsMaster,"Yes",cAcq_LeaseExpiryDate,">0",cAcq_LeaseExpiryDate,"<"&RptDate),0)'),
          (a, "Contractor contracts expiring / expired", 'COUNTIF(cCo_Contractor,"?*")',
           'COUNTIF(cCo_ContractValidity,"Expiring")+COUNTIF(cCo_ContractValidity,"Expired")')]
    a = "PR-PO"
    E += [(a, "CW sites released without PO / WO", 'SUMPRODUCT((SC_CW="Yes")*' + v("V") + ')',
           'SUMPRODUCT((SC_CW="Yes")*' + v("V") + '*(Master_AI<>"Done"))'),
          (a, "PR → PO over SLA (B-2 Max)", 'COUNT(cOrd_PRPOwd)', 'COUNTIF(cOrd_PRPOwd,">"&N(SLA_Max_B2))')]
    a = "GECOL & HLC"
    E += [(a, "GECOL open", 'COUNTIF(Master_AC,"Yes")', 'COUNTIFS(Master_AC,"Yes",Master_BW,"<>Done")'),
          (a, "GECOL meter installed, not powered", 'COUNTIF(cPwr_SiteID,"?*")',
           'SUMPRODUCT((cPwr_InstallationFinishDate>0)*(cPwr_PowerOnDateMeterCharging=""))'),
          (a, "GECOL open over SLA (C-2 Max)", 'COUNTIF(Master_AC,"Yes")',
           "COUNTIF(SLA!" + "{C2O}" + ',">"&N(SLA_Max_C2))'),
          (a, "Fiber open", 'COUNTIF(Master_AD,"Yes")', 'COUNTIFS(Master_AD,"Yes",Master_BX,"<>Done")'),
          (a, "Fiber open over SLA (C-4 Max)", 'COUNTIF(Master_AD,"Yes")',
           "COUNTIF(SLA!" + "{C4O}" + ',">"&N(SLA_Max_C4))'),
          (a, "Fiber Status conflicts with HLC = Yes", 'COUNTIF(cFib_SiteID,"?*")',
           'SUMPRODUCT(((cFib_FiberStatus="No Fiber")+(cFib_FiberStatus="Through MW Link"))*(COUNTIFS(Master_B,'
           'cFib_SiteID,Master_AD,"Yes")>0))')]
    a = "WH"
    E += [(a, "DN approved, not dispatched", 'COUNTIF(cMov_Type,"*DN")',
           'COUNTIFS(cMov_Type,"*DN",cMov_Status,"Approved")'),
          (a, "Faulty / damaged returned (qty)", "SUM(WHReturnedGood)+SUM(WHFaulty)", "SUM(WHFaulty)"),
          (a, "Movement rows failing Check", 'COUNTIF(cMov_Check,"?*")',
           'COUNTIF(cMov_Check,"?*")-COUNTIF(cMov_Check,"OK")'),
          (a, "Serial rows failing Check", 'COUNTIF(cSer_Check,"?*")',
           'COUNTIF(cSer_Check,"?*")-COUNTIF(cSer_Check,"OK")'),
          (a, "Items with negative available", "COUNTA(WHItemNames)", 'COUNTIF(WHAvailable,"<0")')]
    a = "Quality"
    E += [(a, "Open records (QC / Snag / ITP / NCR)", 'COUNTIF(cQI_SiteID,"?*")',
           'COUNTIF(cQI_Status,"Open")+COUNTIF(cQI_Status,"Overdue")'),
          (a, "Blocking open (Hold / NCR / Critical-Major snag)", 'COUNTIF(cQI_SiteID,"?*")',
           'COUNTIF(cQI_Blocking,"Yes")'),
          (a, "Overdue against Target Date", 'COUNTIF(cQI_SiteID,"?*")', 'COUNTIF(cQI_Status,"Overdue")'),
          (a, "NCR open", 'COUNTIF(cQI_RecordType,"NCR")',
           'COUNTIFS(cQI_RecordType,"NCR",cQI_Status,"<>Closed",cQI_Status,"<>N/A")'),
          (a, "CW PAC dated, acceptance conditions open", "SUMPRODUCT(" + v("BV") + ")",
           "SUMPRODUCT(" + v("BV") + '*((Master_BZ="")+(Master_BS="")+((Master_BS="Yes")*(Master_BT<>"Yes"))>0))'),
          (a, "TE PAC dated, TE PAC milestone not achieved", "SUMPRODUCT(" + v("EJ") + ")",
           "SUMPRODUCT(" + v("EJ") + "*(1-" + v("EL") + "))"),
          (a, "MW links outside RSL tolerance", "COUNT(Master_CZ)", 'COUNTIF(Master_CZ,">"&RSLTol)+COUNTIF(Master_CZ,"<"&-RSLTol)')]
    a = "Issues & Risks"
    issue_open = '(Master_GA<>"")*(Master_GD<>"Closed")'
    E += [(a, "Open site issues", 'COUNTIF(Master_GA,"?*")', "SUMPRODUCT(" + issue_open + ")"),
          (a, "Open site issues older than the Expiring alert", "SUMPRODUCT(" + issue_open + ")",
           'IF(' + rd + ',SUMPRODUCT(' + issue_open + '*(Master_GC>0)*(Master_GC<RptDate-ExpiringDays)),0)'),
          ]
    ro = '((cRisk_Status="Open")+(cRisk_Status="Monitoring"))'
    for lvl in ("Critical", "High", "Medium", "Low"):
        E.append((a, "Open risks · " + lvl, 'COUNTIF(cRisk_RiskID,"?*")',
                  'SUMPRODUCT(' + ro + '*(cRisk_RiskLevel="' + lvl + '"))'))
    E += [(a, "Risks with review overdue", 'COUNTIF(cRisk_RiskID,"?*")',
           'IF(' + rd + ',SUMPRODUCT((cRisk_Status<>"Closed")*(cRisk_NextReviewDate>0)*(cRisk_NextReviewDate<RptDate)),0)'),
          (a, "Materialized risks", 'COUNTIF(cRisk_RiskID,"?*")', 'COUNTIF(cRisk_Status,"Materialized")')]
    a = "Schedule Control"
    E += [(a, "Open SLA steps over Max", "COUNT(SLA_Open)", "{OPENOVER}"),
          (a, "Completed SLA steps over Max", "COUNT(SLA_Completed)", "{COMPOVER}"),
          (a, "Pending change requests", 'COUNTIF(cChg_ChangeType,"?*")', 'COUNTIF(cChg_Approval,"Pending")'),
          (a, "Data Quality checks needing action", "27", "DQIssues"),
          (a, "Sites on hold", 'COUNTIF(Master_B,"?*")', 'COUNTIF(Master_FZ,"On Hold")'),
          (a, "Sites cancelled", 'COUNTIF(Master_B,"?*")', 'COUNTIF(Master_FZ,"Cancelled")')]
    return E


# ================================================================== Progress & Plan (D69)
def build_progress(wb, names):
    ws = wb.create_sheet("Progress & Plan")

    def look(c):
        return 'IF([[Site ID]]="","",IFERROR(INDEX(Master_' + c + ',MATCH([[Site ID]],Master_B,0))&"",""))'

    def lookv(c):
        x = 'INDEX(Master_' + c + ',MATCH([[Site ID]],Master_B,0))'
        return 'IF([[Site ID]]="","",IFERROR(IF(' + x + '="","",' + x + '),""))'
    base = [
        Col("No.", "num", T('=IF([[Site ID]]="","",ROW()-{off})'), w=5),
        Col("Site ID", dv="SiteIDs", w=11),
        Col("Tower Owner", "calc", T("=" + look("M")), w=10),
        Col("Owner ID", "calc", T("=" + look("C")), w=11),
        Col("Latitude", "calc", T("=" + lookv("H")), fmt="0.000000", w=10),
        Col("Longitude", "calc", T("=" + lookv("I")), fmt="0.000000", w=10),
        Col("Region", "calc", T("=" + look("E")), w=12),
        Col("Site Type", "calc", T("=" + look("J")), w=8),
        Col("SOW", dv="RefWHScope", w=6),
        Col("Vendor / Contractor Name", "calc",
            T('=IF([[Site ID]]="","",IF([[SOW]]="CW",' + look("AJ") + ',IF([[SOW]]="TE",' + look("CD") + ',"")))'), w=16),
        Col("AAT SE", "calc",
            T('=IF([[Site ID]]="","",IF([[SOW]]="CW",' + look("AM") + ',IF([[SOW]]="TE",' + look("CF") + ',"")))'), w=12),
    ]
    import copy as _c
    daily = [_c.copy(x) for x in base] + [Col("Planned Work", w=24), Col("Actual Work", w=24),
                                          Col("Reason (if not reach target)", w=22)]
    acc = ('IF(OR([[Site ID]]="",[[SOW]]=""),"",IF(COUNTIFS(cPrm_SiteID,[[Site ID]],cPrm_Type,"Access - "&[[SOW]],'
           'cPrm_Status,"Approved",cPrm_ValidFrom,"<="&TomorrowDate,cPrm_ValidTo,">="&TomorrowDate)+IF(IFERROR('
           'INDEX(PATypeHasExpiry,MATCH("Access - "&[[SOW]],PAType,0))&"","")="No",COUNTIFS(cPrm_SiteID,[[Site ID]],'
           'cPrm_Type,"Access - "&[[SOW]],cPrm_Status,"Approved",cPrm_ValidFrom,"<="&TomorrowDate,cPrm_ValidTo,""),0)'
           '+COUNTIFS(cPrm_SiteID,[[Site ID]],cPrm_Type,"Access - "&[[SOW]],cPrm_Validity,"N/A")>0,"Yes","No"))')
    tomorrow = [_c.copy(x) for x in base] + [Col("Access Valid", "calc", T("=" + acc), w=8),
                                             Col("Planned Work", w=24), Col("Remark", w=22)]
    # block title (A:B) with its date in C, header on the next row, 10 rows, no SUBTOTAL row (user layout)
    for blk, cols, r0, nm, f in (("Daily Progress", daily, 1, "TodayDate", "=TODAY()"),
                                 ("Tomorrow Plan", tomorrow, 14, "TomorrowDate",
                                  "=WORKDAY.INTL(TODAY(),1,7,Holidays)")):
        cols[0] = Col("No.", "num", T('=IF([[Site ID]]="","",ROW()-' + str(r0 + 1) + ')'), w=5)
        title(ws, f"A{r0}", blk, 12)
        ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=2)
        style(ws.cell(r0, 3, f), font(10, True), fill(C_CALC), AL_C, DATE_FMT)
        names.add(nm, f"'Progress & Plan'!$C${r0}")
        tname = "tblDailyProgress" if r0 == 1 else "tblTomorrowPlan"
        register(ws, names, tname, "cDP" if r0 == 1 else "cTP", cols, sub_row=None, hdr_row=r0 + 1, nrows=10,
                 header_h=24)
    ws.freeze_panes = "C3"
    return ws


# ================================================================== Data Quality (D68)
def dq_checks():
    def vv(c):
        return v(c)
    reg_site_cols = ["cAcq_SiteID", "cPrm_SiteID", "cOrd_SiteID", "cPwr_SiteID", "cFib_SiteID", "cAB_SiteID",
                     "cQI_SiteID", "cChg_SiteID", "cRdy_SiteID", "cMov_SiteID", "cBLA_SiteID"]
    notin = "+".join("SUMPRODUCT((" + c + "<>\"\")*(COUNTIF(Master_B," + c + ")=0))" for c in reg_site_cols)
    notin += '+SUMPRODUCT((cRF_SiteID<>"")*(COUNTIF(Master_B,cRF_SiteID)=0))'
    hub = '((cRF_SiteType="HUB")+(cRF_SiteType="Terminal"))'
    fiber = '((cRF_SiteType="Core")+(cRF_SiteType="MFN")+(cRF_SiteType="FN"))'
    fe = 'cRF_MWFarendSite'
    has_fe = '(' + fe + '<>"")*(' + fe + '<>"N/A")'
    lists = [("X", "ProjectPhases"), ("Y", "DeliveryBatches"), ("Z", "DeliveryScenario"), ("AG", "DeliveryModel"),
             ("AP", "AnalysisType"), ("AQ", "AnalysisResult"), ("GB", "IssueOwners"), ("GD", "IssueStatus"),
             ("DF", "CutoverResult"), ("DL", "KPIResult"), ("DS", "SSVResult")]
    outlist = "+".join('SUMPRODUCT((Master_' + c + '<>"")*(COUNTIF(' + n + ',Master_' + c + ')=0))' for c, n in lists)
    outlist += ('+SUMPRODUCT((cRF_Region<>"")*(COUNTIF(Region,cRF_Region)=0))'
                '+SUMPRODUCT((cRF_SiteType<>"")*(COUNTIF(SiteType,cRF_SiteType)=0))')
    swap = '((Master_Z="Hot Swap")+(Master_Z="Cold Swap")>0)'
    gap_pairs = [("AW", "V"), ("BP", "AW"), ("CU", None), ("DP", "CU"), ("EL", "DP"), ("ES", "EL"),
                 ("EY", "ES"), ("FM", "EY")]
    gaps = []
    for a, b in gap_pairs:
        if b:
            gaps.append(vv(a) + "*(1-" + vv(b) + ")")
    gaps.append(vv("BZ") + '*(1-' + vv("BP") + ')*(Master_BP<>"N/A")*(SC_CW="Yes")')
    gaps.append(vv("CU") + '*(1-' + vv("BZ") + ')*(Master_BZ<>"N/A")')
    C = [
        ("DQ01", "Master", "Duplicate Site ID (Master or RF & Design)", "Master · RF & Design",
         'SUMPRODUCT((Master_B<>"")*(COUNTIF(Master_B,Master_B)>1))+SUMPRODUCT((cRF_SiteID<>"")*'
         '(COUNTIF(cRF_SiteID,cRF_SiteID)>1))', "High", "Keep one row per Site ID."),
        ("DQ02", "Links", "Site ID in a register but not in Master", "All registers", notin, "High",
         "Add the site to Master or correct the Site ID."),
        ("DQ03", "Links", "Master site without an RF & Design row", "Master",
         'SUMPRODUCT((Master_B<>"")*(COUNTIF(cRF_SiteID,Master_B)=0))', "High", "Register the site in RF & Design."),
        ("DQ04", "Topology", "Uplink / MW Far-end invalid (missing for HUB / Terminal, set for fiber sites, unknown "
                             "site or self link)", "RF & Design",
         'SUMPRODUCT((cRF_SiteID<>"")*((' + hub + '*((' + fe + '="")+(' + fe + '="N/A")>0)+' + fiber + '*' + has_fe
         + '+' + has_fe + '*(COUNTIF(cRF_SiteID,' + fe + ')=0)+(' + fe + '=cRF_SiteID))>0))', "High",
         "HUB / Terminal need a Far-end site; Core / MFN / FN use N/A."),
        ("DQ05", "Acquisition", "More than one selected candidate for a site", "Acquisition",
         'SUMPRODUCT((cAcq_Selected="Yes")*(COUNTIFS(cAcq_SiteID,cAcq_SiteID,cAcq_Selected,"Yes")>1))', "High",
         "Keep one Selected = Yes per site."),
        ("DQ06", "Scope", "Mandatory or scope field blank (name, region, type, phase, batch, scenario, AA:AF, CW "
                          "delivery model)", "Master",
         'SUMPRODUCT((Master_B<>"")*(((Master_D="")+(Master_E="")+(Master_J="")+(Master_X="")+(Master_Y="")+'
         '(Master_Z="")+(Master_AA="")+(Master_AB="")+(Master_AC="")+(Master_AD="")+(Master_AE="")+(Master_AF="")'
         '+(SC_CW="Yes")*(Master_AG=""))>0))', "High", "Fill the field; blank never means N/A."),
        ("DQ07", "Scope", "Scope conflict: GECOL / HLC record for a site whose scope is N/A, or Fiber Status No Fiber "
                          "/ Through MW Link with HLC = Yes", "GECOL · HLC",
         'SUMPRODUCT((cPwr_Status="N/A")*(cPwr_PowerOnDateMeterCharging>0))+SUMPRODUCT((cFib_Status="N/A")*'
         '(cFib_InstallationDate>0))+SUMPRODUCT(((cFib_FiberStatus="No Fiber")+(cFib_FiberStatus="Through MW Link"))*'
         '(COUNTIFS(Master_B,cFib_SiteID,Master_AD,"Yes")>0))', "Medium", "Correct the scope in Master or the record."),
        ("DQ08", "Scope", "TE installed while MW scope is unknown (Configuration blank in As-Built)", "As-Built",
         "SUMPRODUCT(" + vv("CU") + '*(Master_R=""))', "Medium", "Fill Configuration (WL/MW/IP/PWR)."),
        ("DQ09", "KPI", "Cluster ID missing or several KPI source dates in one cluster", "RF & Design · Master",
         'SUMPRODUCT((Master_B<>"")*(SC_TE="Yes")*(((Master_EN="")+(Master_EQ="Multiple Sources"))>0))', "Medium",
         "One KPI Acceptance Source Date per cluster."),
        ("DQ10", "Chronology", "Negative SLA durations (finish before start, any step)", "SLA",
         'COUNTIF(SLA_Completed,"<0")', "High", "Correct the dates shown red in SLA."),
        ("DQ11", "Gates", "Actual date entered but the milestone is not achieved (gate or conditions not met)",
         "Master",
         'SUMPRODUCT((Master_B<>"")*((' + vv("BN") + '*(1-' + vv("BP") + ')+' + vv("CS") + '*(1-' + vv("CU") + ')+'
         + vv("DN") + '*(1-' + vv("DP") + ')+' + vv("EJ") + '*(1-' + vv("EL") + '))>0))', "Medium",
         "Complete the readiness conditions or correct the actual date."),
        ("DQ12", "Chronology", "Milestone achieved after an earlier applicable milestone that is still missing",
         "Master", 'SUMPRODUCT((Master_B<>"")*((' + "+".join(gaps) + ')>0))', "High", "Complete the earlier milestone."),
        ("DQ13", "Status", "Overall Status does not match the milestone chain", "Master",
         'SUMPRODUCT((Master_FZ="Completed")*(1-' + vv("FM") + '))+SUMPRODUCT((Master_FZ="Not Started")*'
         '(Master_FX<>"Not Started"))', "Medium", "Check Change Control and milestone dates."),
        ("DQ14", "Baseline", "Active site (In Progress / On Hold) without an approved baseline", "Baseline",
         'SUMPRODUCT((Master_B<>"")*(((Master_FZ="In Progress")+(Master_FZ="On Hold"))>0)*(COUNTIFS(cBLA_SiteID,'
         'Master_B,cBLA_Current,"Yes")=0))', "Medium", "Approve the proposal (Guide › Approve a baseline)."),
        ("DQ15", "Baseline", "Approved baseline rows without Version / Approval Date, or a site repeated in a version",
         "Baseline",
         'SUMPRODUCT((cBLA_SiteID<>"")*(((cBLA_Version="")+(cBLA_ApprovalDate="")+(COUNTIFS(cBLA_SiteID,cBLA_SiteID,'
         'cBLA_Version,cBLA_Version)>1))>0))', "Medium", "Complete Version and Approval Date."),
        ("DQ16", "Permits", "Permit type without Has Expiry (Setting)", "Permits & HSE · Setting",
         'SUMPRODUCT((cPrm_Type<>"")*(COUNTIFS(PAType,cPrm_Type,PATypeHasExpiry,"")>0))', "High",
         "Set Has Expiry Yes / No for the type in Setting."),
        ("DQ17", "Governance", "N/A without reason or approver", "Permits · Quality · Readiness",
         'COUNTIF(cPrm_Validity,"N/A (incomplete)")+COUNTIF(cQI_Status,"N/A (incomplete)")+SUMPRODUCT((cRdy_OMChecklist'
         '="N/A")*(((cRdy_NAReason="")+(cRdy_NAApprovedBy=""))>0))', "High", "Add N/A Reason and N/A Approved By."),
        ("DQ18", "Warehouse", "WH movement rows failing Check", "WH",
         'COUNTIF(cMov_Check,"?*")-COUNTIF(cMov_Check,"OK")', "High", "See the Check column of the row."),
        ("DQ19", "Warehouse", "Items with negative available quantity", "WH", 'COUNTIF(WHAvailable,"<0")', "High",
         "Correct the movement quantities or statuses."),
        ("DQ20", "Warehouse", "Serial rows failing Check", "WH Serials",
         'COUNTIF(cSer_Check,"?*")-COUNTIF(cSer_Check,"OK")', "Medium", "See the Check column of the row."),
        ("DQ21", "Procurement", "PR-PO order rows failing Check, contracts with date errors, contractor teams without "
                                "a registered contractor", "PR-PO",
         'COUNTIF(cOrd_Check,"?*")-COUNTIF(cOrd_Check,"OK")+COUNTIF(cCo_ContractValidity,"Date error")+COUNTIF('
         'cCo_ContractValidity,"Missing start")+SUMPRODUCT((cTeam_Team<>"")*(cTeam_TeamType="Contractor")*'
         '(COUNTIF(cCo_Contractor,cTeam_Contractor)=0))', "Medium", "See Check / Contract Validity."),
        ("DQ22", "Quality", "PAC dated with an open Major / Critical / unclassified snag", "Quality Inspections",
         'SUMPRODUCT((Master_B<>"")*(((' + vv("EJ") + ')+(' + vv("BV") + '))>0)*(COUNTIFS(cQI_SiteID,Master_B,'
         'cQI_RecordType,"Snag",cQI_Blocking,"Yes")>0))', "High", "Close or classify the snag before PAC."),
        ("DQ23", "Swap", "Swap / relocation control missing (MOP, approved window, Pre-Swap KPI baseline, legacy "
                         "handover)", "Master · Change Control · Readiness",
         'SUMPRODUCT((Master_B<>"")*' + swap + '*' + vv("DN") + '*(((1-' + vv("DC") + ')+(COUNTIFS(cChg_SiteID,'
         'Master_B,cChg_ChangeType,"Cutover Window",cChg_Approval,"Approved")=0)+(1-' + vv("DJ") + '))>0))+'
         'SUMPRODUCT((Master_B<>"")*' + swap + '*' + vv("FV") + '*(COUNTIFS(cRdy_SiteID,Master_B,cRdy_LegacyHandoverDate,'
         '">0")=0))', "Medium", "Record the missing control."),
        ("DQ24", "Issues", "Site issue incomplete (owner, raised date or status missing)", "Master",
         'SUMPRODUCT((Master_B<>"")*(((Master_GA<>"")*(((Master_GB="")+(Master_GC="")+(Master_GD=""))>0)+'
         '(Master_GA="")*(((Master_GB<>"")+(Master_GD<>""))>0))>0))', "Low", "Complete GA:GD."),
        ("DQ25", "Risks", "Risk Register row without Risk ID or with a repeated Risk ID", "Risk Register",
         'SUMPRODUCT((cRisk_RiskStatement<>"")*(((cRisk_RiskID="")+(COUNTIF(cRisk_RiskID,cRisk_RiskID)>1))>0))', "Low",
         "Give each risk a unique ID."),
        ("DQ26", "Lists", "Values outside their drop-down lists", "Master · RF & Design", outlist, "Medium",
         "Re-select the value or add it to the list in Setting."),
        ("DQ27", "Setting", "Region without WL Vendor (Setting) or Master region missing from the table", "Setting",
         'SUMPRODUCT((tblRegionVendor[Region]<>"")*(tblRegionVendor[WL Vendor]=""))+SUMPRODUCT((Master_E<>"")*'
         '(COUNTIF(tblRegionVendor[Region],Master_E)=0))', "High", "Set the WL Vendor for every region."),
    ]
    return C


def build_data_quality(wb, names):
    ws = wb.create_sheet("Data Quality")
    title(ws, "A1", "DATA QUALITY · rows breaking each rule (red = action needed)", 11)
    C = dq_checks()
    n = len(C)
    ws["A2"] = f'=COUNTIF(E4:E{3+n},">0")&" of {n} checks need action"'
    ws["A2"].font = font(10, True, "C00000")
    names.add("DQIssues", f"'Data Quality'!$H$2")
    ws["H2"] = f'=COUNTIF(E4:E{3+n},">0")'
    ws["H2"].font = font(9, color="808080")
    heads = ["Check ID", "Area", "Check", "Sheet", "Rows", "Severity", "How to fix"]
    for i, h in enumerate(heads):
        style(ws.cell(3, 1 + i, h), font(9, True), fill(C_DQ_HDR), AL_C)
    for k, (cid, area, chk, sh, f, sev, fix) in enumerate(C):
        r = 4 + k
        vals = [cid, area, chk, sh, "=IFERROR(" + f + ',"Error")', sev, fix]
        for i, x in enumerate(vals):
            c = ws.cell(r, 1 + i, x)
            style(c, font(9, i == 4), fill(C_CALC) if i == 4 else None, AL_LW if i in (2, 6) else AL_C)
    ws.conditional_formatting.add(f"E4:E{3+n}", CellIsRule(operator="greaterThan", formula=["0"], **red()))
    t = Table(displayName="tblDataQuality", ref=f"A3:G{3+n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    for col, w in zip("ABCDEFG", (8, 11, 60, 22, 7, 8, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    return ws


# ================================================================== Project Status (D65)
ACTS = [("C1", "Survey & Engineering", 12), ("C2", "Site Acquisition & Permits", 8), ("C3", "Procurement & Supply", 5),
        ("C4", "Civil Work & Site Readiness", 25), ("C5", "TE Installation", 20), ("C6", "Integration & On-Air", 12),
        ("C7", "Testing & Acceptance", 13), ("C8", "Training & Handover", 5)]
PSSEL = "PSSel"


def site_item(done, appl, plan, units=None):
    """Per-site item: done mask (or units sum), applicability mask, plan date column (Master letter)."""
    return ("site", done, appl, plan, units)


ITEMS_PS = [
    ("1.1", "C1", "TSSR approved", 30, "Units completed per applicable site", "RF & Design · TSSR Approved Date",
     site_item('(COUNTIFS(cRF_SiteID,Master_B,cRF_TSSRApprovedDate,">="&DateMin)>0)', "1", "U")),
    ("1.2", "C1", "Layout / STA / ST report approved", 25, "Units completed per applicable site",
     "Master AS · plan: baseline AV", site_item(v("AS"), CW, "AV")),
    ("1.3", "C1", "HLD", 15, "0 / 100 per deliverable unit", "Deliverables Log · HLD", ("log", "HLD")),
    ("1.4", "C1", "LLD (per site)", 20, "Units completed per applicable site",
     "RF & Design · LLD Approved Date · plan: Master U", site_item(vr("SC_LLD"), "1", "U")),
    ("1.5", "C1", "Technical document templates", 10, "0 / 100 per deliverable unit", "Deliverables Log · DOC",
     ("log", "DOC")),
    ("2.1", "C2", "Candidate selected", 30, "Units completed per applicable site",
     "Acquisition · Selection Date · plan: Master U", site_item('(COUNTIFS(cAcq_SiteID,Master_B,cAcq_Selected,"Yes",cAcq_SelectionDate,">="&DateMin)>0)', "1", "U")),
    ("2.2", "C2", "Site acquired", 50, "Units completed per applicable site", "Master T · plan: Master U",
     site_item(v("T"), "1", "U")),
    ("2.3", "C2", "CW access permission", 20, "Units completed per applicable site",
     "Permits & HSE · Access - CW (Master AL) · plan: baseline AV", site_item(v("AL"), CW, "AV")),
    ("3.1", "C3", "CW PO / WO issued", 15, "Units completed per applicable site", "Master AI · plan: baseline AV",
     site_item('(Master_AI="Done")', CW, "AV")),
    ("3.2", "C3", "BTS", 14.45, "Received ÷ Contract Qty", "WH · Progress Item Code BTS", ("supply", "BTS")),
    ("3.3", "C3", "PtP", 14.45, "Received ÷ Contract Qty", "WH · Progress Item Code PTP", ("supply", "PTP")),
    ("3.4", "C3", "CPE", 8.5, "Received ÷ Contract Qty", "WH · Progress Item Code CPE", ("supply", "CPE")),
    ("3.5", "C3", "Battery", 17, "Received ÷ Contract Qty", "WH · Progress Item Code BATTERY", ("supply", "BATTERY")),
    ("3.6", "C3", "Cabinet", 17, "Received ÷ Contract Qty", "WH · Progress Item Code CABINET", ("supply", "CABINET")),
    ("3.7", "C3", "Router", 8.5, "Received ÷ Contract Qty", "WH · Progress Item Code ROUTER", ("supply", "ROUTER")),
    ("3.8", "C3", "NMS", 3.4, "0 / 100 per deliverable unit", "Deliverables Log · NMS-S", ("log", "NMS-S")),
    ("3.9", "C3", "Core server", 1.7, "Received ÷ Contract Qty", "WH · Progress Item Code CORE-SERVER",
     ("supply", "CORE-SERVER")),
    ("3.10", "C3", "Billing system", 0, "0 / 100 per deliverable unit", "Deliverables Log · BILL", ("log", "BILL")),
    ("4.1", "C4", "Site adaption items (BD, BE, BF, BG, BJ)", 35, "Units completed (one per item)",
     "Master BD BE BF BG BJ = Done · plan: baseline BO", ("units", ["BD", "BE", "BF", "BG", "BJ"], '(Master_AB="Yes")', "BO")),
    ("4.2", "C4", "Tower erection items (BA, BB, BC)", 10, "Units completed (one per item)",
     "Master BA BB BC = Done · plan: baseline BO", ("units", ["BA", "BB", "BC"], '(Master_AA="Yes")', "BO")),
    ("4.3", "C4", "Power · internal AC / DC (BI)", 8, "Units completed per applicable site", "Master BI = Done",
     ("units", ["BI"], '(Master_AB="Yes")', "BO")),
    ("4.4", "C4", "Power · GECOL Power-On", 12, "Units completed per applicable site", "Master BW = Done · plan: BY",
     site_item('(Master_BW="Done")', '(Master_AC="Yes")', "BY")),
    ("4.5", "C4", "Fiber · site optical fiber (BH)", 6, "Units completed per applicable site", "Master BH = Done",
     ("units", ["BH"], '(Master_AB="Yes")', "BO")),
    ("4.6", "C4", "Fiber · HLC link installed", 9, "Units completed per applicable site", "Master BX = Done · plan: BY",
     site_item('(Master_BX="Done")', '(Master_AD="Yes")', "BY")),
    ("4.7", "C4", "Site RFI", 15, "Units completed per applicable site", "Master BZ · plan: BY",
     site_item(v("BZ"), RFI, "BY")),
    ("4.8", "C4", "CW PAC", 5, "Units completed per applicable site", "Master BV · plan: BY",
     site_item(v("BV"), CW, "BY")),
    ("5.1", "C5", "TE DN approved", 5, "Units completed per applicable site", "Master CH · plan: CT",
     site_item(v("CH"), TE, "CT")),
    ("5.2", "C5", "TE material on site (MOS)", 15, "Units completed per applicable site", "Master CJ · plan: CT",
     site_item(v("CJ"), TE, "CT")),
    ("5.3", "C5", "TE installation items (CK–CO)", 70, "Units completed (one per item)",
     "Master CK–CO = Done · plan: CT", ("units", ["CK", "CL", "CM", "CN", "CO"], TE, "CT")),
    ("5.4", "C5", "HW installation document (HWID)", 10, "Units completed per applicable site", "Master DT · plan: EK",
     site_item(v("DT"), TE, "EK")),
    ("6.1", "C6", "Pilot sites", 10, "0 / 100 per deliverable unit", "Deliverables Log · PILOT", ("log", "PILOT")),
    ("6.2", "C6", "Commissioning", 20, "Units completed per applicable site", "Master CW · plan: DO",
     site_item(v("CW"), TE, "DO")),
    ("6.3", "C6", "MW alignment", 10, "Units completed per applicable site", "Master CX · plan: DO",
     site_item(v("CX"), TE + "*" + MW, "DO")),
    ("6.4", "C6", "Integration", 15, "Units completed per applicable site", "Master DA · plan: DO",
     site_item(v("DA"), TE, "DO")),
    ("6.5", "C6", "On-Air", 35, "Units completed per applicable site", "Master DP · plan: DO",
     site_item(v("DP"), TE, "DO")),
    ("6.6", "C6", "NMS integration", 10, "0 / 100 per deliverable unit", "Deliverables Log · NMS-I", ("log", "NMS-I")),
    ("7.1", "C7", "Site standalone test (EIR) per system", 20, "Units completed (one per system)",
     "Master DU–DX = Done · plan: EK", ("units", ["DU", "DV", "DW", "DX"], TE, "EK")),
    ("7.2", "C7", "Network test (PAT) per system", 20, "Units completed (one per system)",
     "Master DZ–EC = Done · plan: EK", ("units", ["DZ", "EA", "EB", "EC"], TE, "EK")),
    ("7.3", "C7", "SSV", 10, "Units completed per applicable site", "Master DS = Pass · plan: EK",
     site_item('(Master_DS="Pass")', TE, "EK")),
    ("7.4", "C7", "As-Built", 10, "Units completed per applicable site", "Master EF · plan: EK",
     site_item(v("EF"), TE, "EK")),
    ("7.5", "C7", "TE PAC", 20, "Units completed per applicable site", "Master EL · plan: EK",
     site_item(v("EL"), TE, "EK")),
    ("7.6", "C7", "KPI Acceptance", 10, "Units completed per applicable site", "Master ES · plan: ER",
     site_item(v("ES"), TE, "ER")),
    ("7.7", "C7", "NMS acceptance", 10, "0 / 100 per deliverable unit", "Deliverables Log · NMS-A", ("log", "NMS-A")),
    ("8.1", "C8", "Vendor training", 50, "0 / 100 per deliverable unit", "Deliverables Log · TRN", ("log", "TRN")),
    ("8.2", "C8", "O&M Handover (HOTO)", 50, "Units completed per applicable site", "Master EY · plan: EX",
     site_item(v("EY"), TE, "EX")),
]


def ps_formulas(it):
    """Return (done, applicable, plan dates, planned units) formulas."""
    kind = it[0]
    rd = "N(RptDate)"
    if kind == "site":
        _, done, appl, plan, _u = it
        a = [PSSEL] + ([] if appl == "1" else [appl])
        d = "SUMPRODUCT(" + "*".join(a + [done]) + ")"
        ap = "SUMPRODUCT(" + "*".join(a) + ")"
        pdates = "SUMPRODUCT(" + "*".join(a + [v(plan)]) + ")"
        planned = "IF(" + rd + "=0,0,SUMPRODUCT(" + "*".join(a + [v(plan), "(Master_" + plan + "<=RptDate)"]) + "))"
        return d, ap, pdates, planned
    if kind == "units":
        _, cols, appl, plan = it
        a = [PSSEL] + ([] if appl == "1" else [appl])
        d = "SUMPRODUCT(" + "*".join(a) + "*(" + "+".join('(Master_' + c + '="Done")' for c in cols) + "))"
        ap = "SUMPRODUCT(" + "*".join(a) + "*(" + "+".join('(Master_' + c + '<>"N/A")' for c in cols) + "))"
        pdates = "SUMPRODUCT(" + "*".join(a + [v(plan)]) + "*(" + "+".join('(Master_' + c + '<>"N/A")' for c in cols) + "))"
        planned = ("IF(" + rd + "=0,0,SUMPRODUCT(" + "*".join(a + [v(plan), "(Master_" + plan + "<=RptDate)"]) + "*("
                   + "+".join('(Master_' + c + '<>"N/A")' for c in cols) + ")))")
        return d, ap, pdates, planned
    if kind == "log":
        code = '"' + it[1] + '"'
        d = 'COUNTIFS(cLog_Code,' + code + ',cLog_DoneDate,">="&DateMin,cLog_AcceptanceRef,"<>")'
        ap = "COUNTIF(cLog_Code," + code + ")"
        pdates = 'COUNTIFS(cLog_Code,' + code + ',cLog_PlanDate,">0")'
        planned = "IF(" + rd + '=0,0,COUNTIFS(cLog_Code,' + code + ',cLog_PlanDate,">0",cLog_PlanDate,"<="&RptDate))'
        return d, ap, pdates, planned
    if kind == "supply":
        code = '"' + it[1] + '"'
        d = "SUMIF(WHProgCode," + code + ",WHReceived)"
        ap = "SUMIF(WHProgCode," + code + ",WHContractQty)"
        pdates = 'SUMIFS(WHContractQty,WHProgCode,' + code + ',WHPlannedDelivery,">0")'
        planned = "IF(" + rd + '=0,0,SUMIFS(WHContractQty,WHProgCode,' + code + ',WHPlannedDelivery,">0",' \
                  'WHPlannedDelivery,"<="&RptDate))'
        return d, ap, pdates, planned
    raise ValueError(kind)


def build_project_status(wb, names):
    ws = wb.create_sheet("Project Status")
    ws.sheet_view.showGridLines = False
    widths = {"A": 7, "B": 40, "C": 18, "D": 34, "E": 9, "F": 10, "G": 9, "H": 8, "I": 9, "J": 9, "K": 9,
              "L": 9, "M": 9, "N": 12}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    title(ws, "A1", "PROJECT STATUS · weighted physical progress (actual vs planned) · scope ends at HOTO", 13)
    texts = [
        ("Measurement standard", "Physical progress per AACE / PMI practice: completed units of objective "
                                 "deliverables, weighted by relative effort (not money)."),
        ("Weight basis", "Activity and item weights are editable (yellow). Each activity's items total 100; activity "
                         "weights total 100. Weight 0 = Not weighted (excluded from progress and completeness)."),
        ("Credit rules", "Credit only on objective evidence: a valid date, status Done / Pass, or a log row with Done "
                         "Date and Acceptance Ref. Justified N/A removes the unit from the scope."),
        ("Planned progress", "Same weights applied to plan dates up to the Report Date: approved baseline dates, "
                             "planned Site Release (Master U), WH Planned Delivery Date and log Plan Dates."),
        ("Project status", "Scope = sites in Master for the selected phase / batch (cancelled excluded). FAC and "
                           "Final Closure are outside the measured scope."),
    ]
    for k, (a, t) in enumerate(texts):
        r = 3 + k
        style(ws.cell(r, 1, a), font(9, True), fill(C_DQ_HDR), AL_C)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.cell(r, 1).value = a
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=14)
        style(ws.cell(r, 3, t), font(9), fill(C_CALC), AL_LW)
        ws.row_dimensions[r].height = 30
    # selectors
    r = 9
    style(ws.cell(r, 1, "Project Phase"), font(9, True), fill(C_HDR), AL_C)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    style(ws.cell(r, 3), font(9), fill(C_INPUT), AL_C)
    style(ws.cell(r, 4, "Delivery Batch"), font(9, True), fill(C_HDR), AL_C)
    style(ws.cell(r, 5), font(9), fill(C_INPUT), AL_C)
    style(ws.cell(r, 7, "Report Date"), font(9, True), fill(C_HDR), AL_C)
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
    style(ws.cell(r, 8, '=IF(N(RptDate)=0,"Set in Setting",RptDate)'), font(9), fill(C_CALC), AL_C, DATE_FMT)
    for col, lst in (("C", "ProjectPhases"), ("E", "DeliveryBatches")):
        dv = DataValidation(type="list", formula1="=" + lst, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}{r}")
    names.add("PSPhase", "'Project Status'!$C$9")
    names.add("PSBatch", "'Project Status'!$E$9")
    names.add("PSSel", ("(Master!$B$5:$B$504<>\"\")*(Master!$FZ$5:$FZ$504<>\"Cancelled\")*(((PSPhase=\"\")+(Master!"
                       "$X$5:$X$504=PSPhase))>0)*(((PSBatch=\"\")+(Master!$Y$5:$Y$504=PSBatch))>0)").replace("$504", "$" + str(M1)))
    # result row
    r = 11
    heads = ["ACTUAL PROGRESS", "PLANNED PROGRESS", "SPI (Actual ÷ Planned)", "PLAN COVERAGE",
             "SELECTED PHASE CONTRACT SCOPE (sites)", "TOTAL CONTRACT (sites)"]
    pos = [1, 3, 5, 7, 9, 12]
    for p, h in zip(pos, heads):
        ws.merge_cells(start_row=r, start_column=p, end_row=r, end_column=p + 1)
        style(ws.cell(r, p, h), font(9, True, "FFFFFF"), fill(C_TITLE), AL_C)
        ws.merge_cells(start_row=r + 1, start_column=p, end_row=r + 1, end_column=p + 1)
    ws.row_dimensions[r].height = 28
    ws.row_dimensions[r + 1].height = 34
    # positions of summary table rows (set below)
    s0 = 17                    # summary header row
    first_act = s0 + 1
    last_act = s0 + len(ACTS)
    tot = last_act + 1
    big = dict(f=font(16, True, C_TITLE), al=AL_C)
    vals = [f"=F{tot}", f'=IF(ISNUMBER(H{tot}),H{tot},"n/a")', f'=IF(AND(ISNUMBER(A12),ISNUMBER(C12)),IF(C12>0,A12/C12,'
                                                                  f'"n/a"),"n/a")', "=PSCoverage",
            '=IF(PSPhase="","Select phase",IFERROR(INDEX(tblProjectScope[Contract Sites],MATCH(PSPhase,'
            'tblProjectScope[Project Phase],0)),""))',
            '=IFERROR(INDEX(tblProjectScope[Contract Sites],MATCH("Total Contract",tblProjectScope[Project Phase],0)),"")']
    fmts = [PCT_FMT, PCT_FMT, "0.00", PCT_FMT, "#,##0", "#,##0"]
    for p, f, fm in zip(pos, vals, fmts):
        c = ws.cell(r + 1, p, f)
        style(c, font(16, True, C_TITLE), fill(C_CALC), AL_C, fm)
    r = 14
    style(ws.cell(r, 1, "Scope check"), font(9, True), fill(C_DQ_HDR), AL_C)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=14)
    style(ws.cell(r, 3, '=IF(PSPhase="","Select Project Phase to compare registered sites with the contract scope.",'
                        'IF(N(I12)=0,"Contract Sites for "&PSPhase&" not set in Setting (tblProjectScope).",'
                        'IF(SUMPRODUCT(PSSel)=N(I12),"Sites in Master match the contract scope ("&SUMPRODUCT(PSSel)&").",'
                        '"Sites in Master: "&SUMPRODUCT(PSSel)&" · contract scope: "&I12&" · difference: "&'
                        '(SUMPRODUCT(PSSel)-I12))))'), font(9), fill(C_CALC), AL_L)
    # ---- summary by activity
    style(ws.cell(s0 - 1, 1, "SUMMARY BY ACTIVITY"), font(10, True, C_TITLE), None, AL_L, border=False)
    heads = ["#", "Activity", "", "Weight", "Actual %", "Weighted Actual", "Planned %", "Weighted Planned",
             "Variance (A − P)", "SPI"]
    for i, h in enumerate(heads):
        if h:
            hdr(ws, s0, 1 + i, h, h != "Weight")
    ws.merge_cells(start_row=s0, start_column=2, end_row=s0, end_column=3)
    ws.row_dimensions[s0].height = 30
    # ---- detail (computed rows first so we know the item rows)
    d0 = tot + 4
    style(ws.cell(d0 - 1, 1, "DETAIL BY ITEM · yellow = weight · grey = calculated"), font(10, True, C_TITLE), None,
          AL_L, border=False)
    dheads = ["Code", "Item", "Method", "Source (sheet · column) and plan basis", "Done", "Applicable", "Actual %",
              "Weight", "Weighted", "Plan Dates", "Planned Units", "Planned %", "Effective Weight", "Status"]
    for i, h in enumerate(dheads):
        hdr(ws, d0, 1 + i, h, h != "Weight")
    ws.row_dimensions[d0].height = 30
    r = d0
    act_rows = {}
    item_rows = {}
    for aid, aname, aw in ACTS:
        r += 1
        act_rows[aid] = r
        style(ws.cell(r, 1, aid[1:]), font(9, True), fill(C_CALC_HDR), AL_C)
        style(ws.cell(r, 2, aname), font(9, True), fill(C_CALC_HDR), AL_L)
        for c in range(3, 15):
            style(ws.cell(r, c), font(9, True), fill(C_CALC_HDR), AL_C)
        item_rows[aid] = []
        for it in ITEMS_PS:
            if it[1] != aid:
                continue
            r += 1
            item_rows[aid].append(r)
            code, _, name, w, meth, src, spec = it
            d, ap, pdates, planned = ps_formulas(spec)
            style(ws.cell(r, 1, code), font(9), None, AL_C)
            style(ws.cell(r, 2, name), font(9), None, AL_L)
            style(ws.cell(r, 3, meth), font(8), None, AL_LW)
            style(ws.cell(r, 4, src), font(8), None, AL_LW)
            cell(ws, r, 5, "=" + d)
            cell(ws, r, 6, "=" + ap)
            cell(ws, r, 7, f'=IF(H{r}=0,"Not weighted",IF(N(F{r})=0,"N/A",MIN(1,E{r}/F{r})))', PCT_FMT)
            style(ws.cell(r, 8, w), font(9), fill(C_INPUT), AL_C)
            cell(ws, r, 10, "=" + pdates)
            cell(ws, r, 11, "=" + planned)
            cell(ws, r, 12, f'=IF(NOT(ISNUMBER(G{r})),G{r},IF(J{r}<F{r},"n/a",MIN(1,K{r}/F{r})))', PCT_FMT)
            cell(ws, r, 14, f'=IF(H{r}=0,"Not weighted",IF(NOT(ISNUMBER(G{r})),"N/A",IF(G{r}>=1,"Complete",IF(E{r}>0,'
                            f'"In progress","Not started"))))')
            ws.row_dimensions[r].height = 24
    log_hdr = r + 4
    # weighted columns: item share of measured items within the activity
    for aid, rows in item_rows.items():
        ar = act_rows[aid]
        rng = lambda col: f"${col}${rows[0]}:${col}${rows[-1]}"
        meas = f"SUMPRODUCT({rng('H')}*ISNUMBER({rng('G')}))"
        for rr in rows:
            cell(ws, rr, 9, f'=IF(ISNUMBER(G{rr}),H{rr}/MAX(1E-9,{meas})*G{rr},"")', "0.000")
            cell(ws, rr, 13, f'=IF(ISNUMBER(G{rr}),H{rr}/MAX(1E-9,{meas})*INDEX($D${first_act}:$D${last_act},'
                             f'MATCH("{aid[1:]}",$A${first_act}:$A${last_act},0))/100,0)', "0.000")
        cell(ws, ar, 7, f'=IF({meas}=0,"N/A",SUM({rng("I")}))', PCT_FMT, bold=True, f_fill=C_CALC_HDR)
        cell(ws, ar, 8, f"=SUM({rng('H')})", bold=True, f_fill=C_CALC_HDR)
        cell(ws, ar, 10, f"=SUM({rng('J')})", bold=True, f_fill=C_CALC_HDR)
        cell(ws, ar, 12, f'=IF({meas}=0,"N/A",IF(COUNTIF({rng("L")},"n/a")>0,"n/a",SUMPRODUCT({rng("H")}*ISNUMBER('
                         f'{rng("G")}),{rng("L")})/{meas}))', PCT_FMT,
             bold=True, f_fill=C_CALC_HDR)
        cell(ws, ar, 13, f"=SUM({rng('M')})", "0.000", bold=True, f_fill=C_CALC_HDR)
    all_items = [x for v_ in item_rows.values() for x in v_]
    names.add("PSCoverage", f"'Project Status'!$N${tot}")
    # summary rows
    for k, (aid, aname, aw) in enumerate(ACTS):
        rr = first_act + k
        ar = act_rows[aid]
        style(ws.cell(rr, 1, aid[1:]), font(9), None, AL_C)
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
        style(ws.cell(rr, 2, aname), font(9, True), None, AL_L)
        style(ws.cell(rr, 4, aw), font(9), fill(C_INPUT), AL_C)
        cell(ws, rr, 5, f"=G{ar}", PCT_FMT)
        msum = f"SUMPRODUCT($D${first_act}:$D${last_act}*ISNUMBER($E${first_act}:$E${last_act}))"
        cell(ws, rr, 6, f'=IF(ISNUMBER(E{rr}),D{rr}/{msum}*E{rr},"n/a")', PCT_FMT)
        cell(ws, rr, 7, f"=L{ar}", PCT_FMT)
        psum = f"SUMPRODUCT($D${first_act}:$D${last_act}*ISNUMBER($G${first_act}:$G${last_act}))"
        cell(ws, rr, 8, f'=IF(AND(ISNUMBER(E{rr}),ISNUMBER(G{rr})),D{rr}/{psum}*G{rr},"n/a")', PCT_FMT)
        cell(ws, rr, 9, f'=IF(AND(ISNUMBER(E{rr}),ISNUMBER(G{rr})),E{rr}-G{rr},"n/a")', PCT_FMT)
        cell(ws, rr, 10, f'=IF(AND(ISNUMBER(E{rr}),ISNUMBER(G{rr})),IF(G{rr}>0,E{rr}/G{rr},"n/a"),"n/a")', "0.00")
    style(ws.cell(tot, 2, "TOTAL · overall progress"), font(9, True), fill(C_CALC_HDR), AL_L)
    ws.merge_cells(start_row=tot, start_column=2, end_row=tot, end_column=3)
    cell(ws, tot, 4, f"=SUM(D{first_act}:D{last_act})", bold=True, f_fill=C_CALC_HDR)
    cell(ws, tot, 6, f"=SUM(F{first_act}:F{last_act})", PCT_FMT, bold=True, f_fill=C_CALC_HDR)
    cell(ws, tot, 8, f'=IF(OR(COUNT(H{first_act}:H{last_act})=0,COUNT(H{first_act}:H{last_act})<COUNT('
                     f'F{first_act}:F{last_act})),"n/a",SUM(H{first_act}:H{last_act}))', PCT_FMT, bold=True,
         f_fill=C_CALC_HDR)
    cell(ws, tot, 9, f'=IF(AND(ISNUMBER(F{tot}),ISNUMBER(H{tot})),F{tot}-H{tot},"n/a")', PCT_FMT, bold=True,
         f_fill=C_CALC_HDR)
    cell(ws, tot, 10, f'=IF(AND(ISNUMBER(F{tot}),ISNUMBER(H{tot})),IF(H{tot}>0,F{tot}/H{tot},"n/a"),"n/a")', "0.00",
         bold=True, f_fill=C_CALC_HDR)
    items_f = [f"F{x}" for x in all_items]
    items_j = [f"J{x}" for x in all_items]
    style(ws.cell(tot, 13, "Plan coverage"), font(8, True), fill(C_CALC_HDR), AL_C)
    cell(ws, tot, 14, "=IFERROR((" + "+".join(items_j) + ")/MAX(1," + "+".join(items_f) + "),0)", PCT_FMT,
         bold=True, f_fill=C_CALC_HDR)
    style(ws.cell(tot + 1, 1, f'=IF(ROUND(D{tot},6)<>100,"Activity weights total "&D{tot}&" — must be 100.","Weights '
                              f'OK: activities total 100.")'), font(9, True, "C00000"), None, AL_L, border=False)
    # ---- Manual Deliverables Log
    style(ws.cell(log_hdr - 2, 1, "MANUAL DELIVERABLES LOG · one row per deliverable unit"), font(10, True, C_TITLE),
          None, AL_L, border=False)
    cols = [Col("Code", dv="RefProjectStatusCode", w=8), Col("Unit / Deliverable", w=30), Col("Site ID", dv="SiteIDs"),
            Col("Plan Date", dv=("date",), fmt="date"), Col("Done Date", dv=("date",), fmt="date"),
            Col("Remark", w=24), Col("Acceptance Ref", w=14)]
    register(ws, names, "tblProgressLog", "cLog", cols, sub_row=log_hdr - 1, hdr_row=log_hdr, nrows=20)
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    return ws


# ================================================================== finishing touches
def finish(wb, names):
    # Risk level colours (old file rule) and overdue review
    ws = wb["Risk Register"]
    rl = None
    for c in ws[2]:
        if c.value == "Risk Level":
            rl = c.column_letter
        if c.value == "Next Review Date":
            nr = c.column_letter
        if c.value == "Status":
            stc = c.column_letter
    last = ws.max_row
    ws.conditional_formatting.add(f"{rl}3:{rl}{last}", CellIsRule(operator="equal", formula=['"Critical"'], **red()))
    ws.conditional_formatting.add(f"{rl}3:{rl}{last}", CellIsRule(
        operator="equal", formula=['"High"'], fill=PatternFill("solid", fgColor="FF" + C_YEL_BG),
        font=Font(color="FF" + C_YEL_FG)))
    ws.conditional_formatting.add(f"{nr}3:{nr}{last}", FormulaRule(
        formula=[f'AND(N(RptDate)>0,${stc}3<>"Closed",N({nr}3)>0,{nr}3<RptDate)'], **red()))
    # Permits: grey PTW-only columns on non-PTW rows
    ws = wb["Permits & HSE"]
    hdrs = {c.value: c.column_letter for c in ws[2]}
    a, b = hdrs["JSA TRA"], hdrs["Stop Lifted Date"]
    ty = hdrs["Type"]
    ws.conditional_formatting.add(f"{a}3:{b}{ws.max_row}", FormulaRule(
        formula=[f'AND(${ty}3<>"",LEFT(${ty}3,3)<>"PTW")'], fill=PatternFill("solid", fgColor="FFD9D9D9")))
    vcol = hdrs["Validity"]
    ws.conditional_formatting.add(f"{vcol}3:{vcol}{ws.max_row}", CellIsRule(operator="equal", formula=['"Expired"'],
                                                                            **red()))
    ws.conditional_formatting.add(f"{vcol}3:{vcol}{ws.max_row}", CellIsRule(
        operator="equal", formula=['"Expiring"'], fill=PatternFill("solid", fgColor="FF" + C_YEL_BG),
        font=Font(color="FF" + C_YEL_FG)))
    # check columns red when not OK
    for sh, col_name in (("PR-PO", "Check"), ("WH", "Check"), ("WH Serials", "Check")):
        ws = wb[sh]
        hdr_row = 27 if sh == "WH" else (6 if sh == "WH Serials" else 2)
        for c in ws[hdr_row]:
            if c.value == col_name:
                L_ = c.column_letter
                ws.conditional_formatting.add(f"{L_}{hdr_row+1}:{L_}{ws.max_row}", FormulaRule(
                    formula=[f'AND({L_}{hdr_row+1}<>"",{L_}{hdr_row+1}<>"OK")'], **red()))
                break
    # fill Exceptions placeholders that need SLA column letters
    import calc_sheets as cs
    sla = wb["SLA"]
    c0, o0 = 4, 4 + len(cs.STEPS)
    cols = {}
    for k, st in enumerate(cs.STEPS):
        cols[st[0]] = (L(c0 + k), L(o0 + k))
    last = SLA_R1
    ds = wb["Dashboard"]
    open_over = "+".join(f'COUNTIF(SLA!{o}{cs.SLA_R0}:{o}{last},">"&N(SLA!{c}$7))' for c, o in cols.values())
    comp_over = "+".join(f'COUNTIF(SLA!{c}{cs.SLA_R0}:{c}{last},">"&N(SLA!{c}$7))' for c, o in cols.values())
    rep = {"{C2O}": f"{cols['C-2'][1]}{cs.SLA_R0}:{cols['C-2'][1]}{last}",
           "{C4O}": f"{cols['C-4'][1]}{cs.SLA_R0}:{cols['C-4'][1]}{last}",
           "{OPENOVER}": open_over, "{COMPOVER}": comp_over}
    for row in ds.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "{" in c.value:
                for k_, v_ in rep.items():
                    c.value = c.value.replace(k_, v_)
