"""Master sheet: literal copy of Rebuilt's Master (D29 exception: EZ header grey)
with every formula rewritten (D4)."""
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill

from common import NEW_THEME, copy_style_resolved, valid, C_CALC_HDR, fill, M0, M1

FIRST, LAST = M0, M1
B = "$B{r}"


def V(c):
    return "$" + c + "{r}"


def sc(col):
    """Same-row Site Calc cell (Site Calc rows are aligned with Master rows)."""
    return "'Site Calc'!$" + col + "{r}"


def idx(name, keyname, key=B):
    return "INDEX(" + name + ",MATCH(" + key + "," + keyname + ",0))"


def sidx(name, keyname, key=B):
    """Lookup that returns "" when the register has no row for the site (safe inside AND / OR)."""
    return 'IFERROR(' + idx(name, keyname, key) + ',"")'


def pull_text(name, keyname="cRF_SiteID"):
    return '=IF($B{r}="","",IFERROR(' + idx(name, keyname) + '&"",""))'


def pull_val(name, keyname="cRF_SiteID"):
    x = idx(name, keyname)
    return '=IF($B{r}="","",IFERROR(IF(' + x + '="","",' + x + '),""))'


def maxifs(rng, *crit):
    return "_xlfn.MAXIFS(" + rng + "," + ",".join(crit) + ")"


def minifs(rng, *crit):
    return "_xlfn.MINIFS(" + rng + "," + ",".join(crit) + ")"


def date_or_blank(expr):
    return "IF(" + expr + "=0,\"\"," + expr + ")"


# ---------------------------------------------------------------- hold / permit logic
def hold_adjust(cand):
    s = maxifs("cChg_ApprovalDate", "cChg_SiteID,$B{r}", 'cChg_ChangeType,"Suspension"',
               'cChg_Approval,"Approved"', 'cChg_ApprovalDate,"<="&' + cand)
    rr = maxifs("cChg_ApprovalDate", "cChg_SiteID,$B{r}", 'cChg_ChangeType,"Resumption"',
                'cChg_Approval,"Approved"', 'cChg_ApprovalDate,"<="&' + cand)
    rn = minifs("cChg_ApprovalDate", "cChg_SiteID,$B{r}", 'cChg_ChangeType,"Resumption"',
                'cChg_Approval,"Approved"', 'cChg_ApprovalDate,">"&' + s)
    return "IF(" + s + ">" + rr + ",IF(" + rn + ">0," + rn + ',""),' + cand + ")"


def permit_from(ptype):
    """First approval date of a permit type (0 if justified N/A, 99999 if missing)."""
    if ptype.startswith("PTW"):
        cnt = 'COUNTIFS(cPrm_SiteID,$B{r},cPrm_Type,"' + ptype + '",cPrm_HSEValidDate,">0")'
        first = minifs("cPrm_HSEValidDate", "cPrm_SiteID,$B{r}", 'cPrm_Type,"' + ptype + '"')
    else:
        cnt = ('COUNTIFS(cPrm_SiteID,$B{r},cPrm_Type,"' + ptype + '",cPrm_Status,"Approved",'
               'cPrm_ValidFrom,">0")')
        first = minifs("cPrm_ValidFrom", "cPrm_SiteID,$B{r}", 'cPrm_Type,"' + ptype + '"',
                       'cPrm_Status,"Approved"')
    na = ('COUNTIFS(cPrm_SiteID,$B{r},cPrm_Type,"' + ptype + '",cPrm_Status,"N/A",'
          'cPrm_NAReason,"<>",cPrm_NAApprovedBy,"<>")')
    return "IF(" + na + ">0,0,IF(" + cnt + "=0,99999," + first + "))"


def stop_block(ptype):
    return ('COUNTIFS(cPrm_SiteID,$B{r},cPrm_Type,"' + ptype + '",cPrm_StopWork,"Yes",'
            'cPrm_StopEndhelper,">="&DateMax)>0')


def gate(pred, permits=(), stop=None, extra=None, na_test=None):
    return '=IF($B{r}="","",IFERROR(' + gate_body(pred, permits, stop, extra, na_test) + ',""))'


def gate_body(pred, permits=(), stop=None, extra=None, na_test=None):
    parts = [pred] + [permit_from(p) for p in permits] + ([extra] if extra else [])
    cand = parts[0] if len(parts) == 1 else "MAX(" + ",".join(parts) + ")"
    body = 'IF(NOT(' + valid(pred) + '),"",IF(' + cand + '>DateMax,"",'
    if stop:
        body += 'IF(' + stop_block(stop) + ',"",' + hold_adjust(cand) + ')))'
    else:
        body += hold_adjust(cand) + '))'
    if na_test:
        body = 'IF(' + na_test + ',"N/A",' + body + ')'
    return body


def done(c):
    return 'OR(' + V(c) + '="Done",' + V(c) + '="N/A")'


def snag_ok(f, c):
    return ('OR(AND(' + V(f) + '="Yes",' + V(c) + '="Yes"),AND(' + V(f) + '="N/A",' + V(c) + '="N/A"))')


CWSCOPE = ('cOrd_SiteID,$B{r},cOrd_Scope,"<>GECOL",cOrd_Scope,"<>*Dismantle",cOrd_Scope,"<>TE Installation",'
           'cOrd_Scope,"<>ST*",cOrd_OrderStatus,"<>Cancelled",cOrd_OrderStatus,"<>Rejected"')


def cw_order_count(scope, datecol):
    return ('COUNTIFS(cOrd_SiteID,$B{r},cOrd_Scope,"' + scope + '",cOrd_OrderStatus,"<>Cancelled",'
            'cOrd_OrderStatus,"<>Rejected",' + datecol + ',">="&DateMin)')


def proc_status(datecol):
    tower = '(' + cw_order_count("Tower Erection", datecol) + '+' + cw_order_count("CW", datecol) + ')'
    adapt = '(' + cw_order_count("Site Adaption", datecol) + '+' + cw_order_count("CW", datecol) + ')'
    return tower, adapt


def first_order(field, test):
    return ('IFERROR(INDEX(cOrd_' + field + ',MATCH(1,INDEX((cOrd_SiteID=$B{r})*' + test +
            '*(cOrd_OrderStatus<>"Cancelled")*(cOrd_OrderStatus<>"Rejected"),0),0))&"","")')


CW_TEST = '(cOrd_Scope<>"GECOL")*(cOrd_Scope<>"TE Installation")*(cOrd_Scope<>"CW Dismantle")*(cOrd_Scope<>"TE Dismantle")*(cOrd_Scope<>"STA")*(cOrd_Scope<>"ST")'
TE_TEST = '(cOrd_Scope="TE Installation")'

BASELINE = {"AV": "PRECWCompletion", "BO": "CWCompletion", "BY": "SiteRFI", "CT": "TECompletion",
            "DO": "OnAir", "EK": "TEPAC", "ER": "KPIAcceptance", "EX": "HOTO", "FL": "FinalClosure"}


def formulas():
    F = {}
    rf_text = {"C": "SiteOwnerID", "D": "SiteName", "E": "Region", "F": "SubRegion", "G": "ClusterID",
               "J": "SiteType", "K": "SiteScenario", "L": "UplinkSite", "M": "TowerOwner",
               "N": "ExistingorNew", "O": "RTorGF", "P": "TowerType"}
    for c, f in rf_text.items():
        F[c] = pull_text("cRF_" + f)
    F["H"] = pull_val("cRF_Latitude")
    F["I"] = pull_val("cRF_Longitude")
    F["Q"] = pull_val("cRF_TowerHeightm")
    F["R"] = pull_text("cAB_ConfigurationWLMWIPPWR", "cAB_SiteID")
    F["S"] = pull_val("cRF_RFSurveyActualDate")
    F["T"] = ('=IF($B{r}="","",IFERROR(IF(INDEX(cAcq_AcquisitionCompletedDate,MATCH($B{r}&"|Yes",'
              'cAcq_SelectKeyhelper,0))="","",INDEX(cAcq_AcquisitionCompletedDate,MATCH($B{r}&"|Yes",'
              'cAcq_SelectKeyhelper,0))),""))')
    # ---- V Site Release (D32)
    tssr = sidx("cRF_TSSRApprovedDate", "cRF_SiteID")
    lld = sidx("cRF_LLDApprovedDate", "cRF_SiteID")
    los = sidx("cRF_MWLOSSurveyDate", "cRF_SiteID")
    losr = sidx("cRF_LOSResult", "cRF_SiteID")
    lb = sidx("cRF_LinkBudgetDate", "cRF_SiteID")
    caa_n = 'COUNTIFS(cPrm_SiteID,$B{r},cPrm_Type,"Civil Aviation Authority Permit")'
    caa = maxifs("cPrm_ValidFrom", "cPrm_SiteID,$B{r}", 'cPrm_Type,"Civil Aviation Authority Permit"',
                 'cPrm_Status,"Approved"')
    hub = 'OR($J{r}="HUB",$J{r}="Terminal")'
    F["V"] = ('=IF($B{r}="","",IFERROR(IF(AND(COUNTIF(Master_B,$B{r})=1,COUNTIFS(cAcq_SiteID,$B{r},'
              'cAcq_Selected,"Yes")=1,' + valid("$T{r}") + ',' + valid(tssr) + ',' + valid(lld) + ',OR('
              + caa_n + '=0,' + valid(caa) + '),OR(NOT(' + hub + '),AND(' + valid(los) + ',' + losr
              + '="Pass",' + valid(lb) + '))),MAX($T{r},' + tssr + ',' + lld + ',IF(' + caa_n + '>0,' + caa
              + ',0),IF(' + hub + ',MAX(' + los + ',' + lb + '),0)),""),""))')
    F["W"] = gate("$V{r}")
    # ---- procurement
    tw_pr, ad_pr = proc_status("cOrd_PRDate")
    tw_po, ad_po = proc_status("cOrd_POWODate")
    unresolved = 'OR($AA{r}="",$AB{r}="",$AC{r}="",$AD{r}="",$AE{r}="",$AF{r}="")'
    cw_na = 'AND($AA{r}="N/A",$AB{r}="N/A")'
    F["AH"] = ('=IF($B{r}="","",IF(' + unresolved + ',"Pending",IF(OR(' + cw_na + ',$AG{r}="In-House"),'
               '"N/A",IF($AG{r}="","Pending",IF(AND(OR($AA{r}<>"Yes",' + tw_pr + '>0),OR($AB{r}<>"Yes",'
               + ad_pr + '>0)),"Done","Pending")))))')
    F["AI"] = ('=IF($B{r}="","",IF(' + unresolved + ',"Pending",IF(' + cw_na + ',"N/A",IF(AND(OR($AA{r}'
               '<>"Yes",' + tw_po + '>0),OR($AB{r}<>"Yes",' + ad_po + '>0)),"Done","Pending"))))')
    F["AJ"] = '=IF($B{r}="","",' + first_order("Contractor", CW_TEST) + ')'
    F["AK"] = '=IF($B{r}="","",' + first_order("Team", CW_TEST) + ')'
    acc_cw = minifs("cPrm_ValidFrom", "cPrm_SiteID,$B{r}", 'cPrm_Type,"Access - CW"', 'cPrm_Status,"Approved"')
    F["AL"] = '=IF($B{r}="","",IFERROR(' + date_or_blank(acc_cw) + ',""))'
    # ---- AW PRE-CW Completion
    po_date = maxifs("cOrd_POWODate", CWSCOPE)
    analysis_ok = 'OR($AP{r}="N/A",AND(OR($AP{r}="STA",$AP{r}="ST"),OR($AQ{r}="Pass",$AQ{r}="Reinforce")))'
    cw_conds = ('AND(OR($AH{r}="Done",$AH{r}="N/A"),$AI{r}="Done",' + valid("$AL{r}") + ',' + valid("$AO{r}")
                + ',' + analysis_ok + ',' + valid("$AR{r}") + ',' + valid("$AS{r}") + ',' + valid("$AT{r}")
                + ',' + valid("$AU{r}") + ')')
    F["AW"] = ('=IF($B{r}="","",IFERROR(IF(NOT(' + valid("$W{r}") + '),"",IF(' + unresolved + ',"",IF('
               + sc("D") + '="N/A",$W{r},IF(' + cw_conds + ',MAX($W{r},' + po_date + ',$AL{r},$AO{r},$AR{r},'
               '$AS{r},$AU{r}),"")))),""))')
    F["AX"] = gate("$AW{r}", ("Access - CW", "PTW - CW"), stop="PTW - CW", na_test=sc("D") + '="N/A"')
    # ---- BP CW Completion
    tower_ok = 'OR($AA{r}<>"Yes",AND(' + ",".join(done(c) for c in ("BA", "BB", "BC")) + '))'
    adapt_ok = 'OR($AB{r}<>"Yes",AND(' + ",".join(done(c) for c in ("BD", "BE", "BF", "BG", "BH", "BI", "BJ")) + '))'
    F["BP"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("D") + '="N/A","N/A",IF(NOT(' + valid("$AX{r}") + '),"",IF(AND('
               + tower_ok + ',' + adapt_ok + ',' + valid("$BN{r}") + ',N($BN{r})>=N($AX{r})),$BN{r},""))),""))')
    F["BQ"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("F") + '="N/A","N/A",IF(' + sc("D") + '="Yes",'
               + gate_body("$BP{r}") + ',IF(' + sc("D") + '="N/A",' + gate_body("$AW{r}") + ',""))),""))')
    F["BW"] = ('=IF($B{r}="","",IF($AC{r}="N/A","N/A",IFERROR(INDEX(cPwr_Status,MATCH($B{r},cPwr_SiteID,0))&"",'
               'IF($AC{r}="Yes","Pending",""))))')
    F["BX"] = ('=IF($B{r}="","",IF($AD{r}="N/A","N/A",IFERROR(INDEX(cFib_Status,MATCH($B{r},cFib_SiteID,0))&"",'
               'IF($AD{r}="Yes","Pending",""))))')
    # ---- BZ Site RFI
    pwr_on = sidx("cPwr_PowerOnDateMeterCharging", "cPwr_SiteID")
    tx = sidx("cFib_TXLinkTestedDate", "cFib_SiteID")
    cw_pac_ok = 'OR(' + sc("D") + '<>"Yes",AND(' + valid("$BV{r}") + ',' + snag_ok("BS", "BT") + '))'
    F["BZ"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("F") + '="N/A","N/A",IF(NOT(' + valid("$BQ{r}") + '),"",IF(AND('
               + cw_pac_ok + ',OR($BW{r}="Done",$BW{r}="N/A"),OR($BX{r}="Done",$BX{r}="N/A"),OR($AD{r}<>"Yes",'
               + valid(tx) + ')),MAX($BQ{r},IF(' + sc("D") + '="Yes",$BV{r},0),IF($AC{r}="Yes",' + pwr_on
               + ',0),IF($AD{r}="Yes",' + tx + ',0)),""))),""))')
    # ---- CA TE Installation Readiness
    ca_pred = 'IF(' + sc("F") + '="N/A",$V{r},$BZ{r})'
    ca = gate("$CA_PRED{r}", ("Access - TE", "PTW - TE"), stop="PTW - TE", na_test=sc("E") + '="N/A"')
    F["CA"] = ca.replace("$CA_PRED{r}", ca_pred)
    acc_te = minifs("cPrm_ValidFrom", "cPrm_SiteID,$B{r}", 'cPrm_Type,"Access - TE"', 'cPrm_Status,"Approved"')
    F["CB"] = '=IF($B{r}="","",IFERROR(' + date_or_blank(acc_te) + ',""))'
    F["CE"] = '=IF($B{r}="","",' + first_order("Team", TE_TEST) + ')'
    # ---- CU TE Completion
    items = 'IF($AE{r}="Yes",AND(' + ",".join(done(c) for c in ("CK", "CL", "CM", "CN", "CO")) + '),$CL{r}="Done")'
    F["CU"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("E") + '="N/A","N/A",IF(NOT(' + valid("$CA{r}") + '),"",IF(AND('
               + items + ',' + valid("$CS{r}") + ',N($CS{r})>=N($CA{r})),$CS{r},""))),""))')
    psafe = sidx("cRdy_CommissioningPowerSafeDate", "cRdy_SiteID")
    F["CV"] = gate("$CU{r}", extra='IF(' + valid(psafe) + ',' + psafe + ',99999)')
    F["CZ"] = ('=IF(OR($B{r}="",NOT(ISNUMBER($CY{r}))),"",IFERROR(IF(N(' + idx("cAB_RSL", "cAB_SiteID")
               + ')=0,"",$CY{r}-' + idx("cAB_RSL", "cAB_SiteID") + '),""))')
    F["DG"] = ('=IF(OR($B{r}="",NOT(' + valid("$DD{r}") + '),NOT(' + valid("$DE{r}") + ')),"",'
               'IF(INT($DE{r})<INT($DD{r}),"Date Error",IF(INT($DE{r})=INT($DD{r}),0,'
               'NETWORKDAYS.INTL(INT($DD{r})+1,INT($DE{r}),7,Holidays))))')
    F["DH"] = ('=IF(OR($B{r}="",$Z{r}<>"Cold Swap",NOT(' + valid("$DP{r}") + ')),"",'
               'WORKDAY.INTL($DP{r},N(SLA_Max_S5),7,Holidays))')
    F["DI"] = ('=IF($B{r}="","",COUNTIFS(cChg_SiteID,$B{r},cChg_ChangeType,"Cutover Rollback",'
               'cChg_Approval,"Approved"))')
    # ---- DP On-Air
    uplink = sidx("cRdy_UplinkReadyDate", "cRdy_SiteID")
    freq = sidx("cRF_FrequencyLicenseDate", "cRF_SiteID")
    win = 'COUNTIFS(cChg_SiteID,$B{r},cChg_ChangeType,"Cutover Window",cChg_Approval,"Approved",cChg_ApprovalDate,">="&DateMin)'
    swap = 'OR($Z{r}="Cold Swap",$Z{r}="Hot Swap")'
    mw = 'N(' + sc("J") + ')>0'
    F["DP"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("E") + '="N/A","N/A",IF(NOT(' + valid("$CV{r}") + '),"",IF(AND('
               + valid("$DN{r}") + ',N($DN{r})>=N($CV{r}),' + valid("$CW{r}") + ',' + valid("$DA{r}") + ',OR(NOT('
               + mw + '),AND(' + valid("$CX{r}") + ',' + valid(freq) + ')),OR($L{r}="",$L{r}="N/A",' + valid(uplink)
               + '),OR(NOT(' + swap + '),AND(' + valid("$DC{r}") + ',' + win + '>0))),$DN{r},""))),""))')
    F["DQ"] = gate("$DP{r}", stop="PTW - TE")
    # ---- EL TE PAC
    tepac = ('AND(' + valid("$EJ{r}") + ',N($EJ{r})>=N($DQ{r}),OR($DS{r}="N/A",AND($DS{r}="Pass",' + valid("$DR{r}")
             + ')),' + valid("$DT{r}") + ',' + ",".join(done(c) for c in ("DU", "DV", "DW", "DX")) + ',' + valid("$DY{r}")
             + ',' + ",".join(done(c) for c in ("DZ", "EA", "EB", "EC", "ED")) + ',' + valid("$EE{r}") + ','
             + valid("$EF{r}") + ',' + snag_ok("EG", "EH") + ')')
    F["EL"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("E") + '="N/A","N/A",IF(NOT(' + valid("$DQ{r}") + '),"",IF('
               + tepac + ',$EJ{r},""))),""))')
    F["EM"] = gate("$EL{r}")
    F["EN"] = '=IF($B{r}="","",$G{r}&"")'
    F["EQ"] = ('=IF($B{r}="","",IF($EN{r}="","Missing Cluster ID",IF(COUNTIFS(Master_EN,$EN{r},Master_EP,">0")=0,'
               '"No Source",IF(COUNTIFS(Master_EN,$EN{r},Master_EP,">0")>1,"Multiple Sources","Valid"))))')
    src = maxifs("Master_EP", "Master_EN,$EN{r}")
    F["ES"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("E") + '="N/A","N/A",IF(OR(NOT(' + valid("$EM{r}") + '),$EQ{r}'
               '<>"Valid"),"",IF(AND(' + valid(src) + ',' + src + '>=N($EM{r})),' + src + ',""))),""))')
    F["ET"] = gate("$ES{r}")
    F["EW"] = pull_val("cRdy_HOTOPackageValidDate", "cRdy_SiteID")
    ops = sidx("cRdy_OpsAcceptanceValidDate", "cRdy_SiteID")
    F["EY"] = ('=IF($B{r}="","",IFERROR(IF(' + sc("E") + '="N/A","N/A",IF(NOT(' + valid("$ET{r}") + '),"",IF(AND('
               + valid("$EU{r}") + ',' + valid("$EV{r}") + ',' + valid("$EW{r}") + ',' + valid(ops) + ',' + ops
               + '>=N($ET{r})),' + ops + ',""))),""))')
    F["EZ"] = gate("$EY{r}")
    F["FA"] = '=IF(AND($B{r}<>"",' + valid("$BV{r}") + '),$BV{r},"")'
    F["FB"] = '=IF(' + valid("$FA{r}") + ',$FA{r}+CWDLPDays,"")'
    F["FF"] = '=IF(' + valid("$EL{r}") + ',$EL{r},"")'
    F["FG"] = '=IF(' + valid("$FF{r}") + ',EDATE($FF{r},TEDLPMonths),"")'
    F["FK"] = pull_val("cRdy_ObligationsClosureDate", "cRdy_SiteID")
    fin = sidx("cRdy_FinalClosure", "cRdy_SiteID")
    F["FM"] = ('=IF($B{r}="","",IFERROR(IF(NOT(' + valid("$EZ{r}") + '),"",IF(AND(' + valid(fin) + ',' + fin
               + '>=N($EZ{r})),' + fin + ',"")),""))')
    for c, f in BASELINE.items():
        m = maxifs("cBLA_" + f, "cBLA_SiteID,$B{r}", 'cBLA_Current,"Yes"')
        F[c] = '=IF($B{r}="","",IFERROR(' + date_or_blank(m) + ',""))'
    # ---- status
    order = [("FM", "Final Closure"), ("EY", "HOTO"), ("ES", "KPI Acceptance"), ("EL", "TE PAC"),
             ("DP", "On-Air"), ("CU", "TE Completion"), ("BZ", "Site RFI"), ("BP", "CW Completion"),
             ("AW", "PRE-CW Completion"), ("V", "Site Release")]
    cur = '"Not Started"'
    for c, n in reversed(order):
        cur = 'IF(' + valid(V(c)) + ',"' + n + '",' + cur + ')'
    F["FX"] = '=IF($B{r}="","",' + cur + ')'
    nxt = '""'
    for c, n in order:
        nxt = 'IF(AND(' + V(c) + '<>"N/A",NOT(' + valid(V(c)) + ')),"' + n + '",' + nxt + ')'
    F["FY"] = '=IF($B{r}="","",' + nxt + ')'
    hold_now = ('AND(N(RptDate)>0,' + maxifs("cChg_ApprovalDate", "cChg_SiteID,$B{r}", 'cChg_ChangeType,"Suspension"',
                                             'cChg_Approval,"Approved"', 'cChg_ApprovalDate,"<="&RptDate') + '>'
                + maxifs("cChg_ApprovalDate", "cChg_SiteID,$B{r}", 'cChg_ChangeType,"Resumption"',
                         'cChg_Approval,"Approved"', 'cChg_ApprovalDate,"<="&RptDate') + ')')
    F["FZ"] = ('=IF($B{r}="","",IF(COUNTIFS(cChg_SiteID,$B{r},cChg_ChangeType,"Cancellation",cChg_Approval,'
               '"Approved")>0,"Cancelled",IF(' + valid("$FM{r}") + ',"Completed",IF(' + hold_now
               + ',"On Hold",IF($FX{r}<>"Not Started","In Progress","Not Started")))))')
    return F


DV_LISTS = [
    ("X", "ProjectPhases"), ("B", "SiteIDs"), ("Y", "DeliveryBatches"), ("Z", "DeliveryScenario"),
    ("AA:AF BS:BU EG:EH FC:FD FH:FI FN FS", "YesNA"), ("AG", "DeliveryModel"), ("AM CF", "StaffNames"),
    ("AP", "AnalysisType"), ("AQ", "AnalysisResult"), ("BA:BJ CK:CO DU:DX DZ:ED", "CompletionStatus"),
    ("CD FO FT", "ContractorNames"), ("DF", "CutoverResult"), ("DL", "KPIResult"), ("DS", "SSVResult"),
    ("GB", "IssueOwners"), ("GD", "IssueStatus"),
]
DATE_COLS = ("U AN:AO AR:AU AY:AZ BK:BN BR BV CC CG:CJ CP:CS CW:CX DA:DE DJ:DK DM:DN DR DT DY EE:EF EI:EJ "
             "EO:EP EU:EV FE FJ FP:FR FU:FW GC")


def _sqref(spec):
    out = []
    for part in spec.split():
        if ":" in part:
            a, b = part.split(":")
            out.append(f"{a}{FIRST}:{b}{LAST}")
        else:
            out.append(f"{part}{FIRST}:{part}{LAST}")
    return " ".join(out)


def build(wb, names, src_path):
    src = openpyxl.load_workbook(src_path)["Master"]
    ws = wb.create_sheet("Master")
    maxc = CI("GE")
    for c in range(1, maxc + 1):
        col = L(c)
        sd = src.column_dimensions[col]
        ws.column_dimensions[col].width = sd.width
        ws.column_dimensions[col].hidden = sd.hidden
    for r in range(1, 5):
        ws.row_dimensions[r].height = src.row_dimensions[r].height
        for c in range(1, maxc + 1):
            s = src.cell(r, c)
            d = ws.cell(r, c)
            if r <= 3:
                d.value = s.value
            copy_style_resolved(s, d, NEW_THEME)
    for m in src.merged_cells.ranges:
        if m.max_row <= 4:
            ws.merge_cells(str(m))
    # data row styles (row 5 template)
    tmpl = []
    for c in range(1, maxc + 1):
        tmpl.append(src.cell(5, c))
    F = formulas()
    for r in range(FIRST, LAST + 1):
        for c in range(1, maxc + 1):
            d = ws.cell(r, c)
            copy_style_resolved(tmpl[c - 1], d, NEW_THEME)
            col = L(c)
            if c == 1:
                d.value = r - FIRST + 1
            elif col in F:
                d.value = F[col].format(r=r)
    # row 4: SUBTOTAL(3) over each column (spec)
    for c in range(1, maxc + 1):
        col = L(c)
        ws.cell(4, c).value = f"=SUBTOTAL(3,{col}{FIRST}:{col}{LAST})"
    # D29: EZ header grey
    for r in (1, 2, 3):
        cell = ws[f"EZ{r}"]
        if cell.coordinate in ws.merged_cells or True:
            cell.fill = fill(C_CALC_HDR)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:GE{LAST}"
    ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    # validations
    for spec, name in DV_LISTS:
        dv = DataValidation(type="list", formula1="=" + name, allow_blank=True)
        ws.add_data_validation(dv)
        for part in _sqref(spec).split():
            dv.add(part)
    dv = DataValidation(type="date", operator="between", formula1="DateMin", formula2="DateMax", allow_blank=True)
    ws.add_data_validation(dv)
    for part in _sqref(DATE_COLS).split():
        dv.add(part)
    dv = DataValidation(type="decimal", operator="between", formula1="-200", formula2="0", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(_sqref("CY"))
    # names Master_<col>
    for c in range(2, maxc + 1):
        col = L(c)
        names.add("Master_" + col, f"Master!${col}${FIRST}:${col}${LAST}")
    return ws
