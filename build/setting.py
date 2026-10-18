"""Setting sheet (D11, D13, D14, D73): parameters, holidays, staff and all lists."""
import datetime as dt

from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from common import AL_C, AL_L, AL_LW, C_DARK, C_INPUT, DATE_FMT, fill, font, style, title
from lists import L1

PARAMS = [
    ("Report Date", None, "date", "RptDate", "Reporting cut-off date. Open durations, overdue and plan-to-date use it."),
    ("CW DLP period", 90, "calendar days", "CWDLPDays", "CW defects liability period after CW PAC."),
    ("TE DLP period", 12, "months", "TEDLPMonths", "TE defects liability period after TE PAC."),
    ("Earliest valid date", dt.date(2020, 1, 1), "date", "DateMin", "Technical date validity only."),
    ("Latest valid date", dt.date(2035, 12, 31), "date", "DateMax", "Technical date validity only."),
    ("Expiring alert", 30, "calendar days", "ExpiringDays",
     "Permits, leases and contracts show Expiring this many days before expiry; issues older than this are aged."),
    ("Risk level Critical from", 15, "score", "RiskCritical", "Risk Score = Probability × Impact."),
    ("Risk level High from", 10, "score", "RiskHigh", "Below High and at or above Medium = Medium."),
    ("Risk level Medium from", 5, "score", "RiskMedium", "Below Medium = Low."),
    ("MW RSL tolerance", 2, "dB", "RSLTol", "MW link outside tolerance when |RSL Deviation| is above this value."),
]


def _hdr(ws, r, c, text):
    style(ws.cell(r, c, text), font(9, True, "FFFFFF"), fill(C_DARK), AL_C)


def _table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)


def build(wb, names):
    ws = wb.create_sheet("Setting")
    ws.row_dimensions[1].height = 31.5
    ws.row_dimensions[2].height = 24
    # ---- parameters A:D
    title(ws, "A1", "Parameters", 11)
    for i, h in enumerate(["Parameter", "Value", "Unit", "Used in / Rule"]):
        _hdr(ws, 2, 1 + i, h)
    for k, (p, v, u, nm, note) in enumerate(PARAMS):
        r = 3 + k
        style(ws.cell(r, 1, p), font(9, True), None, AL_C)
        c = ws.cell(r, 2, v)
        style(c, font(9), fill(C_INPUT), AL_C, DATE_FMT if u == "date" else "General")
        style(ws.cell(r, 3, u), font(9), None, AL_C)
        style(ws.cell(r, 4, note), font(9), None, AL_LW)
        names.add(nm, f"Setting!$B${r}")
    dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)",
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B3")
    for col, w in zip("ABCD", (26, 14.75, 12, 49)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["E"].width = 2
    # ---- holidays F:G
    title(ws, "F1", "Public Holidays", 11)
    _hdr(ws, 2, 6, "Date")
    _hdr(ws, 2, 7, "Holiday")
    for r in range(3, 23):
        style(ws.cell(r, 6), font(9), fill(C_INPUT), AL_C, DATE_FMT)
        style(ws.cell(r, 7), font(9), None, AL_C)
    _table(ws, "tblHolidays", "F2:G22")
    names.add("Holidays", "tblHolidays[Date]")
    ws.column_dimensions["F"].width = 14.75
    ws.column_dimensions["G"].width = 20
    ws.column_dimensions["H"].width = 2
    dvh = DataValidation(type="date", operator="between", formula1="DateMin", formula2="DateMax", allow_blank=True)
    ws.add_data_validation(dvh)
    dvh.add("F3:F22")
    # ---- staff I:L
    title(ws, "I1", "Project Staff", 11)
    for i, h in enumerate(["Name", "Role", "Region", "Status"]):
        _hdr(ws, 2, 9 + i, h)
    for r in range(3, 13):
        for c in range(9, 13):
            style(ws.cell(r, c), font(9), None, AL_C)
    _table(ws, "tblStaff", "I2:L12")
    names.add("StaffNames", "tblStaff[Name]")
    for col, nm in (("J", "StaffRole"), ("K", "Region"), ("L", "StaffStatus")):
        v = DataValidation(type="list", formula1="=" + nm, allow_blank=True)
        ws.add_data_validation(v)
        v.add(f"{col}3:{col}12")
    for col, w in zip("IJKL", (18, 22, 13, 9)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["M"].width = 2
    # ---- lists from column N, one header row
    title(ws, "N1", "Drop-down Lists · add values inside each table (the list grows automatically)", 11)
    c = 14
    pos = {}
    for tname, nm, hdr, rows in L1:
        width = len(hdr)
        for i, h in enumerate(hdr):
            _hdr(ws, 2, c + i, h)
        n = len(rows)
        for k, row in enumerate(rows):
            vals = row if isinstance(row, list) else [row]
            for i, v in enumerate(vals):
                style(ws.cell(3 + k, c + i, v), font(9), fill(C_INPUT) if (len(hdr) > 1 and i > 0) else None, AL_C)
        ref = f"{L(c)}2:{L(c + width - 1)}{2 + max(n, 1)}"
        _table(ws, tname, ref)
        pos[tname] = (c, 2 + max(n, 1))
        if nm:
            names.add(nm, f"{tname}[{hdr[0]}]")
        # list caption (row 1) shows the defined name used by drop-downs
        cap = ws.cell(1, c, nm or tname)
        cap.font = font(8, True, "595959")
        cap.alignment = AL_C
        for i in range(width):
            ws.column_dimensions[L(c + i)].width = 16
        c += width + 1
    names.add("PATypeHasExpiry", "tblPAType[Has Expiry]")
    names.add("ScopeMatchScope", "tblScopeMatch[Scope]")
    names.add("ScopeMatchSpec", "tblScopeMatch[Required Specialization]")
    names.add("UoMWhole", "tblUoM[Whole Number]")
    for tname, lst in (("tblRegionVendor", "WLVendor"), ("tblPAType", "YesNo"), ("tblUoM", "YesNo"),
                       ("tblProcScope", "YesNo")):
        c0, last = pos[tname]
        v = DataValidation(type="list", formula1="=" + lst, allow_blank=True)
        ws.add_data_validation(v)
        v.add(f"{L(c0 + 1)}3:{L(c0 + 1)}{last}")
    c0, last = pos["tblRegionVendor"]
    v = DataValidation(type="list", formula1="=Region", allow_blank=True)
    ws.add_data_validation(v)
    v.add(f"{L(c0)}3:{L(c0)}{last}")
    c0, last = pos["tblProjectScope"]
    v = DataValidation(type="date", operator="between", formula1="DateMin", formula2="DateMax", allow_blank=True)
    ws.add_data_validation(v)
    v.add(f"{L(c0 + 3)}3:{L(c0 + 3)}{last}")
    for r in range(3, last + 1):
        ws.cell(r, c0 + 3).number_format = DATE_FMT
    ws.freeze_panes = "A3"
    return ws
