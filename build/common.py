"""Shared helpers for the GSDN workbook builder: colours, styles, tables, names."""
import re
import colorsys
from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

# ---------------------------------------------------------------- colours
# Theme palettes (Excel theme index order: lt1, dk1, lt2, dk2, accent1..6)
NSITES = 500                      # site capacity: Master, Site Calc, SLA and Baseline rows
M0, M1 = 5, 5 + NSITES - 1        # Master / Site Calc / Baseline data rows
SLA_R0, SLA_R1 = 13, 13 + NSITES - 1

OLD_THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "5B9BD5", "ED7D31",
             "A5A5A5", "FFC000", "4472C4", "70AD47", "0563C1", "954F72"]
NEW_THEME = ["FFFFFF", "000000", "E8E8E8", "0E2841", "156082", "E97132",
             "196B24", "0F9ED5", "A02B93", "4EA72E", "467886", "96607D"]


def tint(hex6, t):
    """Apply an Excel theme tint to an RGB hex colour."""
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if t < 0:
        l = l * (1 + t)
    elif t > 0:
        l = l * (1 - t) + t
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def theme_rgb(palette, idx, t=0.0):
    return tint(palette[idx], t or 0.0)


def resolve_color(color, palette):
    """Return an explicit ARGB string for an openpyxl Color, or None."""
    if color is None:
        return None
    try:
        if color.type == "rgb":
            v = color.rgb
            return v if isinstance(v, str) else None
        if color.type == "theme":
            return "FF" + theme_rgb(palette, color.theme, color.tint)
        if color.type == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            if color.indexed is not None and color.indexed < len(COLOR_INDEX):
                return COLOR_INDEX[color.indexed]
    except Exception:
        return None
    return None


# Colours of the old file (resolved from its Office 2013 theme)
C_HDR = theme_rgb(OLD_THEME, 5, 0.4)        # register input header
C_CALC_HDR = theme_rgb(OLD_THEME, 2, -0.25)  # Background 2, Darker 25% (D16)
C_CALC = theme_rgb(OLD_THEME, 0, -0.05)     # calculated data cell / subtotal row
C_TITLE = theme_rgb(OLD_THEME, 4, -0.5)     # sheet titles
C_DARK = theme_rgb(OLD_THEME, 4, -0.5)      # Setting header fill
C_INPUT = theme_rgb(OLD_THEME, 7, 0.8)      # yellow input cell
C_SLA_HDR = theme_rgb(OLD_THEME, 0, -0.25)
C_SLA_BLOCK = theme_rgb(OLD_THEME, 1, 0.35)
C_DQ_HDR = theme_rgb(OLD_THEME, 5, 0.6)
C_GROUP = theme_rgb(OLD_THEME, 4, 0.6)      # dashboard section band
C_RED_BG, C_RED_FG = "FFC7CE", "9C0006"
C_YEL_BG, C_YEL_FG = "FFF2CC", "7F6000"
C_GRN_BG, C_GRN_FG = "C6EFCE", "006100"

FONT = "Calibri"
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "[$-en-US]mmmm d, yyyy"
PCT_FMT = "0.0%"


def fill(hex6):
    return PatternFill("solid", fgColor="FF" + hex6)


def font(sz=9, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=sz, bold=bold, color="FF" + color, italic=italic)


AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=False)
AL_LW = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style(cell, f=None, fl=None, al=None, nf=None, border=True):
    if f is not None:
        cell.font = f
    if fl is not None:
        cell.fill = fl
    if al is not None:
        cell.alignment = al
    if nf is not None:
        cell.number_format = nf
    if border:
        cell.border = BORDER


def title(ws, ref, text, sz=12):
    c = ws[ref]
    c.value = text
    c.font = font(sz, True, C_TITLE)
    c.alignment = AL_L


# ---------------------------------------------------------------- names
class Names:
    def __init__(self, wb):
        self.wb = wb
        self.seen = set()

    def add(self, name, ref):
        if name in self.seen:
            raise ValueError("duplicate name " + name)
        self.seen.add(name)
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)


def key(s):
    """Column name -> identifier fragment for defined names."""
    return "".join(ch for ch in s if ch.isalnum())


# ---------------------------------------------------------------- columns
class Col:
    """One register column.

    kind: 'in' (input) or 'calc' (formula, grey header).
    f: formula template using {r} for the row number (A1 references).
    dv: defined name for a list validation, or ('date',) / ('whole',) / ('decimal',)
    fmt: 'date' | 'pct' | 'int' | number format string | None
    """

    def __init__(self, name, kind="in", f=None, dv=None, fmt=None, w=None, nm=None):
        self.name, self.kind, self.f, self.dv, self.fmt, self.w, self.nm = (
            name, kind, f, dv, fmt, w, nm)


def num_fmt(fmt):
    return {"date": DATE_FMT, "pct": PCT_FMT, "int": "0", "dec": "0.00",
            None: "General"}.get(fmt, fmt)


def default_width(col):
    if col.w:
        return col.w
    if col.fmt == "date":
        return 14.75
    n = len(col.name)
    return max(8.5, min(22, n * 0.85 + 2))


def add_dv(ws, col, rng):
    dv = col.dv
    if dv is None:
        return
    if isinstance(dv, tuple):
        kind = dv[0]
        if kind == "date":
            v = DataValidation(type="date", operator="between",
                               formula1="DateMin", formula2="DateMax", allow_blank=True)
        elif kind == "whole":
            v = DataValidation(type="whole", operator="greaterThanOrEqual",
                               formula1="0", allow_blank=True)
        elif kind == "decimal":
            v = DataValidation(type="decimal", operator="between",
                               formula1=str(dv[1]), formula2=str(dv[2]), allow_blank=True)
        else:
            raise ValueError(kind)
    else:
        v = DataValidation(type="list", formula1="=" + dv, allow_blank=True)
    v.showErrorMessage = True
    ws.add_data_validation(v)
    v.add(rng)


def tref(name):
    """Escape a column name for a structured reference."""
    out = ""
    for ch in name:
        if ch in "[]#'":
            out += "'" + ch
        else:
            out += ch
    return out


def register(ws, names, tname, prefix, cols, first_col=1, sub_row=1, hdr_row=2,
             nrows=10, data=None, hdr_fill=None, freeze=None, header_h=39.95):
    """Write an Excel Table register: SUBTOTAL row, header row, data rows.

    Returns dict column-name -> column letter.
    """
    first = hdr_row + 1
    last = hdr_row + max(nrows, len(data or []))
    letters = {c.name: L(first_col + i) for i, c in enumerate(cols)}
    for c in cols:
        if callable(c.f):
            c.f = c.f(letters)
    for i, c in enumerate(cols):
        ci = first_col + i
        col = L(ci)
        h = ws.cell(hdr_row, ci, c.name)
        hf = fill(C_CALC_HDR) if c.kind == "calc" else fill(hdr_fill or C_HDR)
        if c.kind == "num":
            hf = fill(hdr_fill or C_HDR)
        style(h, font(9, False), hf, AL_C)
        # subtotal row (D45): counts non-blank visible values; sub_row=None writes none
        if sub_row:
            s = ws.cell(sub_row, ci)
            if c.fmt in ("date", "int", "dec", "pct") or c.kind == "calc" and c.fmt:
                s.value = f"=SUBTOTAL(102,{col}{first}:{col}{last})"
            else:
                s.value = f"=SUBTOTAL(103,{col}{first}:{col}{last})"
            style(s, font(9), fill(C_CALC), AL_C)
        ws.column_dimensions[col].width = default_width(c)
        for r in range(first, last + 1):
            cell = ws.cell(r, ci)
            if c.f:
                cell.value = c.f.format(r=r)
                fl_ = None if c.kind == "num" else fill(C_CALC)
                style(cell, font(9), fl_, AL_C, num_fmt(c.fmt))
            else:
                style(cell, font(9), None, AL_C, num_fmt(c.fmt))
        add_dv(ws, c, f"{col}{first}:{col}{last}")
        if c.name != "#":
            nm = c.nm or (prefix + "_" + key(c.name))
            names.add(nm, f"{tname}[{tref(c.name)}]")
    if data:
        for ri, row in enumerate(data):
            for ci, v in enumerate(row):
                if v is not None and not cols[ci].f:
                    ws.cell(first + ri, first_col + ci, v)
    ref = f"{L(first_col)}{hdr_row}:{L(first_col + len(cols) - 1)}{last}"
    t = Table(displayName=tname, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    t._initialise_columns()
    for tc, c in zip(t.tableColumns, cols):
        tc.name = c.name
        if c.f:
            tc.calculatedColumnFormula = TableFormula(attr_text=c.f.format(r=first)[1:])
    ws.add_table(t)
    ws.row_dimensions[hdr_row].height = header_h
    if freeze:
        ws.freeze_panes = freeze
    return letters


def copy_style_resolved(src, dst, palette):
    """Copy a cell style from another workbook, resolving theme colours to RGB."""
    f = src.font
    col = resolve_color(f.color, palette) if f.color is not None else None
    dst.font = Font(name=f.name, size=f.sz, bold=f.b, italic=f.i, underline=f.u,
                    strike=f.strike, color=col, vertAlign=f.vertAlign)
    fl = src.fill
    if fl is not None and fl.fill_type == "solid":
        c = resolve_color(fl.fgColor, palette)
        if c and c != "00000000":
            dst.fill = PatternFill("solid", fgColor=c)
    b = src.border

    def side(s):
        if s is None or s.style is None:
            return Side()
        return Side(style=s.style, color=resolve_color(s.color, palette) if s.color else None)
    dst.border = Border(left=side(b.left), right=side(b.right), top=side(b.top),
                        bottom=side(b.bottom))
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


# ---------------------------------------------------------------- formula bits
def valid(x):
    """Excel test: x holds a valid date (between DateMin and DateMax)."""
    return f"AND(N({x})>=DateMin,N({x})<=DateMax)"


def wd(s, e):
    """Signed working days from s to e: working days after s up to e (same day = 0, reversed = negative)."""
    return (f"IF(INT({e})<INT({s}),-NETWORKDAYS.INTL(INT({e})+1,INT({s}),7,Holidays),"
            f"IF(INT({e})=INT({s}),0,NETWORKDAYS.INTL(INT({s})+1,INT({e}),7,Holidays)))")


import re as _re


def T(template):
    """Formula template: [[Column Name]] -> $<letter>{r} of the same table."""
    def build(lt):
        def rep(m):
            return "$" + lt[m.group(1)] + "{r}"
        return _re.sub(r"\[\[([^\]]+)\]\]", rep, template)
    return build


def rf_lookup(field, key="$B{r}"):
    return f"INDEX(cRF_{field},MATCH({key},cRF_SiteID,0))"



_SAME_ROW = re.compile(r"(?<![!\w$])\$([A-Z]{1,3})\{r\}")


def shift_cols(f, n=1):
    """Move same-sheet same-row references ($X{r}) n columns right; other-sheet references (Sheet!$X{r}) stay."""
    from openpyxl.utils import column_index_from_string as _ci
    return _SAME_ROW.sub(lambda m: "$" + L(_ci(m.group(1)) + n) + "{r}", f)
