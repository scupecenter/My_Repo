"""Site Calc (D72), SLA (D64) and Baseline (D63)."""
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L

from common import (AL_C, AL_L, AL_LW, BORDER, C_CALC, C_CALC_HDR, C_HDR, C_RED_BG, C_RED_FG,
                    C_SLA_BLOCK, C_SLA_HDR, C_INPUT, Col, DATE_FMT, PCT_FMT, T, fill, font,
                    register, shift_cols, style, title, valid, wd, M0, M1, SLA_R0)


# ---------------------------------------------------------------- Site Calc
def build_site_calc(wb, names, wh_items):
    ws = wb.create_sheet("Site Calc")
    title(ws, "A1", "SITE CALC · shared per-site results (read only, rows aligned with Master)")
    ws["A2"] = ("Calculated automatically. Do not edit. Each row reads the Master row with the same number.")
    ws["A2"].font = font(9, italic=True)
    # header row 3, SUBTOTAL + filter row 4, data from row 5 (aligned with Master); A = row number
    hdr = ["#", "Site ID", "In Selection", "CW Applies", "TE Applies", "RFI Applies", "LLD Approved Date",
           "Cambium Sector Need", "Intracom Sector Need", "MW Half-link Need", "TE Dispatch Complete Date"]
    for i, h in enumerate(hdr):
        c = ws.cell(3, i + 1, h)
        style(c, font(9), fill(C_CALC_HDR), AL_C)
        ws.cell(4, i + 1).value = (f"=SUBTOTAL(102,{L(i+1)}5:{L(i+1)}{M1})" if i in (6, 7, 8, 9, 10)
                                   else f"=SUBTOTAL(103,{L(i+1)}5:{L(i+1)}{M1})")
        style(ws.cell(4, i + 1), font(9), fill(C_CALC), AL_C)
        ws.column_dimensions[L(i + 1)].width = 5 if i == 0 else 14.75 if i in (6, 10) else 12
    ws.cell(4, 1).value = f"=SUBTOTAL(103,B5:B{M1})"
    ws.row_dimensions[3].height = 39.95
    # unit -> item names (complements) from WH header
    by_unit = {}
    for it in wh_items:
        if it.get("unit"):
            by_unit.setdefault(it["unit"], []).append(it["nm"])

    def net(nm):
        return (f"(SUMIFS({nm},cMov_SiteID,$A{{r}},cMov_Type,\"TE DN\",cMov_Status,\"Dispatched\")"
                f"-SUMIFS({nm},cMov_SiteID,$A{{r}},cMov_Type,\"TE RN\",cMov_Status,\"Received\"))")
    conds = []
    need = {"Cambium Sector": "$G{r}", "Intracom Sector": "$H{r}", "MW Half-link": "$I{r}",
            "Cabinet": "1", "Battery": "2", "LPU Kit": "2"}
    for u, nms in by_unit.items():
        if u == "Router":
            conds.append("(" + "+".join(net(n) for n in nms) + ")>=1")
        else:
            for n in nms:
                conds.append(net(n) + ">=" + need[u])
    vend = ('IFERROR(INDEX(tblRegionVendor[WL Vendor],MATCH(Master!$E{r},tblRegionVendor[Region],0))&"","")')
    sectors = "N(INDEX(cRF_Sectors,MATCH($A{r},cRF_SiteID,0)))"
    far = 'INDEX(cRF_MWFarendSite,MATCH($A{r},cRF_SiteID,0))&""'
    last_dn = ('_xlfn.MAXIFS(cMov_Date,cMov_SiteID,$A{r},cMov_Type,"TE DN",cMov_Status,"Dispatched")')
    F = [
        '=IF(Master!$B{r}="","",Master!$B{r}&"")',
        '=IF($A{r}="","",IF(AND(OR(DashPhase="",Master!$X{r}=DashPhase),OR(DashBatch="",Master!$Y{r}=DashBatch),'
        'OR(DashScenario="",Master!$Z{r}=DashScenario)),1,0))',
        '=IF($A{r}="","",IF(OR(Master!$AA{r}="Yes",Master!$AB{r}="Yes"),"Yes",IF(AND(Master!$AA{r}="N/A",'
        'Master!$AB{r}="N/A"),"N/A","")))',
        '=IF($A{r}="","",IF(OR(Master!$AE{r}="Yes",Master!$AF{r}="Yes"),"Yes",IF(AND(Master!$AE{r}="N/A",'
        'Master!$AF{r}="N/A"),"N/A","")))',
        '=IF($A{r}="","",IF(OR(Master!$AA{r}="Yes",Master!$AB{r}="Yes",Master!$AC{r}="Yes",Master!$AD{r}="Yes"),'
        '"Yes",IF(AND(Master!$AA{r}="N/A",Master!$AB{r}="N/A",Master!$AC{r}="N/A",Master!$AD{r}="N/A"),"N/A","")))',
        '=IF($A{r}="","",IFERROR(IF(INDEX(cRF_LLDApprovedDate,MATCH($A{r},cRF_SiteID,0))="","",'
        'INDEX(cRF_LLDApprovedDate,MATCH($A{r},cRF_SiteID,0))),""))',
        '=IF(OR($A{r}="",$D{r}<>"Yes"),0,IFERROR(IF(' + vend + '="Cambium",' + sectors + ',0),0))',
        '=IF(OR($A{r}="",$D{r}<>"Yes"),0,IFERROR(IF(' + vend + '="Intracom",' + sectors + ',0),0))',
        '=IF(OR($A{r}="",$D{r}<>"Yes"),0,IFERROR(IF(OR(' + far + '="",' + far + '="N/A"),0,1),0)'
        '+COUNTIF(cRF_MWFarendSite,$A{r}))',
        '=IF(OR($A{r}="",$D{r}<>"Yes"),"",IFERROR(IF(AND(' + ",".join(conds) + '),IF(' + last_dn + '=0,"",'
        + last_dn + '),""),""))',
    ]
    F = [shift_cols(f) for f in F]
    fmts = [None, "0", None, None, None, DATE_FMT, "0", "0", "0", DATE_FMT]
    for r in range(M0, M1 + 1):
        style(ws.cell(r, 1, r - M0 + 1), font(9), None, AL_C)
        for i, f in enumerate(F):
            c = ws.cell(r, i + 2, f.format(r=r))
            style(c, font(9), fill(C_CALC), AL_C, fmts[i] or "General")
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:K{M1}"
    for i, n in enumerate(["SC_Site", "SC_Sel", "SC_CW", "SC_TE", "SC_RFI", "SC_LLD", "SC_Cam",
                           "SC_Int", "SC_MW", "SC_Dispatch"]):
        names.add(n, f"'Site Calc'!${L(i+2)}$5:${L(i+2)}${M1}")
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    return ws


# ---------------------------------------------------------------- SLA
def mref(c):
    return "Master!$" + c + "{m}"


ALL = 'Master!$B{m}<>""'
CW = "'Site Calc'!$D{m}=\"Yes\""
TE = "'Site Calc'!$E{m}=\"Yes\""
RFI = "'Site Calc'!$F{m}=\"Yes\""
TOWER = 'Master!$AA{m}="Yes"'
ADAPT = 'AND(Master!$AB{m}="Yes",Master!$AA{m}<>"Yes")'
RT = 'Master!$O{m}="RT"'
GF = 'Master!$O{m}="GF"'
ANALYSIS = 'OR(Master!$AP{m}="STA",Master!$AP{m}="ST")'
MW = "N('Site Calc'!$J{m})>0"
HOT = 'Master!$Z{m}="Hot Swap"'
COLD = 'Master!$Z{m}="Cold Swap"'
SWAP = "OR(" + HOT + "," + COLD + ")"


def AND(*a):
    return "AND(" + ",".join(a) + ")"


# helper (source date) columns at the end of the SLA sheet
CWCRIT = ('cOrd_SiteID,Master!$B{m},cOrd_Scope,"<>GECOL",cOrd_Scope,"<>*Dismantle",cOrd_Scope,'
          '"<>TE Installation",cOrd_Scope,"<>ST*",cOrd_OrderStatus,"<>Cancelled",cOrd_OrderStatus,"<>Rejected"')


def _blank0(x):
    return "IF(" + x + "=0,\"\"," + x + ")"


def _lk(name, keyname):
    x = "INDEX(" + name + ",MATCH(Master!$B{m}," + keyname + ",0))"
    return 'IFERROR(IF(' + x + '="","",' + x + '),"")'


HELPERS = [
    ("CW PR Date", _blank0("_xlfn.MINIFS(cOrd_PRDate," + CWCRIT + ")")),
    ("CW PO / WO Date", _blank0("_xlfn.MAXIFS(cOrd_POWODate," + CWCRIT + ")")),
    ("GECOL PR Date", _blank0('_xlfn.MINIFS(cOrd_PRDate,cOrd_SiteID,Master!$B{m},cOrd_Scope,"GECOL",'
                              'cOrd_OrderStatus,"<>Cancelled")')),
    ("GECOL Installation Finish", _lk("cPwr_InstallationFinishDate", "cPwr_SiteID")),
    ("HLC Request Date", _lk("cFib_RequestDate", "cFib_SiteID")),
    ("HLC Installation Date", _lk("cFib_InstallationDate", "cFib_SiteID")),
    ("Search Ring Received", _lk("cRF_SearchRingReceivedDate", "cRF_SiteID")),
    ("First SAR", _blank0("_xlfn.MINIFS(cAcq_CandidateSubmittedDateSAR,cAcq_SiteID,Master!$B{m},"
                          "cAcq_CandidateSubmittedDateSAR,\">0\")")),
    ("TSSR Submitted", _lk("cRF_TSSRSubmittedDate", "cRF_SiteID")),
    ("TSSR Approved", _lk("cRF_TSSRApprovedDate", "cRF_SiteID")),
    ("Window Requested", _blank0('_xlfn.MAXIFS(cChg_RequestDate,cChg_SiteID,Master!$B{m},cChg_ChangeType,'
                                 '"Cutover Window")')),
    ("Window Approved", _blank0('_xlfn.MAXIFS(cChg_ApprovalDate,cChg_SiteID,Master!$B{m},cChg_ChangeType,'
                                '"Cutover Window",cChg_Approval,"Approved")')),
    ("E2E CW Start", 'IF(Master!$AG{m}="Outsourced",{H0},{H1})'),
    ("TE Dispatch Complete", "IF('Site Calc'!$K{m}=\"\",\"\",'Site Calc'!$K{m})"),
]


def H(i):
    return "{H" + str(i) + "}"


# (code, title, from->to, applies when, min, max, start, end, applicability)
STEPS = [
    ("A-1", "Site Search (First SAR)", "Search Ring Received → first Candidate Submitted (SAR)",
     "Sites with a search ring", 10, 20, H(6), H(7), ALL),
    ("A-3", "TSSR Approval", "TSSR Submitted → TSSR Approved", "All sites", 1, 2, H(8), H(9), ALL),
    ("B-2", "PO / WO (Procurement Lead Time)", "PR-PO: PR Date → PR-PO: PO / WO Date",
     "CW scope · Outsourced", 1, 2, H(0), H(1), AND(CW, 'Master!$AG{m}="Outsourced"')),
    ("B-3", "CW Access Permission", "PR-PO: PO / WO Date → CW Access Date", "CW scope", 1, 2, H(1), mref("AL"), CW),
    ("B-4", "CW Survey", "CW Access → CW Survey Actual", "CW scope", 1, 2, mref("AL"), mref("AO"), CW),
    ("B-5", "STA / ST Analysis", "CW Survey Actual → Report Submitted (Master AR)",
     "CW scope and Analysis Type STA / ST", 1, 2, mref("AO"), mref("AR"), AND(CW, ANALYSIS)),
    ("B-6", "Layout Submission", "Report Submitted (analysis sites) or CW Survey → Layout Submission (AR)",
     "CW scope", 1, 2, "IF(" + ANALYSIS + "," + mref("AR") + "," + mref("AO") + ")", mref("AR"), CW),
    ("B-7", "Layout Approval", "Layout Submission → Approval", "CW scope", 1, 2, mref("AR"), mref("AS"), CW),
    ("B-8", "Tower MOS", "Layout Approval → CW MOS", "Tower = Yes", 1, 2, mref("AS"), mref("AZ"), AND(CW, TOWER)),
    ("B-9", "Adaption MOS", "Layout Approval → CW MOS", "Adaption only", 1, 2, mref("AS"), mref("AZ"), AND(CW, ADAPT)),
    ("B-10a", "Tower Erection RT (GMT)", "CW MOS → CW Actual End", "Tower = Yes and RT", 14, 21,
     mref("AZ"), mref("BN"), AND(TOWER, RT)),
    ("B-10b", "Tower Erection GF (SST)", "CW MOS → CW Actual End", "Tower = Yes and GF", 30, 45,
     mref("AZ"), mref("BN"), AND(TOWER, GF)),
    ("B-11", "Site Adaption", "CW MOS → CW Actual End", "Adaption only", 4, 8, mref("AZ"), mref("BN"), ADAPT),
    ("B-13", "CW PAC", "CW Completion Date → CW PAC Actual", "CW scope", 1, 3, mref("BP"), mref("BV"), CW),
    ("B-15", "FAC", "CW DLP End → CW FAC Actual", "CW scope", 3, 10, mref("FB"), mref("FE"), CW),
    ("C-2", "GECOL (Power)", "GECOL PR Date → GECOL Installation Finish", "GECOL = Yes", 7, 10, H(2), H(3),
     'Master!$AC{m}="Yes"'),
    ("C-4", "HLC (Fiber)", "HLC Request Date → HLC Installation Date", "HLC = Yes", 7, 10, H(4), H(5),
     'Master!$AD{m}="Yes"'),
    ("RFI", "Site RFI", "RFI Readiness → Site RFI", "RFI applies; not Cold / Hot Swap", 1, 1,
     mref("BQ"), mref("BZ"), AND(RFI, "NOT(" + SWAP + ")")),
    ("E2E-B1", "Full CW · Adaption only", "PR Date (Outsourced) / PO / WO Date (In-House) → CW PAC Actual",
     "Adaption only", 11, 23, H(12), mref("BV"), ADAPT),
    ("E2E-B2", "Full CW · with Tower RT", "PR Date (Outsourced) / PO / WO Date (In-House) → CW PAC Actual",
     "Tower = Yes and RT", 22, 38, H(12), mref("BV"), AND(TOWER, RT)),
    ("E2E-B3", "Full CW · with Tower GF", "PR Date (Outsourced) / PO / WO Date (In-House) → CW PAC Actual",
     "Tower = Yes and GF", 38, 62, H(12), mref("BV"), AND(TOWER, GF)),
    ("D-1", "TE Access Permission", "Site Release → TE Access Date", "TE scope", 3, 5, mref("V"), mref("CB"), TE),
    ("D-2", "TE Survey", "TE Access → TE Survey", "TE scope", 1, 2, mref("CB"), mref("CC"), TE),
    ("D-3", "DN Submission", "TE Survey → DN Submitted", "TE scope", 1, 2, mref("CC"), mref("CG"), TE),
    ("D-4", "DN Approval", "DN Submitted → DN Approved", "TE scope", 1, 2, mref("CG"), mref("CH"), TE),
    ("D-5", "TE MOS", "TE Dispatch Complete → TE MOS Actual", "TE scope", 1, 1, H(13), mref("CJ"), TE),
    ("D-6", "TE Installation", "TE MOS → TE Actual End", "TE scope", 3, 5, mref("CJ"), mref("CS"), TE),
    ("D-7", "Commissioning", "TE Actual End → Commissioning", "TE scope", 1, 1, mref("CS"), mref("CW"), TE),
    ("D-8", "MW Alignment", "Commissioning → MW Alignment", "TE scope with MW", 1, 1, mref("CW"), mref("CX"),
     AND(TE, MW)),
    ("D-9", "Integration", "MW Alignment (Commissioning if no MW) → Integration", "TE scope", 1, 1,
     "IF(" + MW + "," + mref("CX") + "," + mref("CW") + ")", mref("DA"), TE),
    ("D-10", "On-Air", "Integration → On-Air Actual", "TE scope", 1, 1, mref("DA"), mref("DN"), TE),
    ("D-11", "SSV", "On-Air → SSV", "TE scope", 3, 5, mref("DN"), mref("DR"), TE),
    ("D-12", "HWID", "On-Air → HWID", "TE scope", 1, 1, mref("DN"), mref("DT"), TE),
    ("D-13", "EIR", "HWID → EIR", "TE scope", 3, 5, mref("DT"), mref("DY"), TE),
    ("D-14", "PAT", "EIR → PAT", "TE scope", 3, 5, mref("DY"), mref("EE"), TE),
    ("D-15", "As-Built", "PAT → As-Built", "TE scope", 1, 2, mref("EE"), mref("EF"), TE),
    ("D-16", "TE PAC", "As-Built (and SSV) → TE PAC Actual", "TE scope", 5, 10,
     "IF(AND(" + valid(mref("EF")) + "," + valid(mref("DR")) + "),MAX(" + mref("EF") + "," + mref("DR") + "),\"\")",
     mref("EJ"), TE),
    ("D-17", "O&M Handover", "KPI Acceptance → O&M Handover (HOTO)", "TE scope", 3, 5, mref("ES"), mref("EY"), TE),
    ("E2E-D1", "TE MOS to On-Air", "TE MOS → On-Air Actual", "TE scope", 8, 12, mref("CJ"), mref("DN"), TE),
    ("E2E-D2", "TE Access to O&M Handover", "TE Access → O&M Handover", "TE scope", 30, 51,
     mref("CB"), mref("EY"), TE),
    ("S-1", "MOP Approval", "MOP Submitted → MOP Approved", "Cold / Hot Swap", 1, 1, mref("DB"), mref("DC"), SWAP),
    ("S-2", "Window Approval", "Change Control (Cutover Window): Request Date → Approval Date", "Cold / Hot Swap",
     1, 1, H(10), H(11), SWAP),
    ("S-3", "Hot Swap Window", "Window Start → Window End", "Hot Swap", 1, 1, mref("DD"), mref("DE"), HOT),
    ("S-4", "Cold Swap Window", "Window Start → Window End", "Cold Swap", 1, 1, mref("DD"), mref("DE"), COLD),
    ("S-5", "Legacy Dismantle", "On-Air (cutover) → TE Dismantle Actual", "Cold Swap", 1, 1,
     mref("DN"), mref("FV"), COLD),
]



def code_key(code):
    return code.replace("-", "")


def build_sla(wb, names):
    ws = wb.create_sheet("SLA")
    title(ws, "A1", "SLA · working days (Friday / Saturday and holidays excluded) · Min / Max typed in rows 6-7 "
                    "are the only source of durations")
    n = len(STEPS)
    c0 = 4                    # completed block starts at D
    o0 = c0 + n               # open block
    h0 = o0 + n + 1           # helper block (one gap column)
    # left headers
    for c, t in ((1, "#"), (2, "Site ID"), (3, "Region")):
        style(ws.cell(2, c, t), font(9, True), fill(C_SLA_HDR), AL_C)
    labels = ["Step", "From → To", "Applies when", "Min (wd)", "Max (wd)", "Count", "Avg (wd)",
              "Within Max %", "Over Max"]
    for i, t in enumerate(labels):
        r = 3 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        style(ws.cell(r, 2, t), font(9, True, "FFFFFF"), fill(C_SLA_BLOCK), AL_C)
        style(ws.cell(r, 3), None, fill(C_SLA_BLOCK), AL_C)
        style(ws.cell(r, 1), None, fill(C_SLA_BLOCK), AL_C)
    ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + n - 1)
    style(ws.cell(2, c0, "COMPLETED · duration (wd) from start to finish"), font(9, True, "FFFFFF"),
          fill(C_SLA_BLOCK), AL_C)
    ws.merge_cells(start_row=2, start_column=o0, end_row=2, end_column=o0 + n - 1)
    style(ws.cell(2, o0, "OPEN · started, not finished: age (wd) to Report Date"), font(9, True, "FFFFFF"),
          fill(C_SLA_BLOCK), AL_C)
    ws.merge_cells(start_row=2, start_column=h0, end_row=2, end_column=h0 + len(HELPERS) - 1)
    style(ws.cell(2, h0, "SOURCE DATES (helper)"), font(9, True), fill(C_CALC_HDR), AL_C)
    last = SLA_R0 + (M1 - M0)
    rngs = {}
    for k, st in enumerate(STEPS):
        code, ttl, ft, ap, mn, mx = st[:6]
        cc, oc = L(c0 + k), L(o0 + k)
        for col, blk in ((cc, "C"), (oc, "O")):
            style(ws.cell(3, c0 + k if blk == "C" else o0 + k, f"{code} · {ttl}"), font(8, True), fill(C_SLA_HDR), AL_C)
            ci = c0 + k if blk == "C" else o0 + k
            style(ws.cell(4, ci, ft), font(8), fill(C_SLA_HDR), AL_C)
            style(ws.cell(5, ci, ap), font(8), fill(C_SLA_HDR), AL_C)
            rng = f"{col}{SLA_R0}:{col}{last}"
            if blk == "C":
                style(ws.cell(6, ci, mn), font(9), fill(C_INPUT), AL_C)
                style(ws.cell(7, ci, mx), font(9), fill(C_INPUT), AL_C)
            else:
                style(ws.cell(6, ci, f"={cc}6"), font(9), fill(C_CALC), AL_C)
                style(ws.cell(7, ci, f"={cc}7"), font(9), fill(C_CALC), AL_C)
            style(ws.cell(8, ci, f"=COUNT({rng})"), font(9), fill(C_CALC), AL_C)
            style(ws.cell(9, ci, f'=IF({col}8=0,"",AVERAGE({rng}))'), font(9), fill(C_CALC), AL_C, "0.0")
            style(ws.cell(10, ci, f'=IF(OR({col}8=0,N({col}7)=0),"",COUNTIF({rng},"<="&{col}7)/{col}8)'),
                  font(9), fill(C_CALC), AL_C, PCT_FMT)
            style(ws.cell(11, ci, f'=IF(N({col}7)=0,0,COUNTIF({rng},">"&{col}7))'), font(9), fill(C_CALC), AL_C)
            style(ws.cell(12, ci, f"=SUBTOTAL(102,{rng})"), font(9), fill(C_CALC), AL_C)
            ws.column_dimensions[col].width = 9.9
        names.add("SLA_Min_" + code_key(code), f"SLA!${cc}$6")
        names.add("SLA_Max_" + code_key(code), f"SLA!${cc}$7")
        rngs[code] = (cc, oc)
    for i, (hn, _) in enumerate(HELPERS):
        ci = h0 + i
        style(ws.cell(3, ci, hn), font(8, True), fill(C_CALC_HDR), AL_C)
        for r in range(4, 12):
            style(ws.cell(r, ci), None, fill(C_CALC_HDR), AL_C)
        style(ws.cell(12, ci, f"=SUBTOTAL(102,{L(ci)}{SLA_R0}:{L(ci)}{last})"), font(9), fill(C_CALC), AL_C)
        ws.column_dimensions[L(ci)].width = 12
    ws.row_dimensions[3].height = 45
    ws.row_dimensions[4].height = 78.75
    ws.row_dimensions[5].height = 60
    style(ws.cell(12, 1, f"=SUBTOTAL(103,B{SLA_R0}:B{last})"), font(9), fill(C_CALC), AL_C)
    hcols = {i: L(h0 + i) for i in range(len(HELPERS))}
    for r in range(SLA_R0, last + 1):
        m = r - SLA_R0 + M0
        hmap = {f"H{i}": f"${hcols[i]}{r}" for i in hcols}
        style(ws.cell(r, 1, r - SLA_R0 + 1), font(9), None, AL_C)
        style(ws.cell(r, 2, f'=IF(Master!$B{m}="","",Master!$B{m}&"")'), font(9), fill(C_CALC), AL_C)
        style(ws.cell(r, 3, f'=IF(Master!$B{m}="","",Master!$E{m}&"")'), font(9), fill(C_CALC), AL_C)
        for k, st in enumerate(STEPS):
            s, e, ap = st[6], st[7], st[8]
            s = s.format(m=m, **hmap)
            e = e.format(m=m, **hmap)
            ap = ap.format(m=m)
            comp = ('=IF(AND(' + ap + ',' + valid(s) + ',' + valid(e) + '),' + wd(s, e) + ',"")')
            opn = ('=IF(AND(' + ap + ',' + valid(s) + ',NOT(' + valid(e) + '),N(RptDate)>0),'
                   + wd(s, "RptDate") + ',"")')
            style(ws.cell(r, c0 + k, comp), font(9), None, AL_C, "0")
            style(ws.cell(r, o0 + k, opn), font(9), None, AL_C, "0")
        for i, (hn, hf) in enumerate(HELPERS):
            f = '=IF(Master!$B{m}="","",IFERROR(' + hf + ',""))'
            f = f.replace("{H0}", f"${hcols[0]}{r}").replace("{H1}", f"${hcols[1]}{r}").format(m=m)
            style(ws.cell(r, h0 + i, f), font(9), fill(C_CALC), AL_C, DATE_FMT)
    ws.freeze_panes = "D13"
    ws.auto_filter.ref = f"A12:{L(h0 + len(HELPERS) - 1)}{last}"
    comp_rng = f"{L(c0)}{SLA_R0}:{L(c0 + n - 1)}{last}"
    open_rng = f"{L(o0)}{SLA_R0}:{L(o0 + n - 1)}{last}"
    names.add("SLA_Completed", f"SLA!${L(c0)}${SLA_R0}:${L(c0 + n - 1)}${last}")
    names.add("SLA_Open", f"SLA!${L(o0)}${SLA_R0}:${L(o0 + n - 1)}${last}")
    # over-max colouring and negative durations (D68 #10)
    for k in range(n):
        for base in (c0, o0):
            col = L(base + k)
            ws.conditional_formatting.add(
                f"{col}{SLA_R0}:{col}{last}",
                FormulaRule(formula=[f'AND(ISNUMBER({col}{SLA_R0}),N({col}$7)>0,{col}{SLA_R0}>{col}$7)'],
                            fill=PatternFill("solid", fgColor="FF" + C_RED_BG),
                            font=Font(color="FF" + C_RED_FG)))
    ws.conditional_formatting.add(comp_rng, CellIsRule(operator="lessThan", formula=["0"],
                                                       fill=PatternFill("solid", fgColor="FFFF0000"),
                                                       font=Font(color="FFFFFFFF", bold=True)))
    return ws, rngs, (c0, o0, n, last)


# ---------------------------------------------------------------- Baseline
LINKS = [("B01", "Site Release", "CW PR / WO", 0, 1), ("B02", "CW PR / WO", "GECOL PR", 0, 1),
         ("B03", "CW PR / WO", "HLC Request", 0, 1), ("B04", "DN Approved", "TE Dispatch Complete", 0, 1),
         ("B05", "TE PAC", "KPI Acceptance", 1, 5), ("B06", "TE DLP End", "TE FAC", 3, 10),
         ("B07", "Last FAC / HOTO", "Final Closure", 1, 3)]
MS9 = ["PRE-CW Completion", "CW Completion", "Site RFI", "TE Completion", "On-Air", "TE PAC",
       "KPI Acceptance", "HOTO", "Final Closure"]


def build_baseline(wb, names):
    ws = wb.create_sheet("Baseline")
    title(ws, "A1", "BASELINE · proposal (calculated) → approved versions (pasted values)")
    ws["A2"] = ("Proposal is calculated from the actual Site Release (or the planned Site Release in Master U) and the "
                "SLA Max durations. To approve: copy Proposal rows, paste as values into Approved with Version and "
                "Approval Date (see Guide).")
    ws["A2"].font = font(9, italic=True)
    ws["A2"].alignment = AL_L

    def m(c):
        return "Master!$" + c + "{r}"

    def wdy(x, n):
        return "WORKDAY.INTL(" + x + "," + n + ",7,Holidays)"
    S = "$B{r}"
    cw = "'Site Calc'!$D{r}"
    te = "'Site Calc'!$E{r}"
    rfi = "'Site Calc'!$F{r}"
    tower = m("AA") + '="Yes"'
    rt = m("O") + '="RT"'
    analysis = "OR(" + m("AP") + '="STA",' + m("AP") + '="ST")'
    mw = "N('Site Calc'!$J{r})>0"
    precw_n = ("BLink01+IF(" + m("AG") + '="Outsourced",SLA_Max_B2,0)+SLA_Max_B3+SLA_Max_B4+IF(' + analysis
               + ",SLA_Max_B5,0)+SLA_Max_B6+SLA_Max_B7")
    cols = [
        Col("Site ID", "calc", '=IF(Master!$B{r}="","",Master!$B{r}&"")', w=11),
        Col("Planned Start", "calc",
            '=IF($A{r}="","",IF(' + valid(m("V")) + ',' + m("V") + ',IF(' + valid(m("U")) + ',' + m("U") + ',"")))',
            fmt="date"),
        Col("Proposed PRE-CW Completion", "calc",
            '=IF(' + S + '="","",IF(' + cw + '="Yes",' + wdy(S, precw_n) + ',IF(' + cw + '="N/A",' + S + ',"")))',
            fmt="date"),
        Col("Proposed CW Completion", "calc",
            '=IF(' + cw + '="N/A","N/A",IF(N($C{r})=0,"",' + wdy("$C{r}", "IF(" + tower + ",SLA_Max_B8+IF(" + rt
                                                                   + ",SLA_Max_B10a,SLA_Max_B10b),SLA_Max_B9+SLA_Max_B11)") + '))',
            fmt="date"),
        Col("Proposed Site RFI", "calc",
            '=IF(' + rfi + '="N/A","N/A",IF(N(' + S + ')=0,"",IFERROR(' + wdy(
                "MAX(IF(" + cw + '="Yes",' + wdy("$D{r}", "SLA_Max_B13") + "," + S + "),IF(" + m("AC") + '="Yes",'
                + wdy(S, "BLink01+BLink02+SLA_Max_C2") + ",0),IF(" + m("AD") + '="Yes",' + wdy(S, "BLink01+BLink03+SLA_Max_C4")
                + ",0))", "SLA_Max_RFI") + ',"")))', fmt="date"),
        Col("Proposed TE Completion", "calc",
            '=IF(' + te + '="N/A","N/A",IF(N(' + S + ')=0,"",IFERROR(' + wdy(
                "IF(" + rfi + '="N/A",' + S + ",$E{r})",
                "SLA_Max_D2+SLA_Max_D3+SLA_Max_D4+BLink04+SLA_Max_D5+SLA_Max_D6") + ',"")))', fmt="date"),
        Col("Proposed On-Air", "calc",
            '=IF(' + te + '="N/A","N/A",IF(N($F{r})=0,"",' + wdy("$F{r}", "SLA_Max_D7+IF(" + mw
                                                                   + ",SLA_Max_D8,0)+SLA_Max_D9+SLA_Max_D10") + '))',
            fmt="date"),
        Col("Proposed TE PAC", "calc",
            '=IF(' + te + '="N/A","N/A",IF(N($G{r})=0,"",' + wdy(
                "$G{r}", "MAX(SLA_Max_D11,SLA_Max_D12+SLA_Max_D13+SLA_Max_D14+SLA_Max_D15)+SLA_Max_D16") + '))',
            fmt="date"),
        Col("Proposed KPI Acceptance", "calc",
            '=IF(' + te + '="N/A","N/A",IF(N($H{r})=0,"",' + wdy("$H{r}", "BLink05") + '))', fmt="date"),
        Col("Proposed HOTO", "calc",
            '=IF(' + te + '="N/A","N/A",IF(N($I{r})=0,"",' + wdy("$I{r}", "SLA_Max_D17") + '))', fmt="date"),
        Col("Proposed Final Closure", "calc",
            '=IF(N(' + S + ')=0,"",IFERROR(' + wdy(
                "MAX(IF(" + cw + '="Yes",' + wdy(wdy("$D{r}", "SLA_Max_B13") + "+CWDLPDays", "SLA_Max_B15") + ",0),IF("
                + te + '="Yes",' + wdy("EDATE($H{r},TEDLPMonths)", "BLink06") + ",0),N($J{r}))", "BLink07")
            + ',""))', fmt="date"),
        Col("Baseline Status", "calc",
            '=IF($A{r}="","",IF(N(' + S + ')=0,"No start date",IF(COUNTIFS(cBLA_SiteID,$A{r},cBLA_Current,"Yes")>0,'
            '"Approved","Proposed")))', w=11),
    ]
    # proposal range aligned with Master rows: A = row number, header row 3, SUBTOTAL + filter row 4, data from 5
    style(ws.cell(3, 1, "#"), font(9), fill(C_HDR), AL_C)
    style(ws.cell(4, 1, f"=SUBTOTAL(103,B5:B{M1})"), font(9), fill(C_CALC), AL_C)
    ws.column_dimensions["A"].width = 5
    for r in range(M0, M1 + 1):
        style(ws.cell(r, 1, r - M0 + 1), font(9), None, AL_C)
    for i, c in enumerate(cols):
        ci = i + 2
        f = shift_cols(c.f)
        style(ws.cell(3, ci, c.name), font(9), fill(C_CALC_HDR), AL_C)
        ws.cell(4, ci, f"=SUBTOTAL({102 if c.fmt else 103},{L(ci)}5:{L(ci)}{M1})")
        style(ws.cell(4, ci), font(9), fill(C_CALC), AL_C)
        ws.column_dimensions[L(ci)].width = c.w or 14.75
        for r in range(M0, M1 + 1):
            style(ws.cell(r, ci, f.format(r=r)), font(9), fill(C_CALC), AL_C,
                  DATE_FMT if c.fmt == "date" else "General")
    ws.row_dimensions[3].height = 39.95
    ws.auto_filter.ref = f"A4:{L(len(cols) + 1)}{M1}"
    from openpyxl.worksheet.table import Table, TableStyleInfo
    # links table (D39) at O3
    lc = len(cols) + 3
    for i, h in enumerate(["Link", "From", "To", "Min (wd)", "Max (wd)"]):
        style(ws.cell(3, lc + i, h), font(9), fill(C_HDR), AL_C)
    for k, (lk, f, to, mn, mx) in enumerate(LINKS):
        r = 4 + k
        for i, v in enumerate([lk, f, to, mn, mx]):
            c = ws.cell(r, lc + i, v)
            style(c, font(9), fill(C_INPUT) if i >= 3 else None, AL_C)
        names.add("BLink" + lk[1:], f"Baseline!${L(lc + 4)}${r}")
    for i, w in enumerate([7, 16, 18, 8, 8]):
        ws.column_dimensions[L(lc + i)].width = w
    t2 = Table(displayName="tblBaselineLinks", ref=f"{L(lc)}3:{L(lc + 4)}{3 + len(LINKS)}")
    t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t2)
    # approved table at U
    ac = lc + 6
    acols = [Col("Site ID", dv="SiteIDs", w=11),
             Col("Site Name", "calc", T('=IF([[Site ID]]="","",IFERROR(INDEX(cRF_SiteName,MATCH([[Site ID]],'
                                        'cRF_SiteID,0))&"",""))'), w=16)]
    acols += [Col(n, dv=("date",), fmt="date") for n in MS9]
    acols += [Col("Version", w=8), Col("Approval Date", dv=("date",), fmt="date"),
              Col("Current", "calc", T('=IF([[Site ID]]="","",IF(N([[Approval Date]])=_xlfn.MAXIFS(cBLA_ApprovalDate,'
                                       'cBLA_SiteID,[[Site ID]]),"Yes","No"))'), w=8)]
    register(ws, names, "tblBaselineApproved", "cBLA", acols, first_col=ac, sub_row=2, hdr_row=3, nrows=10)
    ws.freeze_panes = "C5"
    return ws
