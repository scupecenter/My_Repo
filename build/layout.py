"""Layout set by hand in the workbook (column widths, row heights, title merges), carried into every rebuild."""
from openpyxl.utils import range_boundaries

WIDTHS = {
    "Master": {"B": 6.44, "C": 9.0, "D": 9.0, "E": 9.0, "F": 9.66, "G": 9.0, "H": 9.0, "I": 9.0, "J": 9.0, "K": 9.0,
        "L": 9.0, "M": 11.11, "N": 9.0, "O": 9.0, "P": 9.66, "Q": 9.11, "R": 17.55, "S": 12.66, "T": 13.89, "U":
        10.89, "V": 10.89, "W": 15.33, "X": 11.55, "Y": 12.11, "AA": 7.33, "AB": 8.11, "AC": 6.44, "AD": 6.44, "AE":
        6.44, "AF": 8.0, "AG": 10.11, "AH": 9.0, "AI": 9.0, "AJ": 9.0, "AK": 9.0, "AL": 13.33, "AM": 7.55, "AN":
        11.55, "AO": 10.0, "AP": 7.55, "AQ": 7.55, "AR": 13.33, "AS": 12.44, "AT": 9.0, "AU": 8.33, "AV": 14.33,
        "AW": 14.33, "AX": 11.89, "AY": 11.55, "AZ": 10.0, "BA": 10.0, "BB": 10.11, "BC": 13.44, "BD": 10.0, "BE":
        8.44, "BF": 6.44, "BG": 10.11, "BH": 10.0, "BI": 6.44, "BJ": 7.89, "BK": 11.44, "BL": 9.89, "BM": 10.66,
        "BN": 9.11, "BQ": 11.44, "BR": 13.55, "BS": 6.44, "BT": 7.0, "BU": 10.0, "BV": 11.89, "BW": 8.89, "BX":
        10.33, "BY": 11.0, "BZ": 11.0, "CA": 16.0, "CB": 12.44, "CC": 10.0, "CD": 6.55, "CE": 11.33, "CF": 6.44,
        "CG": 13.33, "CH": 12.44, "CK": 7.0, "CL": 8.0, "CM": 7.55, "CN": 11.0, "CO": 6.55, "CP": 11.44, "CQ": 9.89,
        "CR": 10.66, "CS": 9.11, "CT": 12.11, "CU": 12.11, "CV": 14.44, "CX": 12.44, "CY": 12.33, "CZ": 11.55, "DA":
        14.0, "DC": 12.11, "DD": 11.44, "DE": 10.66, "DF": 6.66, "DG": 11.66, "DI": 9.0, "DJ": 12.11, "DK": 11.66,
        "DL": 9.0, "DM": 12.66, "DN": 11.11, "DO": 10.0, "DP": 10.0, "DQ": 14.55, "DR": 7.89, "DS": 6.44, "DT":
        9.55, "DV": 6.44, "DW": 6.44, "DX": 6.44, "DY": 7.55, "EA": 6.44, "EB": 6.44, "EC": 6.44, "ED": 10.11, "EE":
        7.89, "EF": 11.11, "EG": 6.44, "EH": 7.0, "EI": 12.66, "EJ": 11.11, "EK": 10.0, "EL": 10.0, "EM": 20.0,
        "EN": 9.0, "EO": 13.55, "EQ": 11.11, "ER": 12.44, "ES": 12.44, "ET": 13.89, "EU": 10.55, "EV": 18.11, "EW":
        16.44, "FC": 6.44, "FD": 7.0, "FE": 10.66, "FF": 14.11, "FG": 13.44, "FI": 7.0, "FN": 6.89, "FO": 10.0,
        "FP": 11.55, "FQ": 10.0, "FR": 10.11, "FS": 6.44, "FT": 10.0, "FU": 11.55, "FV": 10.0, "FW": 9.44, "FX":
        9.0, "FZ": 9.0, "GA": 7.33, "GB": 6.44, "GC": 11.11, "GD": 6.44, "GE": 6.89},
    "Dashboard": {"H": 8.66, "I": 14.89, "K": 11.89, "W": 11.55},
    "Project Status": {"B": 32.33, "C": 15.55, "D": 19.55, "G": 14.0, "I": 11.0, "J": 11.11, "M": 7.33},
    "Progress & Plan": {"C": 15.66, "D": 15.66, "G": 10.44, "L": 15.44, "M": 15.44, "N": 25.33},
    "WH Serials": {"B": 12.44},
    "Risk Register": {"A": 6.0, "F": 9.11, "L": 21.0, "P": 9.55},
    "Data Quality": {"D": 29.66},
    "SLA": {"D": 12.0, "J": 10.78, "V": 12.11, "AE": 12.0, "AP": 11.33, "BO": 11.44, "BX": 11.0},
    "Setting": {"B": 15.44, "R": 6.66, "T": 5.33, "AB": 3.33, "AD": 2.89},
}

WIDTHS["WH"] = {"S": 14.11, "AB": 13.44, "AC": 13.22, "AD": 12.11}

HEIGHTS = {
    "SLA": {3: 30.6, 4: 81.6, 5: 30.6},
    "Acquisition": {2: 36.0},
    "PR-PO": {2: 24.0},
}

MERGES = {
    "Guide": ["A1:D1"],
    "Dashboard": ["B1:I2", "B20:E20", "B70:AD70", "B95:S95"],
    "Project Status": ["A1:N2"],
    "Data Quality": ["A1:G1"],
    "WH": ["A1:BH1", "A13:BH13", "A25:BH25"],
    "WH Serials": ["A1:P1"],
    "Baseline": ["A1:M1"],
}


def apply(wb):
    for sheet, cols in WIDTHS.items():
        ws = wb[sheet]
        for col, w in cols.items():
            ws.column_dimensions[col].width = w
    for sheet, rows in HEIGHTS.items():
        for r, h in rows.items():
            wb[sheet].row_dimensions[r].height = h
    for sheet, refs in MERGES.items():
        ws = wb[sheet]
        for ref in refs:
            c1, r1, c2, r2 = range_boundaries(ref)
            for old in list(ws.merged_cells.ranges):       # a hand merge replaces any merge it overlaps
                if old.min_row <= r2 and old.max_row >= r1 and old.min_col <= c2 and old.max_col >= c1:
                    ws.unmerge_cells(str(old))
            busy = [ws.cell(r, c).coordinate for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)
                    if (r, c) != (r1, c1) and ws.cell(r, c).value not in (None, "")]
            assert not busy, (sheet, ref, busy)
            ws.merge_cells(ref)
