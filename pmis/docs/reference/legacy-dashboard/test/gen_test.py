"""Synthetic 500-site GSDN master for dashboard testing: values written over the formulas of the real template."""
import datetime as dt
import random
import sys

import openpyxl
from openpyxl.utils import column_index_from_string as CI

SRC, OUT = sys.argv[1], sys.argv[2]
random.seed(7)
wb = openpyxl.load_workbook(SRC)
m = wb["Master"]
RPT = dt.datetime(2026, 10, 15)
START = dt.datetime(2026, 1, 5)
N = 500

REG = {"Tripoli Area": (32.85, 13.15, .35), "Zawia Area": (32.70, 12.55, .35), "WM Area": (31.9, 12.0, .9),
       "Middle Area": (31.2, 16.6, 1.2), "Benghazi Area": (32.0, 20.2, .6), "GM Area": (32.4, 22.6, .7),
       "Sebha Area": (26.9, 14.4, 1.6)}
REG_W = [24, 10, 9, 16, 18, 10, 13]
OWNERS = ["AAT", "LTT", "Almadar", "Libyana", "Albounya", "LNET"]
TYPES = [("Terminal", 50), ("FN", 22), ("MFN", 12), ("HUB", 12), ("Core", 4)]
SCEN = [("New Build", 70), ("Expansion", 10), ("Cold Swap", 8), ("Hot Swap", 7), ("Relocation", 5)]
CONTR = ["Al-Bunyan Co.", "Sahara Civil", "Medina Works", "Atlas Construct", "Green Mountain Eng."]
VEND = ["Cambium", "Intracom"]
SE = ["A. Salem", "M. Faraj", "K. Omar", "H. Nasser", "R. Ali"]
ISSUES = [("Access denied by landlord", "Tower Owner"), ("GECOL meter pending", "GECOL"),
          ("HLC fiber route not ready", "HLC"), ("Awaiting DN approval", "AAT"),
          ("Material shortage (cabinet)", "Vendor"), ("Security restriction in area", "Authority"),
          ("Design change requested", "AAT"), ("Contractor team not mobilized", "Contractor")]
MS = ["V", "AW", "BP", "BZ", "CU", "DP", "EL", "ES", "EY", "FM"]          # 10 milestones
BL = ["U", "AV", "BO", "BY", "CT", "DO", "EK", "ER", "EX", "FL"]
GATE = [None, "W", "AX", "BQ", "CA", "CV", "DQ", "EM", "ET", "EZ"]
NAMES = ["Site Release", "PRE-CW Completion", "CW Completion", "Site RFI", "TE Completion", "On-Air",
         "TE PAC", "KPI Acceptance", "HOTO", "Final Closure"]
STEP = [0, 18, 32, 8, 22, 9, 25, 15, 12, 120]                               # typical working days per milestone


def pick(pairs):
    return random.choices([p for p, _ in pairs], [w for _, w in pairs])[0]


def wd(d, n):
    while n > 0:
        d += dt.timedelta(days=1)
        if d.weekday() != 4:
            n -= 1
    return d


maxc = CI("GE")
for r in range(5, 5 + N + 1):
    for c in range(1, maxc + 1):
        m.cell(r, c).value = None

for i in range(N):
    r = 5 + i
    v = {}
    reg = random.choices(list(REG), REG_W)[0]
    la, lo, sp = REG[reg]
    v.update(A=i + 1, B=f"GS{1001 + i}", C=f"OW-{random.randint(1000, 9999)}", D=f"{reg.split()[0]} {i + 1:03d}",
             E=reg, F=reg.split()[0] + (" North" if random.random() < .5 else " South"),
             G=f"CL{random.randint(1, 36):02d}", H=round(la + random.gauss(0, sp), 5),
             I=round(lo + random.gauss(0, sp * 1.3), 5), J=pick(TYPES), K=random.choice(["Outdoor", "Indoor"]),
             L=f"GS{1001 + random.randint(0, N - 1)}", M=random.choice(OWNERS), N=random.choice(["New", "Existing"]),
             O=random.choice(["RT", "GF"]), P=random.choice(["GMT", "SST", "Monopole"]), Q=random.choice([18, 24, 30, 36]),
             X="PH1" if i < 380 else "PH2", Y=f"Batch {min(3, i // 130)}", Z=pick(SCEN))
    swap = v["Z"] in ("Cold Swap", "Hot Swap")
    tower = (v["N"] == "New") and not swap
    v.update(AA="Yes" if tower else "N/A", AB="N/A" if swap and random.random() < .7 else "Yes",
             AC="Yes" if random.random() < .35 else "N/A", AD="Yes" if random.random() < .28 else "N/A",
             AE="Yes" if random.random() < .9 else "N/A")
    v["AF"] = "N/A" if v["AE"] == "Yes" else "Yes"
    cw = "Yes" in (v["AA"], v["AB"])
    rfi = cw or "Yes" in (v["AC"], v["AD"])
    v.update(AG=random.choice(["Outsourced", "In-House"]) if cw else "N/A", AJ=random.choice(CONTR) if cw else None,
             AM=random.choice(SE), CD=random.choice(VEND), CF=random.choice(SE))
    # progress depth: earlier batches further along
    base = {"Batch 0": 7.5, "Batch 1": 5.5, "Batch 2": 3.5, "Batch 3": 1.6}[v["Y"]]
    k = max(0, min(10, int(random.gauss(base, 2.0))))
    status = "In Progress" if k else "Not Started"
    if random.random() < .025:
        status = "Cancelled"
    rel = START + dt.timedelta(days=int(random.uniform(0, 120) + (i // 130) * 35))
    plan = rel - dt.timedelta(days=random.randint(-5, 10))
    d = rel
    na = [False] * 10
    if not cw:
        na[1] = na[2] = True
    if not rfi:
        na[3] = True
    for j in range(10):
        late = random.gauss(0.25, .45)
        dur = max(1, int(STEP[j] * (1 + late))) if j else 0
        pdur = STEP[j]
        if na[j]:
            v[MS[j]] = "N/A"
            v[BL[j]] = "N/A" if j else plan
            if GATE[j]:
                v[GATE[j]] = "N/A"
            continue
        d = wd(d, dur) if j else d
        plan = wd(plan, pdur) if j else plan
        v[BL[j]] = plan if j or True else None
        if j < k and d <= RPT:
            v[MS[j]] = d
            if GATE[j]:
                v[GATE[j]] = wd(d, -0) - dt.timedelta(days=random.randint(1, max(2, dur // 2)))
        else:
            if j == k or (j > 0 and all(isinstance(v.get(MS[x]), dt.datetime) or v.get(MS[x]) == "N/A"
                                        for x in range(j))):
                if GATE[j] and random.random() < .62:
                    v[GATE[j]] = min(RPT, wd(d, -dur // 2))
            k = min(k, j)
            break
    done_n = sum(isinstance(v.get(c), dt.datetime) for c in MS)
    names_done = [NAMES[j] for j in range(10) if isinstance(v.get(MS[j]), dt.datetime)]
    v["FX"] = names_done[-1] if names_done else "Not Started"
    nxt = next((NAMES[j] for j in range(10) if v.get(MS[j]) not in ("N/A",) and not isinstance(v.get(MS[j]), dt.datetime)), "")
    v["FY"] = nxt
    if status != "Cancelled":
        if isinstance(v.get("FM"), dt.datetime):
            status = "Completed"
        elif done_n and random.random() < .04:
            status = "On Hold"
        elif not done_n:
            status = "Not Started"
    v["FZ"] = status
    # detail fields
    def dd(c):
        return isinstance(v.get(c), dt.datetime)
    if dd("V"):
        v["S"] = v["V"] - dt.timedelta(days=random.randint(20, 60))
        v["T"] = v["V"] - dt.timedelta(days=random.randint(3, 15))
        v["R"] = "WL/MW/IP/PWR"
    if cw and dd("V"):
        v["AH"] = "Done" if dd("AW") or random.random() < .6 else "Pending"
        v["AI"] = "Done" if dd("AW") else ("Pending" if random.random() < .5 else "Done")
        v["AK"] = "Team " + random.choice("ABCD")
        v["AL"] = v["V"] + dt.timedelta(days=random.randint(3, 12)) if dd("AW") else None
        v["AN"] = v["V"] + dt.timedelta(days=5)
        v["AO"] = v["V"] + dt.timedelta(days=random.randint(4, 14)) if dd("AW") else None
        v["AP"] = random.choice(["STA", "ST", "N/A"])
        v["AQ"] = "Pass" if dd("AW") else None
    if dd("AW"):
        v["AY"] = v["AW"] + dt.timedelta(days=4)
        v["AZ"] = v["AW"] + dt.timedelta(days=random.randint(2, 9))
        for c in ("BA", "BB", "BC"):
            v[c] = ("Done" if dd("BP") or random.random() < .5 else "Pending") if v["AA"] == "Yes" else "N/A"
        for c in ("BD", "BE", "BF", "BG", "BH", "BI", "BJ"):
            v[c] = ("Done" if dd("BP") or random.random() < .55 else "Pending") if v["AB"] == "Yes" else "N/A"
        v["BK"] = v["AW"] + dt.timedelta(days=6)
        v["BL"] = v["AZ"] + dt.timedelta(days=2)
        v["BM"] = v["BL"] + dt.timedelta(days=STEP[2])
        if dd("BP"):
            v["BN"] = v["BP"]
    if dd("BP"):
        v["BR"] = v["BP"] + dt.timedelta(days=3)
        snag = random.random() < .3
        v["BS"] = "Yes" if snag else "N/A"
        v["BT"] = ("Yes" if dd("BZ") or random.random() < .5 else None) if snag else "N/A"
        v["BU"] = "Yes" if random.random() < .8 else "N/A"
        v["BV"] = v["BP"] + dt.timedelta(days=random.randint(2, 12)) if dd("BZ") or random.random() < .5 else None
    v["BW"] = ("Done" if dd("BZ") or (dd("BP") and random.random() < .5) else "Pending") if v["AC"] == "Yes" else "N/A"
    v["BX"] = ("Done" if dd("BZ") or (dd("BP") and random.random() < .4) else "Pending") if v["AD"] == "Yes" else "N/A"
    if dd("BZ") or (v.get("BZ") == "N/A" and dd("V")):
        ref = v["BZ"] if dd("BZ") else v["V"]
        v["CB"] = ref + dt.timedelta(days=random.randint(2, 10)) if (dd("CU") or random.random() < .7) else None
        if v.get("CB"):
            v["CC"] = v["CB"] + dt.timedelta(days=2); v["CE"] = "Team " + random.choice("XYZ")
            v["CG"] = v["CC"] + dt.timedelta(days=2)
            v["CH"] = v["CG"] + dt.timedelta(days=random.randint(1, 6)) if (dd("CU") or random.random() < .6) else None
            v["CI"] = v["CG"] + dt.timedelta(days=6)
            if v.get("CH"):
                v["CJ"] = v["CH"] + dt.timedelta(days=random.randint(2, 8)) if (dd("CU") or random.random() < .6) else None
            for c in ("CK", "CL", "CM", "CN", "CO"):
                v[c] = "Done" if dd("CU") else ("Pending" if v.get("CJ") else None)
            v["CP"] = v["CI"] + dt.timedelta(days=2)
            if v.get("CJ"):
                v["CQ"] = v["CJ"] + dt.timedelta(days=1)
            v["CR"] = v["CP"] + dt.timedelta(days=5)
            if dd("CU"):
                v["CS"] = v["CU"]
    if dd("CU"):
        v["CW"] = v["CU"] + dt.timedelta(days=2) if dd("DP") or random.random() < .6 else None
        if v.get("CW"):
            v["CX"] = v["CW"] + dt.timedelta(days=1)
            rsl = round(random.gauss(-48, 4), 1)
            v["CY"] = rsl; v["CZ"] = round(rsl + 48, 1)
            v["DA"] = v["CX"] + dt.timedelta(days=2) if dd("DP") or random.random() < .6 else None
        if swap:
            v["DB"] = v["CU"]; v["DC"] = v["CU"] + dt.timedelta(days=3)
            v["DF"] = ("Success" if random.random() < .85 else "Rolled Back") if dd("DP") else None
            v["DI"] = 1 if v.get("DF") == "Rolled Back" else 0
        v["DM"] = v["CU"] + dt.timedelta(days=STEP[5])
        if dd("DP"):
            v["DN"] = v["DP"]
    if dd("DP"):
        v["DR"] = v["DP"] + dt.timedelta(days=3) if dd("EL") or random.random() < .7 else None
        v["DS"] = ("Pass" if random.random() < .9 else "Fail") if v.get("DR") else None
        v["DT"] = v["DP"] + dt.timedelta(days=4) if dd("EL") or random.random() < .6 else None
        for c in ("DU", "DV", "DW", "DX"):
            v[c] = "Done" if dd("EL") or random.random() < .5 else "Pending"
        v["DY"] = v["DP"] + dt.timedelta(days=8) if dd("EL") or random.random() < .5 else None
        for c in ("DZ", "EA", "EB", "EC", "ED"):
            v[c] = "Done" if dd("EL") or random.random() < .4 else "Pending"
        v["EE"] = v["DP"] + dt.timedelta(days=14) if dd("EL") or random.random() < .3 else None
        v["EF"] = v["DP"] + dt.timedelta(days=16) if dd("EL") else None
        snag = random.random() < .25
        v["EG"] = "Yes" if snag else "N/A"
        v["EH"] = ("Yes" if dd("EL") else None) if snag else "N/A"
        v["EI"] = v["DP"] + dt.timedelta(days=STEP[6])
        if dd("EL"):
            v["EJ"] = v["EL"]
            v["FF"] = v["EL"]; v["FG"] = v["EL"] + dt.timedelta(days=365)
    if dd("EL"):
        v["EN"] = v["G"]; v["EO"] = v["EL"] + dt.timedelta(days=5) if dd("ES") or random.random() < .5 else None
        v["EQ"] = "OK"
    if dd("ES"):
        v["EP"] = v["ES"]; v["EU"] = v["ES"] + dt.timedelta(days=3)
        v["EV"] = v["ES"] + dt.timedelta(days=6) if dd("EY") or random.random() < .5 else None
        v["EW"] = v["EY"] if dd("EY") else None
    if dd("BV"):
        v["FA"] = v["BV"]; v["FB"] = v["BV"] + dt.timedelta(days=90)
        v["FC"] = "N/A"; v["FD"] = "N/A"
        v["FE"] = v["FB"] + dt.timedelta(days=5) if v["FB"] < RPT and random.random() < .6 else None
    if v["Z"] in ("Cold Swap", "Relocation") and dd("DP"):
        v.update(FN="Yes", FO=random.choice(CONTR), FP=v["DP"] + dt.timedelta(days=20))
        if v["FP"] < RPT and random.random() < .6:
            v["FQ"] = v["FP"] + dt.timedelta(days=3); v["FR"] = v["FQ"] + dt.timedelta(days=4)
        v.update(FS="Yes", FT=random.choice(VEND), FU=v["DP"] + dt.timedelta(days=15))
        if v["FU"] < RPT and random.random() < .7:
            v["FV"] = v["FU"] + dt.timedelta(days=2); v["FW"] = v["FV"] + dt.timedelta(days=3)
    if status in ("In Progress", "On Hold") and random.random() < .2:
        iss, own = random.choice(ISSUES)
        v.update(GA=iss, GB=own, GC=RPT - dt.timedelta(days=random.randint(2, 120)), GD="Open")
    elif random.random() < .05:
        iss, own = random.choice(ISSUES)
        v.update(GA=iss, GB=own, GC=RPT - dt.timedelta(days=random.randint(30, 200)), GD="Closed")
    for c, val in v.items():
        cell = m.cell(r, CI(c))
        cell.value = val
        if isinstance(val, dt.datetime):
            cell.number_format = "yyyy-mm-dd"

# Setting · report date
s = wb["Setting"]
s["B3"] = RPT
s["B3"].number_format = "yyyy-mm-dd"
# Project Status · overall and by activity
p = wb["Project Status"]
p["H9"] = RPT
act = [(.92, .95), (.86, .9), (.71, .78), (.58, .66), (.41, .5), (.33, .42), (.21, .3), (.08, .12)]
wts = [12, 8, 5, 25, 20, 12, 13, 5]
tot_a = sum(w * a for w, (a, b) in zip(wts, act)) / 100
tot_p = sum(w * b for w, (a, b) in zip(wts, act)) / 100
p["A12"], p["C12"], p["E12"], p["G12"] = tot_a, tot_p, tot_a / tot_p, .93
for k, (w, (a, b)) in enumerate(zip(wts, act)):
    rr = 18 + k
    p.cell(rr, 5).value = a; p.cell(rr, 6).value = w * a / 100; p.cell(rr, 7).value = b
    p.cell(rr, 8).value = w * b / 100; p.cell(rr, 9).value = a - b; p.cell(rr, 10).value = a / b
p["F26"], p["H26"], p["I26"], p["J26"] = tot_a, tot_p, tot_a - tot_p, tot_a / tot_p
for row in p.iter_rows(min_row=27):
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.value = None
# SLA · per-site durations (completed block D.., open block AW..)
sl = wb["SLA"]
heads = {}
for c in range(4, 95):
    h = sl.cell(3, c).value
    if h:
        heads[c] = h
mx = {}
for c, h in heads.items():
    try:
        mx[c] = float(sl.cell(7, c).value)
    except (TypeError, ValueError):
        f = str(sl.cell(7, c).value or "")
        mx[c] = None
for rr in range(8, 13):
    for c in range(1, 120):
        if isinstance(sl.cell(rr, c).value, str) and sl.cell(rr, c).value.startswith("="):
            sl.cell(rr, c).value = None
for i in range(N):
    rr = 13 + i
    mr = 5 + i
    sid = m.cell(mr, 2).value
    for c in range(1, 120):
        sl.cell(rr, c).value = None
    sl.cell(rr, 1).value = i + 1
    sl.cell(rr, 2).value = sid
    sl.cell(rr, 3).value = m.cell(mr, 5).value
    vv = {cc: m.cell(mr, CI(cc)).value for cc in MS}
    done = lambda cc: isinstance(vv[cc], dt.datetime)
    for c, h in heads.items():
        code = h.split(" · ")[0]
        open_blk = c >= 49
        if c >= 94:
            continue
        lim = {"A": 20, "B": 6, "C": 10, "R": 3, "E": 30, "D": 4, "S": 5}.get(code[0], 5)
        need = ("V" if code.startswith("A") else "BP" if code.startswith(("B", "E2E-B")) else "BZ"
                if code.startswith(("C", "RFI")) else "DP" if code in ("D-1", "D-2", "D-3", "D-4", "D-5", "D-6", "D-7",
                                                                     "D-8", "D-9", "D-10", "E2E-D1") else "EL"
                if code.startswith("D-1") else "EY" if code in ("D-17", "E2E-D2") else "DP")
        if code.startswith("S-") and m.cell(mr, CI("Z")).value not in ("Cold Swap", "Hot Swap"):
            continue
        if code.startswith("C-2") and m.cell(mr, CI("AC")).value != "Yes":
            continue
        if code.startswith("C-4") and m.cell(mr, CI("AD")).value != "Yes":
            continue
        if code.startswith(("B", "E2E-B")) and vv["BP"] == "N/A":
            continue
        if not open_blk and done(need):
            sl.cell(rr, c).value = max(0, int(random.gauss(lim * .8, lim * .45)))
        if open_blk and not done(need) and m.cell(mr, CI("FZ")).value == "In Progress" and random.random() < .12:
            sl.cell(rr, c).value = max(1, int(random.gauss(lim * 1.1, lim * .7)))
# Risk Register
rk = wb["Risk Register"]
risks = [("Landlord access disputes delay CW start", "Permit / Access", 4, 4, "Escalate via AAT legal"),
         ("GECOL meter installation backlog", "Power", 4, 5, "Weekly GECOL coordination"),
         ("HLC fiber routes late in GM Area", "Fiber", 3, 4, "Microwave fallback design"),
         ("Cabinet supply shortfall Q4", "Supply Chain", 3, 3, "Expedite second shipment"),
         ("Security restrictions south region", "External / Security", 3, 5, "Escort plan with authority"),
         ("Contractor team shortage", "Contractor", 2, 3, "Add two teams"),
         ("Late DN approvals", "Governance", 3, 2, "Daily DN review"),
         ("KPI acceptance criteria disputes", "Acceptance", 2, 4, "Joint KPI workshop"),
         ("Weather window for tower erection", "External / Weather", 2, 2, "Re-sequence GF towers"),
         ("Data quality of RF records", "Data Integrity", 2, 2, "Weekly DQ review")]
for k, (t, cat, pp, ii, mit) in enumerate(risks):
    rr = 3 + k
    sc = pp * ii
    lvl = "Critical" if sc >= 15 else "High" if sc >= 10 else "Medium" if sc >= 5 else "Low"
    vals = [f"R-{k + 1:03d}", t, cat, "Project", random.randint(2, 40), pp, ii, sc, lvl, "Mitigate", mit,
            random.choice(["PMO", "CW Lead", "TE Lead"]), random.choice(SE),
            RPT + dt.timedelta(days=random.randint(-20, 40)), "Open" if k < 8 else "Closed",
            random.choice(["Increasing", "Stable", "Decreasing"]), RPT - dt.timedelta(days=7), RPT + dt.timedelta(days=7)]
    for c, val in enumerate(vals, 1):
        rk.cell(rr, c).value = val
# Data Quality · rows per check
dq = wb["Data Quality"]
for rr in range(4, 31):
    dq.cell(rr, 5).value = random.choice([0, 0, 0, 0, 1, 2, 3, 6])
# Quality Inspections
qi = wb["Quality Inspections"]
for k in range(30):
    rr = 3 + k
    sid = f"GS{1001 + random.randint(0, 300)}"
    sev = random.choice(["Minor", "Minor", "Major", "Critical"])
    closed = random.random() < .6
    vals = {2: f"Q-{k + 1:03d}", 3: random.choice(["Snag", "NCR", "QC", "ITP"]), 4: sid,
            6: random.choice(["CW", "TE"]), 10: sev, 13: RPT - dt.timedelta(days=random.randint(5, 90)),
            14: RPT + dt.timedelta(days=random.randint(-30, 20)), 15: "Cleared" if closed else "Open",
            16: RPT - dt.timedelta(days=random.randint(1, 4)) if closed else None, 22: "Closed" if closed else "Open",
            23: "Yes" if sev == "Critical" and not closed else "No"}
    for c, val in vals.items():
        qi.cell(rr, c).value = val
wb.save(OUT)
print("saved", OUT)
